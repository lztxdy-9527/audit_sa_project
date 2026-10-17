#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""TASK-052：网页审阅记录合并（校验 → 合并索引 → 《审阅记录.xlsx》 → 重新生成站点/底稿）。

命令行：
    python -B scripts/apply_review_notes.py \
        --index <trace_index.json> --review-notes <审阅记录.json> --out-dir <新证据根>

只读原索引、原底稿与机器证据；校验审阅记录（版本绑定、记录存在、状态枚举、必填、长度、注入），
把审阅字段合并进索引副本，生成《审阅记录.xlsx》、合并索引、以及用合并索引重新生成的站点与底稿。
**不覆盖**原索引/原底稿；未填写记录保持 `待审阅`；审阅记录版本 = 记录文件 SHA 短标识。

边界：审阅记录是**人类输入**；本工具不采集、不虚构、不覆盖机器证据；不构成审计结论或终判。
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
TEST_ROOT = ROOT / "data/runtime_initial_judgment_recording_tests"
REVIEW_VERSION = "TASK-052-review-1.0"
GENERATOR = "scripts/apply_review_notes.py"
STATUSES = ("待审阅", "已审阅-相符", "已审阅-不符", "需补充资料")
RECORD_TYPES = ("case", "ledger", "voucher", "source")
MAX_ISSUE = 2000
MAX_SUGGESTION = 2000
MAX_REVIEWER = 100
DATE_PATTERN = re.compile(r"\A\d{4}-\d{2}-\d{2}\Z")
HEADER_FILL = PatternFill("solid", fgColor="F6F8FA")
THIN = Side(style="thin", color="D0D7DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


class MergeError(Exception):
    def __init__(self, kind, message, field=None):
        super().__init__(message)
        self.kind = kind  # "usage" → exit 2；"content" → exit 1
        self.field = field


def sha_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha_file(path: Path) -> str:
    return sha_bytes(path.read_bytes())


def rel(path: Path) -> str:
    try:
        return Path(path).resolve().relative_to(ROOT.resolve()).as_posix()
    except (ValueError, OSError):
        return str(path)


def ident(path: Path) -> dict:
    raw = path.read_bytes()
    return {"relative": rel(path), "byte_size": len(raw), "sha256": sha_bytes(raw)}


def load_module(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def load_index(path: Path) -> dict:
    if not path.is_file():
        raise MergeError("usage", "索引不存在：%s" % path)
    try:
        return json.loads(path.read_bytes())
    except ValueError as exc:
        raise MergeError("content", "索引不是合法内容：%s" % exc, field="trace_index.json")


def record_key(record_type, record_id):
    return (record_type, record_id)


def index_records(index):
    records = {}
    for case in index.get("cases") or []:
        records[record_key("case", case.get("case_id"))] = case
    for row in index.get("ledger") or []:
        records[record_key("ledger", row.get("ledger_id"))] = row
    for row in index.get("vouchers") or []:
        records[record_key("voucher", row.get("voucher_id"))] = row
    for row in index.get("source_rows") or []:
        record_id = "%s!%s" % (row.get("sheet"), row.get("excel_row"))
        records[record_key("source", record_id)] = row
    return records


def validate_review(review, index):
    if type(review) is not dict:
        raise MergeError("content", "审阅记录必须是对象", field="review_notes")
    if review.get("review_version") != REVIEW_VERSION:
        raise MergeError("content", "审阅记录版本不受支持",
                         field="review_notes.review_version")
    source = (index.get("meta") or {}).get("source_identity") or {}
    expected = (source.get("sha256") or "")[:12]
    if review.get("source_version") != expected:
        raise MergeError("content",
                         "审阅结论对应旧资料版本（记录=%s，当前索引=%s）"
                         % (review.get("source_version"), expected),
                         field="review_notes.source_version")
    entries = review.get("records")
    if type(entries) is not list:
        raise MergeError("content", "records 必须是数组", field="review_notes.records")
    records = index_records(index)
    seen = set()
    for position, entry in enumerate(entries):
        at = "review_notes.records[%d]" % position
        if type(entry) is not dict:
            raise MergeError("content", "记录必须是对象", field=at)
        record_type = entry.get("record_type")
        record_id = entry.get("record_id")
        if record_type not in RECORD_TYPES:
            raise MergeError("content", "record_type 非法：%r" % record_type,
                             field=at + ".record_type")
        if type(record_id) is not str or record_key(record_type, record_id) not in records:
            raise MergeError("content", "未知记录：%s %r" % (record_type, record_id),
                             field=at + ".record_id")
        if record_key(record_type, record_id) in seen:
            raise MergeError("content", "同一记录重复出现", field=at + ".record_id")
        seen.add(record_key(record_type, record_id))
        status = entry.get("status")
        if status not in STATUSES:
            raise MergeError("content", "审阅状态非法：%r" % status, field=at + ".status")
        for field in ("issue", "suggestion", "reviewer", "date"):
            value = entry.get(field)
            if value is not None and type(value) is not str:
                raise MergeError("content", "%s 必须是字符串" % field, field=at + "." + field)
            if type(value) is str and ("<" in value or ">" in value):
                raise MergeError("content", "%s 不得包含 HTML 片段" % field,
                                 field=at + "." + field)
        issue = entry.get("issue") or ""
        suggestion = entry.get("suggestion") or ""
        reviewer = entry.get("reviewer") or ""
        date = entry.get("date") or ""
        if status != "待审阅":
            if not issue.strip():
                raise MergeError("content", "非待审阅必须填写审计问题描述",
                                 field=at + ".issue")
            if not reviewer.strip():
                raise MergeError("content", "非待审阅必须填写审阅人", field=at + ".reviewer")
            if not DATE_PATTERN.match(date):
                raise MergeError("content", "非待审阅必须填写审阅日期（YYYY-MM-DD）",
                                 field=at + ".date")
        if len(issue) > MAX_ISSUE:
            raise MergeError("content", "审计问题描述超长", field=at + ".issue")
        if len(suggestion) > MAX_SUGGESTION:
            raise MergeError("content", "建议处理超长", field=at + ".suggestion")
        if len(reviewer) > MAX_REVIEWER:
            raise MergeError("content", "审阅人超长", field=at + ".reviewer")
    return entries


def merge_index(index, entries):
    records = index_records(index)
    for entry in entries:
        record = records[record_key(entry["record_type"], entry["record_id"])]
        record["review"] = {
            "status": entry.get("status"),
            "issue": entry.get("issue") or None,
            "suggestion": entry.get("suggestion") or None,
            "reviewer": entry.get("reviewer") or None,
            "date": entry.get("date") or None,
        }
    meta = index.setdefault("meta", {})
    meta["review"] = {
        "status": "已合并审阅记录",
        "review_version": REVIEW_VERSION,
        "note": "审阅记录为人类输入；未填写记录保持待审阅；机器证据未被覆盖。",
    }
    return index


def build_review_workbook(index, entries, review_identity, out_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "审阅记录"
    headers = ["编号", "事项或记录", "审阅状态", "审计问题描述", "建议处理",
               "审阅人", "审阅日期", "对应资料版本", "审阅记录版本"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
    records = index_records(index)
    order = [("case", case.get("case_id")) for case in index.get("cases") or []]
    order += [("ledger", row.get("ledger_id")) for row in index.get("ledger") or []]
    order += [("voucher", row.get("voucher_id")) for row in index.get("vouchers") or []]
    order += [("source", "%s!%s" % (row.get("sheet"), row.get("excel_row")))
              for row in index.get("source_rows") or []]
    for record_type, record_id in order:
        record = records[record_key(record_type, record_id)]
        review = record.get("review") or {}
        label = record.get("title") or record.get("ledger_id") or record.get(
            "voucher_id") or ("%s!%s" % (record.get("sheet"), record.get("excel_row")))
        ws.append([record.get("case_id") or record_id,
                   str(label),
                   review.get("status") or "待审阅",
                   review.get("issue") or "",
                   review.get("suggestion") or "",
                   review.get("reviewer") or "",
                   review.get("date") or "",
                   ((index.get("meta") or {}).get("source_identity") or {}).get("sha256", "")[:12],
                   review_identity["sha256"][:12]])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [16, 40, 14, 40, 30, 12, 12, 14, 14]
    for index_, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index_)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), ws.max_row)
    wb.properties.created = __import__("datetime").datetime(2026, 10, 17)
    wb.properties.modified = __import__("datetime").datetime(2026, 10, 17)
    wb.save(out_path)
    workbook_module = load_module(ROOT / "scripts/build_excel_workbook.py",
                                  "task052_workbook_normalize")
    workbook_module.normalize_workbook(out_path)


def main(argv=None):
    parser = argparse.ArgumentParser(description="TASK-052 审阅记录合并")
    parser.add_argument("--index", required=True)
    parser.add_argument("--review-notes", dest="review_notes", required=True)
    parser.add_argument("--out-dir", dest="out_dir", required=True)
    args = parser.parse_args(argv)
    index_path = Path(args.index)
    review_path = Path(args.review_notes)
    out_dir = Path(args.out_dir)
    try:
        if not index_path.is_file():
            raise MergeError("usage", "索引不存在：%s" % index_path)
        if not review_path.is_file():
            raise MergeError("usage", "审阅记录不存在：%s" % review_path)
        try:
            out_dir.resolve().relative_to(TEST_ROOT.resolve())
        except ValueError:
            raise MergeError("usage", "--out-dir 必须位于 %s 之下" % TEST_ROOT)
        if out_dir.exists():
            raise MergeError("usage", "--out-dir 必须从未存在：%s" % out_dir)
        if not out_dir.name.startswith("task052_"):
            raise MergeError("usage", "--out-dir 目录名必须以 task052_ 开头")
        index = load_index(index_path)
        try:
            review = json.loads(review_path.read_bytes())
        except ValueError as exc:
            raise MergeError("content", "审阅记录不是合法内容：%s" % exc,
                             field="review_notes")
        entries = validate_review(review, index)
        merged = merge_index(json.loads(index_path.read_bytes()), entries)
        review_identity = ident(review_path)
        out_dir.mkdir(parents=True, exist_ok=False)
        merged_path = out_dir / "trace_index.json"
        payload = json.dumps(merged, ensure_ascii=False, indent=2,
                             sort_keys=True).encode("utf-8") + b"\n"
        merged_path.write_bytes(payload)
        workbook_path = out_dir / "审阅记录.xlsx"
        build_review_workbook(merged, entries, review_identity, workbook_path)
        site_dir = out_dir / "task049_site"
        workbook_dir = out_dir / "task050_workbook"
        for command in (
                [sys.executable, "-B", str(ROOT / "scripts/build_drilldown_site.py"),
                 "--index", str(merged_path), "--out-dir", str(site_dir)],
                [sys.executable, "-B", str(ROOT / "scripts/build_excel_workbook.py"),
                 "--index", str(merged_path), "--out-dir", str(workbook_dir)]):
            completed = subprocess.run(command, cwd=str(ROOT),
                                       stdout=subprocess.PIPE, stderr=subprocess.PIPE)
            if completed.returncode != 0:
                raise MergeError("content", "重新生成失败：%s"
                                 % completed.stdout.decode("utf-8", "replace")[-300:],
                                 field="regenerate")
        if not (workbook_dir / "证据链底稿.xlsx").is_file():
            raise MergeError("content", "底稿未生成", field="底稿")
        if not (site_dir / "index.html").is_file():
            raise MergeError("content", "站点未生成", field="site")
        source = {
            "task": "TASK-052",
            "out_dir": rel(out_dir),
            "index": ident(index_path),
            "review_notes": review_identity,
            "merged_index": ident(merged_path),
            "review_workbook": ident(workbook_path),
            "site_index": ident(site_dir / "index.html"),
            "site_pages": len(list(site_dir.glob("*.html"))),
            "workbook": ident(workbook_dir / "证据链底稿.xlsx"),
            "review_version": REVIEW_VERSION,
            "source_version": (review.get("source_version") or ""),
            "note": "审阅记录为人类输入；未填写记录保持待审阅；不覆盖原索引/原底稿/机器证据。",
        }
        (out_dir / "merge_source.json").write_bytes(
            json.dumps(source, ensure_ascii=False, indent=2,
                       sort_keys=True).encode("utf-8") + b"\n")
        print(json.dumps({"status": "PASS", "out": str(out_dir),
                          "records": len(entries),
                          "merged_index": source["merged_index"]["byte_size"],
                          "workbook": source["review_workbook"]["byte_size"],
                          "site_pages": source["site_pages"]}, ensure_ascii=False))
        return 0
    except MergeError as exc:
        print(json.dumps({"status": "ERROR", "kind": exc.kind, "error": str(exc),
                          "field": exc.field}, ensure_ascii=False))
        return 1 if exc.kind == "content" else 2


if __name__ == "__main__":
    raise SystemExit(main())
