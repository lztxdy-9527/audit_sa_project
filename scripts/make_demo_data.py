#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""TASK-002 合成数据生成器（开发期造数与独立答案册）。

本文件只负责“按已知场景构造数据并登记答案”，不实现 R01—R04，不做任何审计判断，
也不导入或调用未来任何检测规则。答案来自“正在构造哪种已知场景”的登记信息。

输出（全部写入 --out-dir 指定的新目录，目录必须位于 data/generated 内且尚不存在）：
  input/erp_export.xlsx                  主演示工作簿（suppliers/payables/payments/allocations）
  fixtures/*.xlsx                        22 份小型验收工作簿（C01—C13、C06b/C07b/C09b/C10b、C14a—C14e）
  evaluation/ground_truth.json           主演示答案册（与输入分离）
  evaluation/ground_truth.md             人工可读的答案说明
  evaluation/fixture_expectations.json   小工作簿的独立预期
  manifest.json                          合成声明、种子、哈希、版本、实际行数、业务内容指纹
  checks/                                校验报告输出目录（由 check_demo_data.py 写入）

退出码：0 成功；2 参数或路径问题（越界、目录已存在等，写入任何内容前停止）。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import random
import sys
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
GENERATED_ROOT = (PROJECT_ROOT / "data" / "generated").resolve()
SPEC_PATH = PROJECT_ROOT / "docs" / "SPEC_V1.md"
SELF_PATH = Path(__file__).resolve()

COMPANY_NAME = "澄明商贸有限公司"
PERIOD_START = date(2026, 9, 1)
PERIOD_END = date(2026, 9, 30)

FIXTURE_AP_DATE = "2026-09-05"
FIXTURE_PAY_DATE = "2026-09-10"
FIXTURE_REFUND_DATE = "2026-09-20"

SHEET_COLUMNS: dict[str, list[str]] = {
    "suppliers": ["supplier_id", "supplier_name"],
    "payables": ["ap_id", "ap_number", "supplier_id", "ap_date", "payable_amount", "currency"],
    "payments": [
        "payment_id",
        "transaction_type",
        "original_payment_id",
        "payee_supplier_id",
        "payment_date",
        "amount",
        "currency",
        "settlement_status",
        "bank_reference",
    ],
    "allocations": [
        "allocation_id",
        "payment_id",
        "ap_id",
        "allocated_amount",
        "original_allocation_id",
    ],
}
SHEET_ORDER = ["suppliers", "payables", "payments", "allocations"]

SYNTHETIC_NOTE = (
    "本工作簿为合成数据，仅用于采购付款审计辅助工具的开发与验收演示，"
    "不代表任何真实企业、真实交易、真实银行回单或真实审计发现。"
)

SUPPLIER_NAMES = [
    "青禾文具有限公司",
    "朗行办公设备有限公司",
    "远溪纸品有限公司",
    "珩通耗材有限公司",
    "沐阳商贸有限公司",
    "岑石印务有限公司",
    "星屿电子科技有限公司",
    "澜庭清洁用品有限公司",
    "翊铭办公家具有限公司",
    "芷兰包装材料有限公司",
    "承砚文化用品有限公司",
    "汇茂数码有限公司",
    "渝川五金有限公司",
    "霖和办公用品有限公司",
    "恒玖文体有限公司",
    "若水纸业有限公司",
    "岐峰贸易有限公司",
    "砚山科技服务有限公司",
    "泊远日化有限公司",
    "观澜实业有限公司",
]

# ---------------------------------------------------------------------------
# 基础工具
# ---------------------------------------------------------------------------


def fail(message: str) -> "None":
    print(f"ERROR: {message}", file=sys.stderr)
    sys.exit(2)


def yuan(cents: int) -> str:
    """整数分 -> 两位小数文本（本版金额一律以文本写入，避免 Excel 浮点求和）。"""
    if cents < 0:
        raise ValueError("本版金额必须为正数")
    return f"{cents // 100}.{cents % 100:02d}"


def iso(value: date) -> str:
    return value.isoformat()


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def json_text(payload) -> str:
    return json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


class Table:
    """按固定列顺序累积记录，并返回该记录在 Excel 中的行号（表头占第 1 行）。"""

    def __init__(self, name: str, columns: list[str] | None = None):
        self.name = name
        self.columns = list(columns if columns is not None else SHEET_COLUMNS[name])
        self.rows: list[list] = []

    def add(self, record: dict) -> int:
        row = []
        for column in self.columns:
            if column not in record:
                raise KeyError(f"{self.name} 记录缺少字段 {column}: {record}")
            row.append(record[column])
        self.rows.append(row)
        return len(self.rows) + 1


class Dataset:
    def __init__(self):
        self.tables = {name: Table(name) for name in SHEET_ORDER}
        self.counters = {"ap": 0, "pm": 0, "al": 0, "bank": 0, "ap_number": 0}
        self.clues: list[dict] = []
        self.non_hit_scenarios: list[dict] = []
        self.composition: dict[str, int] = {}

    def table(self, name: str) -> Table:
        return self.tables[name]

    def next_ap_id(self) -> str:
        self.counters["ap"] += 1
        return f"AP{self.counters['ap']:05d}"

    def next_payment_id(self) -> str:
        self.counters["pm"] += 1
        return f"PM{self.counters['pm']:05d}"

    def next_allocation_id(self) -> str:
        self.counters["al"] += 1
        return f"AL{self.counters['al']:05d}"

    def next_bank_reference(self) -> str:
        self.counters["bank"] += 1
        return f"BK2026{self.counters['bank']:08d}"

    def next_ap_number(self) -> str:
        self.counters["ap_number"] += 1
        return f"CG2026{self.counters['ap_number']:05d}"


def add_supplier(ds: Dataset, supplier_id: str, supplier_name: str) -> int:
    return ds.table("suppliers").add(
        {"supplier_id": supplier_id, "supplier_name": supplier_name}
    )


def add_payable(
    ds: Dataset,
    supplier_id: str,
    amount_cents: int,
    ap_date: date,
    ap_number: str | None = None,
) -> tuple[str, int]:
    ap_id = ds.next_ap_id()
    if ap_number is None:
        ap_number = ds.next_ap_number()
    row = ds.table("payables").add(
        {
            "ap_id": ap_id,
            "ap_number": ap_number,
            "supplier_id": supplier_id,
            "ap_date": iso(ap_date),
            "payable_amount": yuan(amount_cents),
            "currency": "CNY",
        }
    )
    return ap_id, row


def add_payment(
    ds: Dataset,
    *,
    transaction_type: str,
    payee_supplier_id: str,
    amount_cents: int,
    payment_date: date,
    settlement_status: str,
    original_payment_id: str | None = None,
    bank_reference: str | None = None,
    auto_bank_reference: bool = True,
) -> tuple[str, int]:
    payment_id = ds.next_payment_id()
    if auto_bank_reference and settlement_status == "SETTLED":
        bank_reference = ds.next_bank_reference()
    row = ds.table("payments").add(
        {
            "payment_id": payment_id,
            "transaction_type": transaction_type,
            "original_payment_id": original_payment_id,
            "payee_supplier_id": payee_supplier_id,
            "payment_date": iso(payment_date),
            "amount": yuan(amount_cents),
            "currency": "CNY",
            "settlement_status": settlement_status,
            "bank_reference": bank_reference,
        }
    )
    return payment_id, row


def add_allocation(
    ds: Dataset,
    payment_id: str,
    ap_id: str | None,
    amount_cents: int,
    original_allocation_id: str | None = None,
) -> tuple[str, int]:
    allocation_id = ds.next_allocation_id()
    row = ds.table("allocations").add(
        {
            "allocation_id": allocation_id,
            "payment_id": payment_id,
            "ap_id": ap_id,
            "allocated_amount": yuan(amount_cents),
            "original_allocation_id": original_allocation_id,
        }
    )
    return allocation_id, row


# ---------------------------------------------------------------------------
# 主演示数据集（1000 条资金事件）
# ---------------------------------------------------------------------------


def build_normal_single_payments(ds: Dataset, rng: random.Random, count: int = 963) -> None:
    """1. 普通一次付款：一张应付 + 一笔足额付款 + 一条分摊。"""
    shared_ap_number = "CG-2026-09-0417"
    for index in range(count):
        if index == 0:
            cents = 1  # 0.01
        elif index == 1:
            cents = 123456  # 1234.56
        else:
            cents = rng.randrange(10000, 500001)  # 100.00 ~ 5000.00
        supplier_id = f"S{3 + (index % 8):03d}"
        ap_date = PERIOD_START + timedelta(days=rng.randrange(0, 20))
        payment_date = min(ap_date + timedelta(days=rng.randrange(1, 6)), PERIOD_END)
        ap_number = shared_ap_number if index in (0, 1) else None
        ap_id, _ = add_payable(ds, supplier_id, cents, ap_date, ap_number=ap_number)
        payment_id, _ = add_payment(
            ds,
            transaction_type="PAY",
            payee_supplier_id=supplier_id,
            amount_cents=cents,
            payment_date=payment_date,
            settlement_status="SETTLED",
        )
        add_allocation(ds, payment_id, ap_id, cents)
    ds.composition["normal_single_payment_groups"] = count
    ds.non_hit_scenarios.append(
        {
            "scenario": "普通一次付款（963 组）",
            "cases": "C01",
            "detail": "单笔足额付款、关联与主体一致，四规则范围内均不命中。",
        }
    )
    ds.non_hit_scenarios.append(
        {
            "scenario": "两张应付共用同一展示单号（普通组内 2 张）",
            "cases": "C01 变体说明",
            "detail": "ap_number 只是展示编号，两张单号相同、供应商不同、ap_id 不同，"
            "关联不依赖展示编号，四规则范围内均不命中。",
        }
    )


def build_installments(ds: Dataset, groups: int = 5) -> None:
    """2. 正常分期付款：应付 5000.00，分付 2000.00 + 3000.00。"""
    for _ in range(groups):
        ap_id, _ = add_payable(ds, "S011", 500000, date(2026, 9, 4))
        first, _ = add_payment(
            ds,
            transaction_type="PAY",
            payee_supplier_id="S011",
            amount_cents=200000,
            payment_date=date(2026, 9, 8),
            settlement_status="SETTLED",
        )
        second, _ = add_payment(
            ds,
            transaction_type="PAY",
            payee_supplier_id="S011",
            amount_cents=300000,
            payment_date=date(2026, 9, 15),
            settlement_status="SETTLED",
        )
        add_allocation(ds, first, ap_id, 200000)
        add_allocation(ds, second, ap_id, 300000)
    ds.composition["installment_groups"] = groups
    ds.non_hit_scenarios.append(
        {
            "scenario": "正常分期付款（5 组）",
            "cases": "C02",
            "detail": "两笔分期合计足额，任一笔均不等于整张应付金额，不构成现存重复足额或超额。",
        }
    )


def build_combined_payments(ds: Dataset, groups: int = 5) -> None:
    """3. 正常合并付款：一笔 10000.00 分至两张应付 6000.00 + 4000.00。"""
    for _ in range(groups):
        first_ap, _ = add_payable(ds, "S016", 600000, date(2026, 9, 3))
        second_ap, _ = add_payable(ds, "S016", 400000, date(2026, 9, 3))
        payment_id, _ = add_payment(
            ds,
            transaction_type="PAY",
            payee_supplier_id="S016",
            amount_cents=1000000,
            payment_date=date(2026, 9, 9),
            settlement_status="SETTLED",
        )
        add_allocation(ds, payment_id, first_ap, 600000)
        add_allocation(ds, payment_id, second_ap, 400000)
    ds.composition["combined_payment_groups"] = groups
    ds.non_hit_scenarios.append(
        {
            "scenario": "正常合并付款（5 组）",
            "cases": "C03",
            "detail": "一笔资金事件分摊到两张应付单，只有一笔资金事件，不能把整笔金额重复累计。",
        }
    )


def build_full_refund_of_duplicate(ds: Dataset) -> None:
    """4. 全额退回其中一次重复付款：2 笔 PAY + 1 笔 REFUND + 1 张应付 + 3 条分摊。"""
    ap_id, _ = add_payable(ds, "S017", 500000, date(2026, 9, 5))
    first, _ = add_payment(
        ds,
        transaction_type="PAY",
        payee_supplier_id="S017",
        amount_cents=500000,
        payment_date=date(2026, 9, 10),
        settlement_status="SETTLED",
    )
    second, _ = add_payment(
        ds,
        transaction_type="PAY",
        payee_supplier_id="S017",
        amount_cents=500000,
        payment_date=date(2026, 9, 11),
        settlement_status="SETTLED",
    )
    add_allocation(ds, first, ap_id, 500000)
    second_allocation, _ = add_allocation(ds, second, ap_id, 500000)
    refund, _ = add_payment(
        ds,
        transaction_type="REFUND",
        payee_supplier_id="S017",
        amount_cents=500000,
        payment_date=date(2026, 9, 20),
        settlement_status="SETTLED",
        original_payment_id=second,
    )
    add_allocation(ds, refund, ap_id, 500000, original_allocation_id=second_allocation)
    ds.composition["full_refund_of_duplicate_groups"] = 1
    ds.non_hit_scenarios.append(
        {
            "scenario": "重复付款后全额退款（1 组）",
            "cases": "C06",
            "detail": "被退那一笔净额为零，退款后不存在重复足额付款，也没有净超额；原付款与退款记录均保留。",
        }
    )


def build_partial_refund(ds: Dataset) -> None:
    """5. 正常部分退款：应付 5000.00，支付 5000.00，退回 2000.00，净付 3000.00。"""
    ap_id, _ = add_payable(ds, "S017", 500000, date(2026, 9, 5))
    payment_id, _ = add_payment(
        ds,
        transaction_type="PAY",
        payee_supplier_id="S017",
        amount_cents=500000,
        payment_date=date(2026, 9, 10),
        settlement_status="SETTLED",
    )
    allocation_id, _ = add_allocation(ds, payment_id, ap_id, 500000)
    refund_id, _ = add_payment(
        ds,
        transaction_type="REFUND",
        payee_supplier_id="S017",
        amount_cents=200000,
        payment_date=date(2026, 9, 22),
        settlement_status="SETTLED",
        original_payment_id=payment_id,
    )
    add_allocation(ds, refund_id, ap_id, 200000, original_allocation_id=allocation_id)
    ds.composition["partial_refund_groups"] = 1
    ds.non_hit_scenarios.append(
        {
            "scenario": "正常部分退款（1 组）",
            "cases": "C06b 同类",
            "detail": "净额 3000.00 低于应付 5000.00，不构成超额；退款有效关联原付款分摊。",
        }
    )


def build_non_settled(ds: Dataset) -> None:
    """6. 非结算付款：PENDING、FAILED、VOIDED 各一笔，各有独立应付及分摊。"""
    for status in ("PENDING", "FAILED", "VOIDED"):
        ap_id, _ = add_payable(ds, "S018", 500000, date(2026, 9, 6))
        payment_id, _ = add_payment(
            ds,
            transaction_type="PAY",
            payee_supplier_id="S018",
            amount_cents=500000,
            payment_date=date(2026, 9, 12),
            settlement_status=status,
            bank_reference=None,
        )
        add_allocation(ds, payment_id, ap_id, 500000)
    ds.composition["non_settled_groups"] = 3
    ds.non_hit_scenarios.append(
        {
            "scenario": "非结算付款（3 组：PENDING / FAILED / VOIDED）",
            "cases": "C10",
            "detail": "有记录不代表资金已经支付；非结算事件不计入资金计算，也不计入已结算流水唯一性。",
        }
    )


def build_unpaid_payables(ds: Dataset) -> None:
    """7. 尚未付款的应付：无资金事件和分摊。"""
    for _ in range(3):
        add_payable(ds, "S019", 500000, date(2026, 9, 7))
    ds.composition["unpaid_payable_count"] = 3
    ds.non_hit_scenarios.append(
        {
            "scenario": "尚未付款的应付（3 张）",
            "cases": "无对应案例，属正常未付款",
            "detail": "没有资金事件和分摊，四规则范围内均不命中；应付总额本来就不要求等于付款总额。",
        }
    )


def build_planted_clue_groups(ds: Dataset) -> None:
    """8. 预埋当前规则线索 9 组：3 组现存重复足额、2 组只超额、2 组缺失应付关联、2 组收款对象不同。"""
    # 3 组现存重复足额：各命中 R01 与 R02 一条
    for _ in range(3):
        ap_id, ap_row = add_payable(ds, "S001", 500000, date(2026, 9, 5))
        pay_a, pay_a_row = add_payment(
            ds,
            transaction_type="PAY",
            payee_supplier_id="S001",
            amount_cents=500000,
            payment_date=date(2026, 9, 10),
            settlement_status="SETTLED",
        )
        pay_b, pay_b_row = add_payment(
            ds,
            transaction_type="PAY",
            payee_supplier_id="S001",
            amount_cents=500000,
            payment_date=date(2026, 9, 12),
            settlement_status="SETTLED",
        )
        alloc_a, alloc_a_row = add_allocation(ds, pay_a, ap_id, 500000)
        alloc_b, alloc_b_row = add_allocation(ds, pay_b, ap_id, 500000)
        ds.clues.append(
            {
                "rule_id": "R01",
                "hit_key": {"rule_id": "R01", "ap_id": ap_id, "payee_supplier_id": "S001"},
                "expected_amounts_cents": {
                    "payable_amount": 500000,
                    "net_per_payment": {pay_a: 500000, pay_b: 500000},
                    "net_total": 1000000,
                },
                "related_ids": {
                    "payable_id": ap_id,
                    "payment_ids": [pay_a, pay_b],
                    "allocation_ids": [alloc_a, alloc_b],
                    "refund_ids": [],
                },
                "source_rows": {
                    "payables": {ap_id: ap_row},
                    "payments": {pay_a: pay_a_row, pay_b: pay_b_row},
                    "allocations": {alloc_a: alloc_a_row, alloc_b: alloc_b_row},
                },
                "reason": "该应付单在截至日存在两笔不同银行流水的已结算原付款，"
                "每笔净额均精确等于应付金额，属现存重复足额付款线索。",
            }
        )
        ds.clues.append(
            {
                "rule_id": "R02",
                "hit_key": {"rule_id": "R02", "ap_id": ap_id},
                "expected_amounts_cents": {
                    "payable_amount": 500000,
                    "paid_total": 1000000,
                    "refund_total": 0,
                    "net_total": 1000000,
                    "overpayment": 500000,
                },
                "related_ids": {
                    "payable_id": ap_id,
                    "payment_ids": [pay_a, pay_b],
                    "allocation_ids": [alloc_a, alloc_b],
                    "refund_ids": [],
                },
                "source_rows": {
                    "payables": {ap_id: ap_row},
                    "payments": {pay_a: pay_a_row, pay_b: pay_b_row},
                    "allocations": {alloc_a: alloc_a_row, alloc_b: alloc_b_row},
                },
                "reason": "该应付单净付款超过应付金额 5000.00 元（按唯一应付单汇总一次）。",
            }
        )

    # 2 组只超额：各命中 R02 一条，超额 500.00
    for _ in range(2):
        ap_id, ap_row = add_payable(ds, "S001", 500000, date(2026, 9, 5))
        pay_a, pay_a_row = add_payment(
            ds,
            transaction_type="PAY",
            payee_supplier_id="S001",
            amount_cents=300000,
            payment_date=date(2026, 9, 10),
            settlement_status="SETTLED",
        )
        pay_b, pay_b_row = add_payment(
            ds,
            transaction_type="PAY",
            payee_supplier_id="S001",
            amount_cents=250000,
            payment_date=date(2026, 9, 12),
            settlement_status="SETTLED",
        )
        alloc_a, alloc_a_row = add_allocation(ds, pay_a, ap_id, 300000)
        alloc_b, alloc_b_row = add_allocation(ds, pay_b, ap_id, 250000)
        ds.clues.append(
            {
                "rule_id": "R02",
                "hit_key": {"rule_id": "R02", "ap_id": ap_id},
                "expected_amounts_cents": {
                    "payable_amount": 500000,
                    "paid_total": 550000,
                    "refund_total": 0,
                    "net_total": 550000,
                    "overpayment": 50000,
                },
                "related_ids": {
                    "payable_id": ap_id,
                    "payment_ids": [pay_a, pay_b],
                    "allocation_ids": [alloc_a, alloc_b],
                    "refund_ids": [],
                },
                "source_rows": {
                    "payables": {ap_id: ap_row},
                    "payments": {pay_a: pay_a_row, pay_b: pay_b_row},
                    "allocations": {alloc_a: alloc_a_row, alloc_b: alloc_b_row},
                },
                "reason": "合计超额但没有两笔各自足额的付款，只构成应付净付款超额线索。",
            }
        )

    # 2 组缺失应付关联：各命中 R03 一条，不创建对应应付
    for missing_kind, ap_value in (
        ("未填写关联", None),
        ("填写的标识不在台账", "AP90101"),
    ):
        payment_id, payment_row = add_payment(
            ds,
            transaction_type="PAY",
            payee_supplier_id="S001",
            amount_cents=500000,
            payment_date=date(2026, 9, 10),
            settlement_status="SETTLED",
        )
        allocation_id, allocation_row = add_allocation(ds, payment_id, ap_value, 500000)
        ds.clues.append(
            {
                "rule_id": "R03",
                "hit_key": {"rule_id": "R03", "allocation_id": allocation_id},
                "expected_amounts_cents": {
                    "allocated_amount": 500000,
                    "paid_amount": 500000,
                    "refunded_amount": 0,
                    "net_amount": 500000,
                },
                "related_ids": {
                    "payable_id": ap_value,
                    "payment_ids": [payment_id],
                    "allocation_ids": [allocation_id],
                    "refund_ids": [],
                    "missing_kind": missing_kind,
                },
                "source_rows": {
                    "payables": {},
                    "payments": {payment_id: payment_row},
                    "allocations": {allocation_id: allocation_row},
                },
                "reason": f"已结算付款分摊在所提供的应付台账中未找到关联对象（{missing_kind}）；"
                "该金额保留在总体勾稽中，不参加 R01、R02。",
            }
        )

    # 2 组收款对象不同：各命中 R04 一条，不预设实际违规
    for _ in range(2):
        ap_id, ap_row = add_payable(ds, "S001", 500000, date(2026, 9, 5))
        payment_id, payment_row = add_payment(
            ds,
            transaction_type="PAY",
            payee_supplier_id="S002",
            amount_cents=500000,
            payment_date=date(2026, 9, 10),
            settlement_status="SETTLED",
        )
        allocation_id, allocation_row = add_allocation(ds, payment_id, ap_id, 500000)
        ds.clues.append(
            {
                "rule_id": "R04",
                "hit_key": {"rule_id": "R04", "allocation_id": allocation_id},
                "expected_amounts_cents": {
                    "allocated_amount": 500000,
                    "payable_amount": 500000,
                    "net_amount": 500000,
                },
                "related_ids": {
                    "payable_id": ap_id,
                    "payment_ids": [payment_id],
                    "allocation_ids": [allocation_id],
                    "refund_ids": [],
                    "payable_supplier_id": "S001",
                    "payee_supplier_id": "S002",
                },
                "source_rows": {
                    "payables": {ap_id: ap_row},
                    "payments": {payment_id: payment_row},
                    "allocations": {allocation_id: allocation_row},
                },
                "reason": "实际收款对象与应付单供应商不同，需核实是否为授权代收、主体映射问题或付款错误；"
                "不得仅凭对象差异判定违规。",
            }
        )

    ds.composition["planted_clue_groups"] = 9
    ds.composition["planted_settled_pay"] = 14
    ds.composition["planted_payables"] = 7
    ds.composition["planted_allocations"] = 14


def build_main_dataset(seed: int) -> Dataset:
    ds = Dataset()
    rng = random.Random(seed)

    for index, name in enumerate(SUPPLIER_NAMES, start=1):
        add_supplier(ds, f"S{index:03d}", name)

    build_normal_single_payments(ds, rng, count=963)
    build_installments(ds, groups=5)
    build_combined_payments(ds, groups=5)
    build_full_refund_of_duplicate(ds)
    build_partial_refund(ds)
    build_non_settled(ds)
    build_unpaid_payables(ds)
    build_planted_clue_groups(ds)

    expected_rows = {"suppliers": 20, "payables": 993, "payments": 1000, "allocations": 1005}
    actual_rows = {name: len(ds.table(name).rows) for name in SHEET_ORDER}
    if actual_rows != expected_rows:
        raise RuntimeError(f"行数与规格不符：{actual_rows} != {expected_rows}")

    counts: dict[str, int] = {}
    for clue in ds.clues:
        counts[clue["rule_id"]] = counts.get(clue["rule_id"], 0) + 1
    expected_counts = {"R01": 3, "R02": 5, "R03": 2, "R04": 2}
    if counts != expected_counts:
        raise RuntimeError(f"答案登记数量与规格不符：{counts} != {expected_counts}")

    return ds


# ---------------------------------------------------------------------------
# 小型验收工作簿（22 份，按规格第 6 节手工案例忠实展开）
# ---------------------------------------------------------------------------


def AP(ap_id: str, supplier_id: str, payable_amount: str, ap_number: str | None = None) -> dict:
    return {
        "ap_id": ap_id,
        "ap_number": ap_number if ap_number is not None else ap_id,
        "supplier_id": supplier_id,
        "ap_date": FIXTURE_AP_DATE,
        "payable_amount": payable_amount,
        "currency": "CNY",
    }


def PAY(
    payment_id: str,
    payee_supplier_id: str,
    amount: str,
    bank_reference: str | None,
    status: str = "SETTLED",
) -> dict:
    return {
        "payment_id": payment_id,
        "transaction_type": "PAY",
        "original_payment_id": None,
        "payee_supplier_id": payee_supplier_id,
        "payment_date": FIXTURE_PAY_DATE,
        "amount": amount,
        "currency": "CNY",
        "settlement_status": status,
        "bank_reference": bank_reference,
    }


def REF(
    refund_id: str,
    original_payment_id: str,
    payee_supplier_id: str,
    amount: str,
    bank_reference: str | None,
    status: str = "SETTLED",
) -> dict:
    return {
        "payment_id": refund_id,
        "transaction_type": "REFUND",
        "original_payment_id": original_payment_id,
        "payee_supplier_id": payee_supplier_id,
        "payment_date": FIXTURE_REFUND_DATE,
        "amount": amount,
        "currency": "CNY",
        "settlement_status": status,
        "bank_reference": bank_reference,
    }


def A(
    allocation_id: str,
    payment_id: str,
    ap_id: str | None,
    allocated_amount: str,
    original_allocation_id: str | None = None,
) -> dict:
    return {
        "allocation_id": allocation_id,
        "payment_id": payment_id,
        "ap_id": ap_id,
        "allocated_amount": allocated_amount,
        "original_allocation_id": original_allocation_id,
    }


def fixture_suppliers() -> list[dict]:
    return [
        {"supplier_id": "S001", "supplier_name": "甲供应商"},
        {"supplier_id": "S002", "supplier_name": "乙收款对象"},
    ]


def c01_records() -> dict:
    return {
        "payables": [AP("AP01", "S001", "5000.00")],
        "payments": [PAY("P01", "S001", "5000.00", "B001")],
        "allocations": [A("A01", "P01", "AP01", "5000.00")],
    }


def build_fixture_records(case_id: str) -> tuple[dict, dict]:
    """返回 (四表记录, 各表列定义)。只有 C14a 按定义故意缺列。"""
    columns = {name: list(SHEET_COLUMNS[name]) for name in SHEET_ORDER}

    if case_id == "C01":
        records = c01_records()
    elif case_id == "C02":
        records = {
            "payables": [AP("AP01", "S001", "5000.00")],
            "payments": [
                PAY("P01", "S001", "2000.00", "B001"),
                PAY("P02", "S001", "3000.00", "B002"),
            ],
            "allocations": [A("A01", "P01", "AP01", "2000.00"), A("A02", "P02", "AP01", "3000.00")],
        }
    elif case_id == "C03":
        records = {
            "payables": [AP("AP01", "S001", "6000.00"), AP("AP02", "S001", "4000.00")],
            "payments": [PAY("P01", "S001", "10000.00", "B001")],
            "allocations": [A("A01", "P01", "AP01", "6000.00"), A("A02", "P01", "AP02", "4000.00")],
        }
    elif case_id == "C04":
        records = {
            "payables": [AP("AP01", "S001", "5000.00")],
            "payments": [
                PAY("P01", "S001", "5000.00", "B001"),
                PAY("P02", "S001", "5000.00", "B002"),
            ],
            "allocations": [A("A01", "P01", "AP01", "5000.00"), A("A02", "P02", "AP01", "5000.00")],
        }
    elif case_id == "C05":
        records = {
            "payables": [AP("AP01", "S001", "5000.00")],
            "payments": [
                PAY("P01", "S001", "3000.00", "B001"),
                PAY("P02", "S001", "2500.00", "B002"),
            ],
            "allocations": [A("A01", "P01", "AP01", "3000.00"), A("A02", "P02", "AP01", "2500.00")],
        }
    elif case_id in ("C06", "C06b"):
        refund_amount = "5000.00" if case_id == "C06" else "1000.00"
        records = {
            "payables": [AP("AP01", "S001", "5000.00")],
            "payments": [
                PAY("P01", "S001", "5000.00", "B001"),
                PAY("P02", "S001", "5000.00", "B002"),
                REF("R01", "P02", "S001", refund_amount, "BR001"),
            ],
            "allocations": [
                A("A01", "P01", "AP01", "5000.00"),
                A("A02", "P02", "AP01", "5000.00"),
                A("AR01", "R01", "AP01", refund_amount, "A02"),
            ],
        }
    elif case_id == "C07":
        records = {
            "payables": [],
            "payments": [PAY("P01", "S001", "5000.00", "B001")],
            "allocations": [A("A01", "P01", None, "5000.00")],
        }
    elif case_id == "C07b":
        records = {
            "payables": [AP("AP01", "S001", "3000.00")],
            "payments": [PAY("P01", "S001", "5000.00", "B001")],
            "allocations": [
                A("A01", "P01", "AP01", "3000.00"),
                A("A02", "P01", None, "2000.00"),
            ],
        }
    elif case_id == "C08":
        records = {
            "payables": [],
            "payments": [PAY("P01", "S001", "5000.00", "B001")],
            "allocations": [A("A01", "P01", "AP_MISSING", "5000.00")],
        }
    elif case_id == "C09":
        records = {
            "payables": [AP("AP01", "S001", "5000.00")],
            "payments": [PAY("P01", "S002", "5000.00", "B001")],
            "allocations": [A("A01", "P01", "AP01", "5000.00")],
        }
    elif case_id == "C09b":
        records = {
            "payables": [AP("AP01", "S001", "5000.00")],
            "payments": [
                PAY("P01", "S002", "5000.00", "B001"),
                REF("R01", "P01", "S002", "5000.00", "BR001"),
            ],
            "allocations": [
                A("A01", "P01", "AP01", "5000.00"),
                A("AR01", "R01", "AP01", "5000.00", "A01"),
            ],
        }
    elif case_id == "C10":
        records = {
            "payables": [AP("AP01", "S001", "5000.00")],
            "payments": [
                PAY("P01", "S001", "5000.00", None, "PENDING"),
                PAY("P02", "S001", "5000.00", None, "FAILED"),
                PAY("P03", "S001", "5000.00", None, "VOIDED"),
            ],
            "allocations": [
                A("A01", "P01", "AP01", "5000.00"),
                A("A02", "P02", "AP01", "5000.00"),
                A("A03", "P03", "AP01", "5000.00"),
            ],
        }
    elif case_id == "C10b":
        records = {
            "payables": [AP("AP01", "S001", "5000.00")],
            "payments": [
                PAY("P01", "S001", "5000.00", "B001"),
                PAY("P02", "S001", "5000.00", "B002"),
                REF("R01", "P02", "S001", "5000.00", None, "PENDING"),
            ],
            "allocations": [
                A("A01", "P01", "AP01", "5000.00"),
                A("A02", "P02", "AP01", "5000.00"),
                A("AR01", "R01", "AP01", "5000.00", "A02"),
            ],
        }
    elif case_id == "C11":
        records = {
            "payables": [AP("AP01", "S001", "5000.00")],
            "payments": [
                PAY("P01", "S001", "5000.00", "B001"),
                PAY("P02", "S001", "5000.00", "B001"),
            ],
            "allocations": [A("A01", "P01", "AP01", "5000.00"), A("A02", "P02", "AP01", "5000.00")],
        }
    elif case_id == "C12":
        records = {
            "payables": [AP("AP01", "S001", "6000.00"), AP("AP02", "S001", "4000.00")],
            "payments": [PAY("P01", "S001", "10000.00", "B001")],
            "allocations": [A("A01", "P01", "AP01", "6000.00"), A("A02", "P01", "AP02", "3000.00")],
        }
    elif case_id == "C13":
        records = {
            "payables": [AP("AP01", "S001", "5000.00")],
            "payments": [
                PAY("P01", "S001", "5000.00", "B001"),
                REF("R01", "P01", "S001", "6000.00", "BR001"),
            ],
            "allocations": [
                A("A01", "P01", "AP01", "5000.00"),
                A("AR01", "R01", "AP01", "6000.00", "A01"),
            ],
        }
    elif case_id == "C14a":
        records = c01_records()
        records["payables"] = [
            {key: value for key, value in row.items() if key != "payable_amount"}
            for row in records["payables"]
        ]
        columns["payables"] = [c for c in SHEET_COLUMNS["payables"] if c != "payable_amount"]
    elif case_id == "C14b":
        records = c01_records()
        records["payments"].append(dict(records["payments"][0]))
    elif case_id == "C14c":
        records = c01_records()
        records["payments"][0]["amount"] = "N/A"
    elif case_id == "C14d":
        records = c01_records()
        records["payments"][0]["payment_date"] = "2026-09-31"
    elif case_id == "C14e":
        records = c01_records()
        records["payments"][0]["amount"] = "5000.001"
        records["allocations"][0]["allocated_amount"] = "5000.001"
    else:
        raise KeyError(f"未定义的 fixture 案例：{case_id}")

    for name in SHEET_ORDER:
        records.setdefault(name, [])
    records["suppliers"] = fixture_suppliers()
    return records, columns


FIXTURE_CASES = [
    "C01",
    "C02",
    "C03",
    "C04",
    "C05",
    "C06",
    "C07",
    "C08",
    "C09",
    "C10",
    "C11",
    "C12",
    "C13",
    "C06b",
    "C07b",
    "C09b",
    "C10b",
    "C14a",
    "C14b",
    "C14c",
    "C14d",
    "C14e",
]

FIXTURE_EXPECTATIONS: dict[str, dict] = {
    "C01": {
        "expected_quality_status": "pass",
        "expected_rule_findings": [],
        "reason": "一次足额付款：关联和主体一致，无超额，四规则范围内不命中。",
    },
    "C02": {
        "expected_quality_status": "pass",
        "expected_rule_findings": [],
        "reason": "分期付款：两笔合计足额，任一笔均不等于整张应付金额，不判重复或超额。",
    },
    "C03": {
        "expected_quality_status": "pass",
        "expected_rule_findings": [],
        "reason": "合并付款：只有一笔资金事件，两个关联行不能让整笔金额重复累计。",
    },
    "C04": {
        "expected_quality_status": "pass",
        "expected_rule_findings": [
            {
                "rule_id": "R01",
                "hit_key": {"rule_id": "R01", "ap_id": "AP01", "payee_supplier_id": "S001"},
                "expected_amounts_cents": {
                    "payable_amount": 500000,
                    "net_per_payment": {"P01": 500000, "P02": 500000},
                },
                "reason": "两笔不同流水的已结算付款，各自净额均等于应付金额。",
            },
            {
                "rule_id": "R02",
                "hit_key": {"rule_id": "R02", "ap_id": "AP01"},
                "expected_amounts_cents": {
                    "payable_amount": 500000,
                    "paid_total": 1000000,
                    "refund_total": 0,
                    "net_total": 1000000,
                    "overpayment": 500000,
                },
                "reason": "净付款 10000.00 超过应付 5000.00，超额 5000.00。",
            },
        ],
        "reason": "两笔现存重复足额付款，同时构成应付净付款超额；两条命中不能相加成损失。",
    },
    "C05": {
        "expected_quality_status": "pass",
        "expected_rule_findings": [
            {
                "rule_id": "R02",
                "hit_key": {"rule_id": "R02", "ap_id": "AP01"},
                "expected_amounts_cents": {
                    "payable_amount": 500000,
                    "paid_total": 550000,
                    "refund_total": 0,
                    "net_total": 550000,
                    "overpayment": 50000,
                },
                "reason": "合计超额 500.00，但没有两笔各自足额的付款。",
            }
        ],
        "reason": "非重复足额的超付款，仅 R02 命中。",
    },
    "C06": {
        "expected_quality_status": "pass",
        "expected_rule_findings": [],
        "reason": "重复付款后全额退款：被退那一笔净额为零，不存在现存重复足额或净超额。",
    },
    "C07": {
        "expected_quality_status": "pass",
        "expected_rule_findings": [
            {
                "rule_id": "R03",
                "hit_key": {"rule_id": "R03", "allocation_id": "A01"},
                "expected_amounts_cents": {"allocated_amount": 500000, "net_amount": 500000},
                "reason": "已结算付款分摊未填写应付关联。",
            }
        ],
        "reason": "未填写应付关联：记录有效保留，只有关联信息缺失。",
    },
    "C08": {
        "expected_quality_status": "pass",
        "expected_rule_findings": [
            {
                "rule_id": "R03",
                "hit_key": {"rule_id": "R03", "allocation_id": "A01"},
                "expected_amounts_cents": {"allocated_amount": 500000, "net_amount": 500000},
                "reason": "填写的应付标识不在所提供台账中。",
            }
        ],
        "reason": "应付标识不在台账：缺失关联不是可删除的无效资金事件。",
    },
    "C09": {
        "expected_quality_status": "pass",
        "expected_rule_findings": [
            {
                "rule_id": "R04",
                "hit_key": {"rule_id": "R04", "allocation_id": "A01"},
                "expected_amounts_cents": {"allocated_amount": 500000, "net_amount": 500000},
                "reason": "实际收款对象 S002 与应付单供应商 S001 不同。",
            }
        ],
        "reason": "收款对象不同：需人工核实代收授权、主体映射与回单，不自动判错。",
    },
    "C10": {
        "expected_quality_status": "pass",
        "expected_rule_findings": [],
        "reason": "非结算状态：有记录不代表资金已支付，已结算付款和净资金均为零。",
    },
    "C11": {
        "expected_quality_status": "blocked",
        "expected_primary_codes": ["SETTLED_BANK_REFERENCE_DUPLICATE"],
        "reason": "同一银行流水重复导入：尚不能认定有两笔真实银行付款，Gate 2 不放行。",
    },
    "C12": {
        "expected_quality_status": "blocked",
        "expected_primary_codes": ["ALLOCATION_SUM_MISMATCH"],
        "reason": "主表 10000.00 与分摊 9000.00 相差 1000.00，不能只分析残余部分。",
    },
    "C13": {
        "expected_quality_status": "blocked",
        "expected_primary_codes": ["REFUND_EXCEEDS_ORIGINAL"],
        "reason": "退款分摊超过原付款分摊 1000.00，不能算成正常负净付款。",
    },
    "C06b": {
        "expected_quality_status": "pass",
        "expected_rule_findings": [
            {
                "rule_id": "R02",
                "hit_key": {"rule_id": "R02", "ap_id": "AP01"},
                "expected_amounts_cents": {
                    "payable_amount": 500000,
                    "paid_total": 1000000,
                    "refund_total": 100000,
                    "net_total": 900000,
                    "overpayment": 400000,
                },
                "reason": "部分退款后净额 9000.00 仍超额 4000.00；P02 净额 4000.00 不满足重复足额。",
            }
        ],
        "reason": "部分退款后仅 R02 命中。",
    },
    "C07b": {
        "expected_quality_status": "pass",
        "expected_rule_findings": [
            {
                "rule_id": "R03",
                "hit_key": {"rule_id": "R03", "allocation_id": "A02"},
                "expected_amounts_cents": {"allocated_amount": 200000, "net_amount": 200000},
                "reason": "付款 5000.00 中未匹配应付的 2000.00 部分未填写关联。",
            }
        ],
        "reason": "部分未匹配：R03 涉及分摊额为 2000.00，不是整笔 5000.00。",
    },
    "C09b": {
        "expected_quality_status": "pass",
        "expected_rule_findings": [
            {
                "rule_id": "R04",
                "hit_key": {"rule_id": "R04", "allocation_id": "A01"},
                "expected_amounts_cents": {
                    "allocated_amount": 500000,
                    "refunded_amount": 500000,
                    "net_amount": 0,
                },
                "reason": "历史收款对象差异不因退款消失，仍列示 S001 与 S002。",
            }
        ],
        "reason": "退款后仍保留 R04 历史事实。",
    },
    "C10b": {
        "expected_quality_status": "pass",
        "expected_rule_findings": [
            {
                "rule_id": "R01",
                "hit_key": {"rule_id": "R01", "ap_id": "AP01", "payee_supplier_id": "S001"},
                "expected_amounts_cents": {"payable_amount": 500000, "net_total": 1000000},
                "reason": "待结算退款不能抵减，两笔已结算付款各自净额仍为 5000.00。",
            },
            {
                "rule_id": "R02",
                "hit_key": {"rule_id": "R02", "ap_id": "AP01"},
                "expected_amounts_cents": {
                    "payable_amount": 500000,
                    "paid_total": 1000000,
                    "refund_total": 0,
                    "net_total": 1000000,
                    "overpayment": 500000,
                },
                "reason": "未结算退款不计入资金计算，净超额仍为 5000.00。",
            },
        ],
        "reason": "非结算退款不抵减，R01、R02 各一条。",
    },
    "C14a": {
        "expected_quality_status": "blocked",
        "expected_primary_codes": ["REQUIRED_COLUMN_MISSING"],
        "reason": "payables 缺 payable_amount 列，接收结构检查失败，Gate 1 不放行。",
    },
    "C14b": {
        "expected_quality_status": "blocked",
        "expected_primary_codes": ["PRIMARY_KEY_DUPLICATE"],
        "reason": "payments 中 P01 主键重复，不能自动去重或当作两笔付款。",
    },
    "C14c": {
        "expected_quality_status": "blocked",
        "expected_primary_codes": ["AMOUNT_INVALID"],
        "reason": "关键金额为文本 N/A，无法解析，隔离并阻断，不能补零。",
    },
    "C14d": {
        "expected_quality_status": "blocked",
        "expected_primary_codes": ["DATE_INVALID"],
        "reason": "2026-09-31 是无效日期，不能滚动修正为十月日期。",
    },
    "C14e": {
        "expected_quality_status": "blocked",
        "expected_primary_codes": ["AMOUNT_SCALE_INVALID"],
        "reason": "金额超过两位小数，阻断，不能舍入到 5000.00 通过。",
    },
}


# ---------------------------------------------------------------------------
# 写盘
# ---------------------------------------------------------------------------


def write_workbook(path: Path, sheet_specs: dict, *, title: str, description: str) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in SHEET_ORDER:
        columns, rows = sheet_specs[name]
        worksheet = workbook.create_sheet(title=name)
        worksheet.append(list(columns))
        for row in rows:
            worksheet.append(list(row))
    workbook.properties.title = title
    workbook.properties.creator = "scripts/make_demo_data.py（TASK-002 合成数据生成器）"
    workbook.properties.description = description
    workbook.properties.keywords = "synthetic;audit-demo;TASK-002"
    workbook.properties.category = "合成数据"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def read_back_sheets(path: Path) -> dict:
    """重新读取已落盘的工作簿，返回四表表头与有序行内容。"""
    workbook = load_workbook(path, data_only=True, read_only=True)
    result: dict[str, dict] = {}
    for name in SHEET_ORDER:
        worksheet = workbook[name]
        raw_rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        header = raw_rows[0] if raw_rows else []
        result[name] = {"columns": list(header), "rows": raw_rows[1:]}
    workbook.close()
    return result


def content_fingerprint(read_back: dict) -> str:
    payload = {
        "sheets": {
            name: {"columns": read_back[name]["columns"], "rows": read_back[name]["rows"]}
            for name in SHEET_ORDER
        }
    }
    return sha256_bytes(json_text(payload).encode("utf-8"))


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="TASK-002 合成数据生成器（开发期）")
    parser.add_argument("--seed", type=int, default=20260930, help="随机种子，默认 20260930")
    parser.add_argument("--out-dir", required=True, help="输出目录，必须位于 data/generated 内且不存在")
    return parser.parse_args(argv)


def resolve_out_dir(raw: str) -> Path:
    target = Path(raw).expanduser()
    if not target.is_absolute():
        target = (Path.cwd() / target).resolve()
    else:
        target = target.resolve()
    try:
        target.relative_to(GENERATED_ROOT)
    except ValueError:
        fail(f"输出目录必须位于 {GENERATED_ROOT} 之内：{target}")
    if target == GENERATED_ROOT:
        fail("输出目录不能是 data/generated 本身")
    if target.exists():
        fail(f"输出目录已存在，不覆盖任何内容：{target}")
    return target


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    target = resolve_out_dir(args.out_dir)
    if not SPEC_PATH.exists():
        fail(f"未找到业务规格文件：{SPEC_PATH}")

    GENERATED_ROOT.mkdir(parents=True, exist_ok=True)

    dataset = build_main_dataset(args.seed)

    input_dir = target / "input"
    fixtures_dir = target / "fixtures"
    evaluation_dir = target / "evaluation"
    checks_dir = target / "checks"
    for directory in (input_dir, fixtures_dir, evaluation_dir, checks_dir):
        directory.mkdir(parents=True, exist_ok=True)

    main_path = input_dir / "erp_export.xlsx"
    write_workbook(
        main_path,
        {name: (dataset.table(name).columns, dataset.table(name).rows) for name in SHEET_ORDER},
        title=f"{COMPANY_NAME} 2026-09 采购付款合成演示数据",
        description=SYNTHETIC_NOTE,
    )

    for case_id in FIXTURE_CASES:
        records, columns = build_fixture_records(case_id)
        sheet_specs = {
            name: (
                columns[name],
                [[record.get(column) for column in columns[name]] for record in records[name]],
            )
            for name in SHEET_ORDER
        }
        write_workbook(
            fixtures_dir / f"{case_id}.xlsx",
            sheet_specs,
            title=f"TASK-002 小型验收工作簿 {case_id}（合成数据）",
            description=f"{SYNTHETIC_NOTE} 案例编号：{case_id}。",
        )

    read_back = read_back_sheets(main_path)
    row_counts = {name: len(read_back[name]["rows"]) for name in SHEET_ORDER}
    fingerprint = content_fingerprint(read_back)
    input_sha256 = sha256_file(main_path)

    counts: dict[str, int] = {}
    for clue in dataset.clues:
        counts[clue["rule_id"]] = counts.get(clue["rule_id"], 0) + 1
    overpayment_total = 0
    for clue in dataset.clues:
        if clue["rule_id"] == "R02":
            overpayment_total += clue["expected_amounts_cents"]["overpayment"]
    expected_hit_counts = {
        "R01": counts.get("R01", 0),
        "R02": counts.get("R02", 0),
        "R03": counts.get("R03", 0),
        "R04": counts.get("R04", 0),
        "total": sum(counts.values()),
    }

    settled_pay = sum(
        1
        for row in dataset.table("payments").rows
        if row[dataset.table("payments").columns.index("transaction_type")] == "PAY"
        and row[dataset.table("payments").columns.index("settlement_status")] == "SETTLED"
    )
    settled_refund = sum(
        1
        for row in dataset.table("payments").rows
        if row[dataset.table("payments").columns.index("transaction_type")] == "REFUND"
        and row[dataset.table("payments").columns.index("settlement_status")] == "SETTLED"
    )
    non_settled_pay = sum(
        1
        for row in dataset.table("payments").rows
        if row[dataset.table("payments").columns.index("transaction_type")] == "PAY"
        and row[dataset.table("payments").columns.index("settlement_status")] != "SETTLED"
    )

    spec_sha256 = sha256_file(SPEC_PATH)
    generator_sha256 = sha256_file(SELF_PATH)

    def dependency_versions() -> dict:
        from importlib.metadata import PackageNotFoundError, version

        result = {}
        for name in ("openpyxl", "et-xmlfile", "pandas", "numpy", "chromadb"):
            try:
                result[name] = version(name)
            except PackageNotFoundError:
                result[name] = "NOT_INSTALLED"
        return result

    manifest = {
        "task": "TASK-002",
        "artifact": "synthetic audit demo dataset",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "synthetic_declaration": SYNTHETIC_NOTE,
        "company": COMPANY_NAME,
        "period": {"start": iso(PERIOD_START), "end": iso(PERIOD_END)},
        "seed": args.seed,
        "spec": {"path": "docs/SPEC_V1.md", "sha256": spec_sha256},
        "generator": {"path": "scripts/make_demo_data.py", "sha256": generator_sha256},
        "python_version": sys.version.split()[0],
        "dependencies": dependency_versions(),
        "row_counts": row_counts,
        "payment_composition": {
            "settled_pay": settled_pay,
            "settled_refund": settled_refund,
            "non_settled_pay": non_settled_pay,
        },
        "group_composition": dataset.composition,
        "input_file": {"path": "input/erp_export.xlsx", "sha256": input_sha256},
        "business_content_fingerprint": fingerprint,
        "expected_hit_counts": expected_hit_counts,
        "expected_overpayment_total_cents": overpayment_total,
        "notes": [
            "本目录全部为合成数据，仅用于开发与验收演示，不代表真实企业或真实审计发现。",
            "生成器不含任何筛查规则实现；答案按构造登记，不调用检测逻辑推导。",
            "答案册只用于开发期验证，不得作为检测程序输入或人类判断界面的预填内容。",
            "主工作簿不含答案页、异常标签列、隐藏标签或透露答案的批注。",
        ],
    }

    ground_truth = {
        "version": "TASK-002-groundtruth-1.0",
        "synthetic_declaration": SYNTHETIC_NOTE,
        "company": COMPANY_NAME,
        "period": {"start": iso(PERIOD_START), "end": iso(PERIOD_END)},
        "seed": args.seed,
        "spec": {"path": "docs/SPEC_V1.md", "sha256": spec_sha256},
        "main_input": {"path": "input/erp_export.xlsx", "sha256": input_sha256},
        "row_counts": row_counts,
        "expected_hit_counts": expected_hit_counts,
        "expected_overpayment_total_cents": overpayment_total,
        "measure_note": "九组场景、十二条命中与超额金额是不同统计口径，不能互相替代或相加称为损失；"
        "命中只是待核实线索，不是舞弊、损失或控制缺陷结论。",
        "usage_limit": "本答案册只用于开发期验证生成器与校验脚本，不得作为检测程序输入，"
        "也不得作为人类初步判断界面的预填内容。",
        "hits": dataset.clues,
        "non_hit_scenarios": dataset.non_hit_scenarios,
    }

    fixture_expectations = {
        "version": "TASK-002-fixtures-1.0",
        "synthetic_declaration": SYNTHETIC_NOTE,
        "seed": args.seed,
        "summary": {
            "total": len(FIXTURE_CASES),
            "valid_expected": sum(
                1 for c in FIXTURE_CASES if FIXTURE_EXPECTATIONS[c]["expected_quality_status"] == "pass"
            ),
            "blocked_expected": sum(
                1
                for c in FIXTURE_CASES
                if FIXTURE_EXPECTATIONS[c]["expected_quality_status"] == "blocked"
            ),
        },
        "notes": [
            "阻断不等于零命中：存在质量阻断时规则预期标为“未执行”。",
            "正常 fixture 可以有业务规则预期；质量通过不代表不存在业务线索。",
        ],
        "fixtures": [
            {
                "case_id": case_id,
                "workbook": f"fixtures/{case_id}.xlsx",
                "expected_quality_status": FIXTURE_EXPECTATIONS[case_id]["expected_quality_status"],
                "expected_primary_codes": FIXTURE_EXPECTATIONS[case_id].get("expected_primary_codes", []),
                "rule_expectation_status": (
                    "未执行"
                    if FIXTURE_EXPECTATIONS[case_id]["expected_quality_status"] == "blocked"
                    else "已登记"
                ),
                "expected_rule_findings": FIXTURE_EXPECTATIONS[case_id].get(
                    "expected_rule_findings", []
                ),
                "reason": FIXTURE_EXPECTATIONS[case_id]["reason"],
            }
            for case_id in FIXTURE_CASES
        ],
    }

    (target / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (evaluation_dir / "ground_truth.json").write_text(
        json.dumps(ground_truth, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (evaluation_dir / "fixture_expectations.json").write_text(
        json.dumps(fixture_expectations, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (evaluation_dir / "ground_truth.md").write_text(
        render_ground_truth_markdown(ground_truth, fixture_expectations), encoding="utf-8"
    )

    print(f"输出目录: {target}")
    print("四表实际行数: " + ", ".join(f"{k}={v}" for k, v in row_counts.items()))
    print(f"资金事件构成: 已结算 PAY={settled_pay}, 已结算 REFUND={settled_refund}, 非结算 PAY={non_settled_pay}")
    print(f"预期命中计数: {expected_hit_counts}")
    print(f"R02 预期超额合计: {overpayment_total} 分（{overpayment_total / 100:.2f} 元）")
    print(f"主输入 SHA-256: {input_sha256}")
    print(f"业务内容指纹: {fingerprint}")
    print(f"fixtures: {len(FIXTURE_CASES)} 份（有效预期 "
          f"{fixture_expectations['summary']['valid_expected']}，阻断预期 "
          f"{fixture_expectations['summary']['blocked_expected']}）")
    return 0


def render_ground_truth_markdown(ground_truth: dict, fixture_expectations: dict) -> str:
    lines: list[str] = []
    lines.append("# 主演示答案册（开发期验证用）")
    lines.append("")
    lines.append(f"- 版本：{ground_truth['version']}")
    lines.append(f"- 合成声明：{ground_truth['synthetic_declaration']}")
    lines.append(f"- 模拟企业：{ground_truth['company']}；期间：{ground_truth['period']['start']} 至 {ground_truth['period']['end']}")
    lines.append(f"- 种子：{ground_truth['seed']}")
    lines.append(f"- 主输入：{ground_truth['main_input']['path']}")
    lines.append(f"- 主输入 SHA-256：`{ground_truth['main_input']['sha256']}`")
    lines.append(f"- 规格 SHA-256：`{ground_truth['spec']['sha256']}`")
    lines.append("")
    lines.append("## 预期命中计数（场景标签核对用，不是检测引擎结果）")
    lines.append("")
    lines.append("| 规则 | 预期条数 |")
    lines.append("| --- | --- |")
    for rule_id in ("R01", "R02", "R03", "R04"):
        lines.append(f"| {rule_id} | {ground_truth['expected_hit_counts'][rule_id]} |")
    lines.append(f"| 合计 | {ground_truth['expected_hit_counts']['total']} |")
    lines.append("")
    lines.append(f"R02 按唯一应付汇总的预期超额合计：{ground_truth['expected_overpayment_total_cents'] / 100:.2f} 元"
                 f"（{ground_truth['expected_overpayment_total_cents']} 分）。")
    lines.append("")
    lines.append(f"> {ground_truth['measure_note']}")
    lines.append("")
    lines.append("## 逐条预期")
    lines.append("")
    for index, clue in enumerate(ground_truth["hits"], start=1):
        lines.append(f"### {index}. {clue['rule_id']} — {json.dumps(clue['hit_key'], ensure_ascii=False)}")
        lines.append("")
        lines.append(f"- 理由：{clue['reason']}")
        lines.append(f"- 预期整数分金额：`{json.dumps(clue['expected_amounts_cents'], ensure_ascii=False)}`")
        lines.append(f"- 相关 ID：`{json.dumps(clue['related_ids'], ensure_ascii=False)}`")
        lines.append(f"- 实际工作表行号：`{json.dumps(clue['source_rows'], ensure_ascii=False)}`")
        lines.append("")
    lines.append("## 四规则范围内不应命中的场景")
    lines.append("")
    for scenario in ground_truth["non_hit_scenarios"]:
        lines.append(f"- **{scenario['scenario']}**（对照 {scenario['cases']}）：{scenario['detail']}")
    lines.append("")
    lines.append("## 使用限制")
    lines.append("")
    lines.append(f"{ground_truth['usage_limit']}")
    lines.append("")
    lines.append("## 小型验收工作簿预期概览")
    lines.append("")
    lines.append(
        f"共 {fixture_expectations['summary']['total']} 份：预期有效 "
        f"{fixture_expectations['summary']['valid_expected']} 份，预期阻断 "
        f"{fixture_expectations['summary']['blocked_expected']} 份（阻断时规则预期标为“未执行”）。"
    )
    lines.append("")
    lines.append("| 案例 | 预期质量状态 | 预期主要代码／规则 | 说明 |")
    lines.append("| --- | --- | --- | --- |")
    for fixture in fixture_expectations["fixtures"]:
        if fixture["expected_quality_status"] == "blocked":
            expectation = ", ".join(fixture["expected_primary_codes"])
        else:
            rule_ids = [item["rule_id"] for item in fixture["expected_rule_findings"]]
            expectation = ", ".join(rule_ids) if rule_ids else "四规则均不命中"
        lines.append(
            f"| {fixture['case_id']} | {fixture['expected_quality_status']} | {expectation} | {fixture['reason']} |"
        )
    lines.append("")
    return "\n".join(lines) + "\n"


if __name__ == "__main__":
    sys.exit(main())
