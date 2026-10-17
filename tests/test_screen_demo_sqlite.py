#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""TASK-006 只读 SQL 规则核心开发验证。

用法：

    & 'D:\\harness work\\audit_sa_project\\.venv\\Scripts\\python.exe' -B `
      'D:\\harness work\\audit_sa_project\\tests\\test_screen_demo_sqlite.py' `
      --out-dir 'D:\\harness work\\audit_sa_project\\data\\sqlite_tests\\task006_01'

每个输入：新副本 → 真实 checker CLI（独立 stdout／stderr）→ 冻结标准化 API → 冻结装载 API →
规则核心 → 保存 screening_report.json。**所有核心输出保存完成后**才读取只读答案册对照。
对照不把答案条目复制成核心返回值；核心不接收答案路径、不读 evaluation、不导入生成器。

退出码：0 全部符合预期；1 断言或保护失败；2 参数或环境问题。
"""

from __future__ import annotations

import argparse
import builtins
import copy
import hashlib
import importlib.util
import io
import json
import shutil
import sqlite3
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SQLITE_TESTS_ROOT = (PROJECT_ROOT / "data" / "sqlite_tests").resolve()
CHECKER = PROJECT_ROOT / "scripts" / "check_demo_data.py"
NORMALIZER = PROJECT_ROOT / "scripts" / "normalize_demo_data.py"
STORE = PROJECT_ROOT / "scripts" / "store_demo_sqlite.py"
SCREEN = PROJECT_ROOT / "scripts" / "screen_demo_sqlite.py"
RULES_SQL = [PROJECT_ROOT / "sql" / "rules" / f"r0{index}.sql" for index in (1, 2, 3, 4)]

DEMO_DIR = PROJECT_ROOT / "data" / "generated" / "demo_20260930_a"
MAIN_INPUT = DEMO_DIR / "input" / "erp_export.xlsx"
FIXTURE_DIR = DEMO_DIR / "fixtures"
EVALUATION_DIR = (DEMO_DIR / "evaluation").resolve()
GROUND_TRUTH = EVALUATION_DIR / "ground_truth.json"
FIXTURE_EXPECTATIONS = EVALUATION_DIR / "fixture_expectations.json"

SHEET_COLUMNS = {
    "suppliers": ["supplier_id", "supplier_name"],
    "payables": ["ap_id", "ap_number", "supplier_id", "ap_date", "payable_amount", "currency"],
    "payments": [
        "payment_id", "transaction_type", "original_payment_id", "payee_supplier_id",
        "payment_date", "amount", "currency", "settlement_status", "bank_reference",
    ],
    "allocations": ["allocation_id", "payment_id", "ap_id", "allocated_amount", "original_allocation_id"],
}

PROTECTED_FILES = [
    CHECKER,
    NORMALIZER,
    STORE,
    PROJECT_ROOT / "sql" / "demo_schema.sql",
    PROJECT_ROOT / "docs" / "SPEC_V1.md",
    PROJECT_ROOT / "docs" / "AUDIT_CONTRACT_DRAFT.md",
    PROJECT_ROOT / "docs" / "EXEC_REPORT_005_R2.md",
    PROJECT_ROOT / "docs" / "SQLITE_LOAD_GUIDE.md",
    PROJECT_ROOT / "tests" / "test_store_demo_sqlite.py",
    MAIN_INPUT,
    *sorted(FIXTURE_DIR.glob("*.xlsx")),
]

ALLOCATION_RECORDS = {
    "suppliers": [{"supplier_id": "S001", "supplier_name": "甲供应商"},
                  {"supplier_id": "S002", "supplier_name": "乙收款对象"}],
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fail(message: str) -> None:
    print(f"RUNTIME_ERROR: {message}", file=sys.stderr)
    sys.exit(2)


def load_module(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def hit_key_tuple(hit_key: dict) -> tuple:
    return tuple(sorted((key, str(value)) for key, value in hit_key.items()))


def dumps_key(hit_key: dict) -> str:
    return json.dumps(hit_key, ensure_ascii=False, sort_keys=True)


class Case:
    def __init__(self, name: str, case_dir: Path, input_path: Path | None = None):
        self.name = name
        self.case_dir = case_dir
        self.input_path = input_path or (case_dir / "input.xlsx")
        self.checker_exit: int | None = None
        self.report: dict | None = None
        self.normalize_outcome: str | None = None
        self.normalize_message = ""
        self.load_outcome: str | None = None
        self.screen_outcome: str | None = None
        self.screen_message = ""
        self.screening: dict | None = None
        self.db_path: Path | None = None

    def findings(self, rule_id: str | None = None) -> list[dict]:
        assert self.screening is not None
        return [item for item in self.screening["findings"] if rule_id is None or item["rule_id"] == rule_id]


class Suite:
    def __init__(self, container: Path, normalizer, store, screen):
        self.container = container
        self.normalizer = normalizer
        self.store = store
        self.screen = screen
        self.cases_dir = container / "cases"
        self.cases_dir.mkdir(parents=True, exist_ok=True)
        self.cli_calls = 0
        self.screen_calls = 0
        self.cases: list[Case] = []
        self.failures: list[dict] = []
        self.command_log: list[dict] = []
        self.outputs_saved = 0
        self.answer_book_first_read_at_outputs: int | None = None
        self.opened_paths_under_guard: list[str] = []

    def checker_call(self, name: str, workbook: Path, case_dir: Path) -> tuple[int, dict | None]:
        self.cli_calls += 1
        report_path = case_dir / "checker_report.json"
        stdout_log = case_dir / f"{self.cli_calls:03d}_{name}.checker.stdout.txt"
        stderr_log = case_dir / f"{self.cli_calls:03d}_{name}.checker.stderr.txt"
        command = [sys.executable, "-B", str(CHECKER), "--workbook", str(workbook), "--report", str(report_path)]
        with stdout_log.open("w", encoding="utf-8") as out, stderr_log.open("w", encoding="utf-8") as err:
            exit_code = subprocess.run(command, stdout=out, stderr=err).returncode
        payload = json.loads(report_path.read_text(encoding="utf-8")) if report_path.exists() else None
        self.command_log.append(
            {
                "call_index": self.cli_calls,
                "case": name,
                "command": command,
                "exit_code": exit_code,
                "report": str(report_path),
                "stdout_log": str(stdout_log),
                "stderr_log": str(stderr_log),
            }
        )
        return exit_code, payload

    def prepare(self, name: str, source: Path | None, build=None) -> Case:
        case_dir = self.cases_dir / name
        case_dir.mkdir(parents=True, exist_ok=True)
        case = Case(name, case_dir)
        if build is not None:
            build(case.input_path)
        else:
            shutil.copyfile(source, case.input_path)
        exit_code, report = self.checker_call(name, case.input_path, case_dir)
        case.checker_exit = exit_code
        case.report = report
        self.cases.append(case)
        return case

    def pipeline(self, case: Case, *, screen_again: bool = False) -> Case:
        """checker → 标准化 → 装载 → 规则；每一步的产物都保存到案例目录。"""
        if case.report is None or case.report.get("result") != "pass":
            case.normalize_outcome = "upstream_blocked"
            case.normalize_message = "checker 未通过，未进入规则阶段"
            return case
        try:
            normalized = self.normalizer.normalize_workbook(case.input_path, case.report)
            case.normalize_outcome = "normalized"
        except self.normalizer.NormalizationBlocked as exc:
            case.normalize_outcome = "upstream_blocked"
            case.normalize_message = str(exc)
            return case
        except self.normalizer.NormalizationError as exc:
            case.normalize_outcome = "normalization_error"
            case.normalize_message = str(exc)
            return case
        (case.case_dir / "normalized.json").write_text(
            json.dumps(normalized, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8"
        )
        try:
            load_report = self.store.load_normalized_to_sqlite(normalized, case.case_dir / "demo.sqlite")
            case.load_outcome = "loaded"
        except (self.store.SqliteLoadBlocked, self.store.SqliteLoadError) as exc:
            case.load_outcome = "load_failed"
            case.normalize_message = str(exc)
            return case
        (case.case_dir / "load_report.json").write_text(
            json.dumps(load_report, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8"
        )
        case.db_path = case.case_dir / "demo.sqlite"
        case.load_report_payload = load_report
        self.run_screen(case, save=True)
        return case

    def run_screen(self, case: Case, *, save: bool) -> Case:
        self.screen_calls += 1
        try:
            result = self.screen.screen_sqlite(case.db_path, case.load_report_payload)
            case.screen_outcome = "computed"
            case.screening = result
            if save:
                (case.case_dir / "screening_report.json").write_text(
                    json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8"
                )
                self.outputs_saved += 1
        except self.screen.RuleScreenBlocked as exc:
            case.screen_outcome = "interface_blocked"
            case.screen_message = str(exc)
        except self.screen.RuleScreenError as exc:
            case.screen_outcome = "runtime_error"
            case.screen_message = str(exc)
        return case

    def screen_raw(self, db_path, load_report, **kwargs):
        """直接调用核心（用于接口边界与确定性检查）。"""
        self.screen_calls += 1
        return self.screen.screen_sqlite(db_path, load_report, **kwargs)

    def judge(self, case: Case, *, outcome: str, kind: str, checks: dict) -> Case:
        problems: list[str] = []
        if case.screen_outcome != outcome:
            problems.append(f"规则结果期望 {outcome}，实际 {case.screen_outcome}（{case.screen_message}）")
        for label, predicate in checks.items():
            try:
                ok = bool(predicate(case))
            except Exception as exc:  # noqa: BLE001
                ok = False
                problems.append(f"{label} 断言时异常：{exc!r}")
            if not ok:
                problems.append(f"断言失败：{label}")
        record = {
            "case": case.name,
            "kind": kind,
            "checker_exit": case.checker_exit,
            "normalize_outcome": case.normalize_outcome,
            "screen_outcome": case.screen_outcome,
            "message": case.screen_message,
            "result": "PASS" if not problems else "FAIL",
            "problems": problems,
        }
        if problems:
            self.failures.append(record)
        print(f"[{'PASS' if not problems else 'FAIL'}] {case.name} screen={case.screen_outcome}"
              + ("" if not problems else "  | " + " ; ".join(problems)))
        return case


# ---------------------------------------------------------------------------
# 对照工具（只在核心输出保存完成后使用）
# ---------------------------------------------------------------------------


def compare_amounts(expected: dict, actual: dict, path: str, problems: list[str]) -> None:
    for key, value in expected.items():
        if key not in actual:
            problems.append(f"{path}.amounts_cents 缺少 {key}")
            continue
        if isinstance(value, dict):
            if not isinstance(actual[key], dict):
                problems.append(f"{path}.amounts_cents.{key} 应为对象")
                continue
            compare_amounts(value, actual[key], f"{path}.amounts_cents.{key}", problems)
        elif actual[key] != value:
            problems.append(f"{path}.amounts_cents.{key} 实际 {actual[key]!r} != 预期 {value!r}")


def compare_registered_ids(expected: dict, actual: dict, path: str, problems: list[str]) -> None:
    for key, value in expected.items():
        if key not in actual:
            problems.append(f"{path}.related_ids 缺少 {key}")
            continue
        actual_value = actual[key]
        if isinstance(value, list):
            if sorted(map(str, actual_value)) != sorted(map(str, value)):
                problems.append(f"{path}.related_ids.{key} 实际 {actual_value!r} != 预期 {value!r}")
        elif actual_value != value:
            problems.append(f"{path}.related_ids.{key} 实际 {actual_value!r} != 预期 {value!r}")


def compare_source_rows(expected: dict, actual: dict, path: str, problems: list[str]) -> None:
    for sheet, rows in expected.items():
        actual_rows = actual.get(sheet, {})
        for record_id, excel_row in rows.items():
            if actual_rows.get(record_id) != excel_row:
                problems.append(
                    f"{path}.source_rows.{sheet}[{record_id}] 实际 {actual_rows.get(record_id)!r} != 预期 {excel_row!r}"
                )


def compare_findings(expected_hits: list[dict], case: Case, label: str) -> list[str]:
    problems: list[str] = []
    actual_findings = case.findings()
    expected_by_key = {hit_key_tuple(item["hit_key"]): item for item in expected_hits}
    actual_by_key = {hit_key_tuple(item["hit_key"]): item for item in actual_findings}
    if len(actual_by_key) != len(actual_findings):
        problems.append("实际命中的结构化键存在重复")
    missing = sorted(set(expected_by_key) - set(actual_by_key))
    extra = sorted(set(actual_by_key) - set(expected_by_key))
    if missing:
        problems.append(f"{label} 漏报命中键：{missing}")
    if extra:
        problems.append(f"{label} 多报命中键：{extra}")
    for key in sorted(set(expected_by_key) & set(actual_by_key)):
        expected = expected_by_key[key]
        actual = actual_by_key[key]
        path = f"{label}:{dumps_key(dict(key))}"
        if "expected_amounts_cents" in expected:
            compare_amounts(expected["expected_amounts_cents"], actual["amounts_cents"], path, problems)
        if "related_ids" in expected:
            compare_registered_ids(expected["related_ids"], actual["related_ids"], path, problems)
        if "source_rows" in expected:
            compare_source_rows(expected["source_rows"], actual["source_rows"], path, problems)
        if actual.get("fact", "") in ("", None):
            problems.append(f"{path} 缺少固定事实描述")
    return problems


def build_workbook(path: Path, tables: dict) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in SHEET_COLUMNS:
        worksheet = workbook.create_sheet(title=sheet)
        worksheet.append(SHEET_COLUMNS[sheet])
        for record in tables.get(sheet, []):
            worksheet.append([record.get(column) for column in SHEET_COLUMNS[sheet]])
    workbook.save(path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="TASK-006 只读 SQL 规则核心开发验证")
    parser.add_argument("--out-dir", required=True, help="data/sqlite_tests 下尚不存在的测试容器")
    args = parser.parse_args(argv)

    container = Path(args.out_dir)
    container = container.resolve() if container.is_absolute() else (Path.cwd() / container).resolve()
    try:
        container.relative_to(SQLITE_TESTS_ROOT)
    except ValueError:
        fail(f"测试容器必须位于 {SQLITE_TESTS_ROOT} 之内：{container}")
    if container == SQLITE_TESTS_ROOT:
        fail("测试容器不能是 data/sqlite_tests 本身")
    if container.exists():
        fail(f"测试容器已存在，拒绝覆盖历史：{container}")
    for required in (CHECKER, NORMALIZER, STORE, SCREEN, MAIN_INPUT, *RULES_SQL):
        if not required.exists():
            fail(f"缺少必需输入：{required}")

    hashes_before = {str(path): sha256(path) for path in PROTECTED_FILES}
    container.mkdir(parents=True)
    normalizer = load_module(NORMALIZER, "normalizer_under_test")
    store = load_module(STORE, "store_under_test")
    screen = load_module(SCREEN, "screen_under_test")
    suite = Suite(container, normalizer, store, screen)
    print(f"测试容器: {container}")
    print(f"Python {sys.version.split()[0]} / SQLite {sqlite3.sqlite_version} / 规则核心 {screen.RULE_CORE_VERSION}")
    print()

    # ---------- 0. 答案隔离：静态与导入边界 ----------
    forbidden = ("ground_truth", "fixture_expectations", "evaluation", "make_demo_data")
    static_hits: list[str] = []
    for path in [SCREEN, *RULES_SQL]:
        text = path.read_text(encoding="utf-8")
        for token in forbidden:
            if token in text:
                static_hits.append(f"{path.name}:{token}")
    print(
        f"[{'PASS' if not static_hits else 'FAIL'}] ANSWER_ISOLATION_static 核心与规则 SQL 未引用答案册／生成器"
        + ("" if not static_hits else f"  违规：{static_hits}")
    )
    if static_hits:
        suite.failures.append({"case": "ANSWER_ISOLATION_static", "kind": "isolation", "result": "FAIL",
                               "problems": static_hits})

    # ---------- 1. 主文件与 22 份 fixture：先跑核心并保存输出 ----------
    all_inputs = [("MAIN", MAIN_INPUT)] + [
        (path.stem, path) for path in sorted(FIXTURE_DIR.glob("*.xlsx"))
    ]
    for name, path in all_inputs:
        case = suite.prepare(name, path)
        suite.pipeline(case)

    main_case = next(item for item in suite.cases if item.name == "MAIN")
    suite.judge(
        main_case, outcome="computed", kind="expected_success",
        checks={
            "状态与模式": lambda c: c.screening["batch_status"] == "DEV_RULES_COMPUTED"
            and c.screening["mode"] == "DEVELOPMENT_TEST_ONLY"
            and c.screening["next_step_allowed"] is False,
            "命中数 3/5/2/2": lambda c: c.screening["rule_hit_counts"]
            == {"R01": 3, "R02": 5, "R03": 2, "R04": 2, "total": 12},
            "R02 超额合计 1600000 分": lambda c: c.screening["statistics"]["r02_overpayment_total_cents"] == 1600000,
            "总体勾稽与装载报告一致": lambda c: c.screening["overall_reconciliation"]["matches_load_report"] is True
            and c.screening["overall_reconciliation"]["master_net"] == 267249329,
            "只读保护": lambda c: c.screening["read_only"] == {"uri_mode": "ro", "query_only": 1}
            and c.screening["database_unchanged"] is True,
            "命中键唯一": lambda c: len({dumps_key(f["hit_key"]) for f in c.findings()}) == len(c.findings()),
            "无 AI／等级字段": lambda c: not any(
                token in json.dumps(c.findings(), ensure_ascii=False)
                for token in ("severity", "severity_level", "ai_suggestion", "risk_level", "conclusion")
            ),
        },
    )

    # ---------- 2. 六个小补充场景（预期按卡固定，不依赖答案册） ----------
    def build_s1(path: Path) -> None:
        shutil.copyfile(FIXTURE_DIR / "C07.xlsx", path)
        workbook = load_workbook(path)
        workbook["payments"].append(["R01", "REFUND", "P01", "S001", "2026-09-20", "5000.00", "CNY", "SETTLED", "BR001"])
        workbook["allocations"].append(["AR01", "R01", None, "5000.00", "A01"])
        workbook.save(path)

    s1 = suite.prepare("S1_full_refund_keeps_r03", None, build=build_s1)
    suite.pipeline(s1)
    suite.judge(
        s1, outcome="computed", kind="expected_success",
        checks={
            "仍有一条 R03 且净额为 0": lambda c: [
                (item["amounts_cents"]["allocated_amount"], item["amounts_cents"]["refunded_amount"],
                 item["amounts_cents"]["net_amount"]) for item in c.findings("R03")
            ] == [(500000, 500000, 0)],
            "退款保持原空关联": lambda c: [
                item["evidence"]["related_refunds"][0]["refund_payment_id"] for item in c.findings("R03")
            ] == ["R01"],
            "其他规则无命中": lambda c: c.findings("R01") == [] and c.findings("R02") == [] and c.findings("R04") == [],
        },
    )

    def build_s2(path: Path) -> None:
        build_workbook(
            path,
            {
                **ALLOCATION_RECORDS,
                "payables": [{"ap_id": "AP01", "ap_number": "AP01", "supplier_id": "S001",
                              "ap_date": "2026-09-05", "payable_amount": "5000.00", "currency": "CNY"}],
                "payments": [
                    {"payment_id": "PA1", "transaction_type": "PAY", "original_payment_id": None,
                     "payee_supplier_id": "S001", "payment_date": "2026-09-10", "amount": "5000.00",
                     "currency": "CNY", "settlement_status": "SETTLED", "bank_reference": "B001"},
                    {"payment_id": "PA2", "transaction_type": "PAY", "original_payment_id": None,
                     "payee_supplier_id": "S001", "payment_date": "2026-09-11", "amount": "5000.00",
                     "currency": "CNY", "settlement_status": "SETTLED", "bank_reference": "B002"},
                    {"payment_id": "PA3", "transaction_type": "PAY", "original_payment_id": None,
                     "payee_supplier_id": "S001", "payment_date": "2026-09-12", "amount": "5000.00",
                     "currency": "CNY", "settlement_status": "SETTLED", "bank_reference": "B003"},
                ],
                "allocations": [
                    {"allocation_id": "A01", "payment_id": "PA1", "ap_id": "AP01", "allocated_amount": "5000.00"},
                    {"allocation_id": "A02", "payment_id": "PA2", "ap_id": "AP01", "allocated_amount": "5000.00"},
                    {"allocation_id": "A03", "payment_id": "PA3", "ap_id": "AP01", "allocated_amount": "5000.00"},
                ],
            },
        )

    s2 = suite.prepare("S2_three_full_payments", None, build=build_s2)
    suite.pipeline(s2)
    suite.judge(
        s2, outcome="computed", kind="expected_success",
        checks={
            "R01 一条组线索列出三笔付款": lambda c: len(c.findings("R01")) == 1
            and c.findings("R01")[0]["related_ids"]["payment_ids"] == ["PA1", "PA2", "PA3"]
            and c.findings("R01")[0]["amounts_cents"]["net_total"] == 1500000,
            "R02 一条超额 1000000 分": lambda c: [item["amounts_cents"]["overpayment"] for item in c.findings("R02")]
            == [1000000],
        },
    )

    def build_s3(path: Path) -> None:
        shutil.copyfile(FIXTURE_DIR / "C05.xlsx", path)
        workbook = load_workbook(path)
        workbook["payments"]["D3"] = "S002"
        workbook.save(path)

    s3 = suite.prepare("S3_payee_diff_with_overpay", None, build=build_s3)
    suite.pipeline(s3)
    suite.judge(
        s3, outcome="computed", kind="expected_success",
        checks={
            "R02 一条超额 50000 分": lambda c: [item["amounts_cents"]["overpayment"] for item in c.findings("R02")]
            == [50000],
            "第二笔原分摊另命中 R04 且分摊额 250000 分": lambda c: [
                (item["hit_key"]["allocation_id"], item["amounts_cents"]["allocated_amount"])
                for item in c.findings("R04")
            ] == [("A02", 250000)],
            "没有 R01": lambda c: c.findings("R01") == [],
        },
    )

    def build_s4(path: Path) -> None:
        shutil.copyfile(FIXTURE_DIR / "C10b.xlsx", path)
        workbook = load_workbook(path)
        for refund_id, status in (("R02", "FAILED"), ("R03", "VOIDED")):
            workbook["payments"].append(
                [refund_id, "REFUND", "P02", "S001", "2026-09-21", "5000.00", "CNY", status, None]
            )
            workbook["allocations"].append(
                [f"AR{refund_id[1:]}", refund_id, "AP01", "5000.00", "A02"]
            )
        workbook.save(path)

    s4 = suite.prepare("S4_non_settled_refunds", None, build=build_s4)
    suite.pipeline(s4)
    suite.judge(
        s4, outcome="computed", kind="expected_success",
        checks={
            "两笔足额原付款仍命中 R01／R02": lambda c: len(c.findings("R01")) == 1
            and [item["amounts_cents"]["overpayment"] for item in c.findings("R02")] == [500000],
            "三种非结算退款均不抵减": lambda c: all(
                item["counted_in_net"] is False for item in c.findings("R02")[0]["evidence"]["related_refunds"]
            ) and c.findings("R02")[0]["amounts_cents"]["refund_total"] == 0,
            "退款证据含状态": lambda c: sorted(
                item["settlement_status"] for item in c.findings("R02")[0]["evidence"]["related_refunds"]
            ) == ["FAILED", "PENDING", "VOIDED"],
        },
    )

    def build_s5(path: Path) -> None:
        build_workbook(
            path,
            {
                "suppliers": [
                    {"supplier_id": " C' ", "supplier_name": "供应商甲"},
                    {"supplier_id": "B| C' ", "supplier_name": "供应商乙"},
                ],
                "payables": [
                    {"ap_id": "A|B", "ap_number": "A|B", "supplier_id": " C' ", "ap_date": "2026-09-05",
                     "payable_amount": "5000.00", "currency": "CNY"},
                    {"ap_id": "A", "ap_number": "A", "supplier_id": "B| C' ", "ap_date": "2026-09-05",
                     "payable_amount": "5000.00", "currency": "CNY"},
                ],
                "payments": [
                    {"payment_id": "0001", "transaction_type": "PAY", "original_payment_id": None,
                     "payee_supplier_id": " C' ", "payment_date": "2026-09-10", "amount": "5000.00",
                     "currency": "CNY", "settlement_status": "SETTLED", "bank_reference": "B001"},
                    {"payment_id": "0002", "transaction_type": "PAY", "original_payment_id": None,
                     "payee_supplier_id": " C' ", "payment_date": "2026-09-11", "amount": "5000.00",
                     "currency": "CNY", "settlement_status": "SETTLED", "bank_reference": "B002"},
                    {"payment_id": "0003", "transaction_type": "PAY", "original_payment_id": None,
                     "payee_supplier_id": "B| C' ", "payment_date": "2026-09-12", "amount": "5000.00",
                     "currency": "CNY", "settlement_status": "SETTLED", "bank_reference": "B003"},
                    {"payment_id": "0004", "transaction_type": "PAY", "original_payment_id": None,
                     "payee_supplier_id": "B| C' ", "payment_date": "2026-09-13", "amount": "5000.00",
                     "currency": "CNY", "settlement_status": "SETTLED", "bank_reference": "B004"},
                ],
                "allocations": [
                    {"allocation_id": "AL01", "payment_id": "0001", "ap_id": "A|B", "allocated_amount": "5000.00"},
                    {"allocation_id": "AL02", "payment_id": "0002", "ap_id": "A|B", "allocated_amount": "5000.00"},
                    {"allocation_id": "AL03", "payment_id": "0003", "ap_id": "A", "allocated_amount": "5000.00"},
                    {"allocation_id": "AL04", "payment_id": "0004", "ap_id": "A", "allocated_amount": "5000.00"},
                ],
            },
        )

    s5 = suite.prepare("S5_key_collision_guard", None, build=build_s5)
    suite.pipeline(s5)
    suite.judge(
        s5, outcome="computed", kind="expected_success",
        checks={
            "R01 两条且键不碰撞": lambda c: sorted(
                (item["hit_key"]["ap_id"], item["hit_key"]["payee_supplier_id"]) for item in c.findings("R01")
            ) == [("A", "B| C' "), ("A|B", " C' ")],
            "R02 两条各超额 500000 分": lambda c: sorted(
                item["amounts_cents"]["overpayment"] for item in c.findings("R02")
            ) == [500000, 500000],
            "R03／R04 无命中": lambda c: c.findings("R03") == [] and c.findings("R04") == [],
            "空格与前导零原样保留": lambda c: all(
                pid.startswith("000") for item in c.findings("R01") for pid in item["related_ids"]["payment_ids"]
            ) and all(
                item["hit_key"]["payee_supplier_id"] in (" C' ", "B| C' ") for item in c.findings("R01")
            ),
        },
    )

    def build_s6(path: Path) -> None:
        shutil.copyfile(FIXTURE_DIR / "C01.xlsx", path)
        workbook = load_workbook(path)
        for sheet in ("payments", "allocations"):
            del workbook[sheet]
            worksheet = workbook.create_sheet(sheet)
            worksheet.append(SHEET_COLUMNS[sheet])
        workbook.save(path)

    s6 = suite.prepare("S6_no_funds", None, build=build_s6)
    suite.pipeline(s6)
    suite.judge(
        s6, outcome="computed", kind="expected_success",
        checks={
            "四规则空数组": lambda c: c.screening["rule_hit_counts"]["total"] == 0,
            "资金净额为 0 且应付保留": lambda c: c.screening["overall_reconciliation"]["master_net"] == 0
            and c.screening["overall_reconciliation"]["settled_pay_master_total"] == 0,
            "不因退款分组为空而阻断": lambda c: c.screen_outcome == "computed",
        },
    )

    # ---------- 3. 核心输出全部保存完成后，才开始读取只读答案册并对照 ----------
    suite.answer_book_first_read_at_outputs = suite.outputs_saved
    ground_truth = json.loads(GROUND_TRUTH.read_text(encoding="utf-8"))
    fixture_expectations = json.loads(FIXTURE_EXPECTATIONS.read_text(encoding="utf-8"))
    fixtures_by_id = {item["case_id"]: item for item in fixture_expectations["fixtures"]}

    main_problems = compare_findings(ground_truth["hits"], main_case, "MAIN")
    actual_overpayment = main_case.screening["statistics"]["r02_overpayment_total_cents"]
    if actual_overpayment != ground_truth["expected_overpayment_total_cents"]:
        main_problems.append(
            f"R02 超额合计 实际 {actual_overpayment} != 答案册 {ground_truth['expected_overpayment_total_cents']}"
        )
    suite.judge(
        main_case, outcome="computed", kind="expected_success",
        checks={"逐条对照答案册（键／金额／相关 ID／原行号）": lambda c: not main_problems},
    )

    # ---------- 4. fixture 对照 ----------
    for name in sorted(fixtures_by_id):
        case = next(item for item in suite.cases if item.name == name)
        expectation = fixtures_by_id[name]
        if expectation["expected_quality_status"] == "blocked":
            problems: list[str] = []
            if case.normalize_outcome != "upstream_blocked":
                problems.append(f"上游应阻断，实际 {case.normalize_outcome}")
            if case.screening is not None or case.screen_outcome is not None:
                problems.append("上游阻断时不得调用规则核心")
            if (case.case_dir / "demo.sqlite").exists():
                problems.append("上游阻断时不得建立数据库")
            if case.checker_exit != 1 or not (case.report and case.report.get("issue_counts_by_code")):
                problems.append("未保留真实 checker 问题")
            if problems:
                suite.failures.append({"case": name, "kind": "expected_upstream_block", "result": "FAIL",
                                       "problems": problems})
            print(f"[{'PASS' if not problems else 'FAIL'}] {name} upstream_blocked"
                  + ("" if not problems else "  | " + " ; ".join(problems)))
            continue
        problems = compare_findings(expectation["expected_rule_findings"], case, name)
        checks = {
            "逐条对照 fixture 期望": lambda c, p=problems: not p,
            "上游质量通过": lambda c: c.checker_exit == 0,
            "只读保护": lambda c: c.screening["database_unchanged"] is True,
        }
        if name == "C03":
            checks["C03 合并付款：无命中且整笔付款只计一次"] = (
                lambda c: c.screening["rule_hit_counts"]["total"] == 0
                and c.screening["overall_reconciliation"]["settled_pay_master_total"] == 1000000
            )
        if name == "C06":
            checks["C06 全额退款后无 R01／R02"] = lambda c: c.findings("R01") == [] and c.findings("R02") == []
        if name == "C06b":
            checks["C06b 部分退款：仅 R02 且超额 400000 分"] = lambda c: [
                item["amounts_cents"]["overpayment"] for item in c.findings("R02")
            ] == [400000]
        if name == "C07b":
            checks["C07b 分摊缺失：R03 涉及 200000 分而非整笔"] = lambda c: [
                item["amounts_cents"]["allocated_amount"] for item in c.findings("R03")
            ] == [200000]
        if name == "C09b":
            checks["C09b 全额退款后仍保留 R04 且净额为 0"] = lambda c: [
                (item["amounts_cents"]["refunded_amount"], item["amounts_cents"]["net_amount"])
                for item in c.findings("R04")
            ] == [(500000, 0)]
        if name == "C10":
            checks["C10 非结算：无命中且已结算为 0"] = (
                lambda c: c.screening["rule_hit_counts"]["total"] == 0
                and c.screening["overall_reconciliation"]["settled_pay_master_total"] == 0
            )
        if name == "C10b":
            checks["C10b 未结算退款不抵减"] = lambda c: [
                item["amounts_cents"]["refund_total"] for item in c.findings("R02")
            ] == [0]
        suite.judge(case, outcome="computed", kind="expected_success", checks=checks)

    # ---------- 5. 接口边界与确定性 ----------
    valid_case = next(item for item in suite.cases if item.name == "C04" and item.screen_outcome == "computed")

    missing_path = container / "missing_case" / "not_there.sqlite"
    missing_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        suite.screen_raw(missing_path, valid_case.load_report_payload)
        missing_outcome, missing_message = "computed", ""
    except screen.RuleScreenBlocked as exc:
        missing_outcome, missing_message = "blocked", str(exc)
    except screen.RuleScreenError as exc:
        missing_outcome, missing_message = "error", str(exc)
    print(f"[{'PASS' if missing_outcome == 'blocked' and not missing_path.exists() else 'FAIL'}] "
          f"I1_missing_database blocked={missing_outcome == 'blocked'} 未创建={not missing_path.exists()}")
    if missing_outcome != "blocked" or missing_path.exists():
        suite.failures.append({"case": "I1_missing_database", "kind": "interface", "result": "FAIL",
                               "problems": [f"outcome={missing_outcome} message={missing_message}"]})

    mismatch_dir = container / "mismatch_case"
    mismatch_dir.mkdir(parents=True, exist_ok=True)
    mismatch_db = mismatch_dir / "demo.sqlite"
    shutil.copyfile(valid_case.db_path, mismatch_db)
    with mismatch_db.open("ab") as handle:
        handle.write(b"-- tampered-by-test")
    mismatch_report = dict(valid_case.load_report_payload)
    mismatch_report["database_path"] = str(mismatch_db)
    try:
        suite.screen_raw(mismatch_db, mismatch_report)
        mismatch_outcome, mismatch_message = "computed", ""
    except screen.RuleScreenBlocked as exc:
        mismatch_outcome, mismatch_message = "blocked", str(exc)
    except screen.RuleScreenError as exc:
        mismatch_outcome, mismatch_message = "error", str(exc)
    mismatch_ok = mismatch_outcome == "blocked" and "database_sha256" in mismatch_message
    print(f"[{'PASS' if mismatch_ok else 'FAIL'}] I2_hash_mismatch blocked={mismatch_outcome == 'blocked'} "
          f"reason={mismatch_message[:60]}")
    if not mismatch_ok:
        suite.failures.append({"case": "I2_hash_mismatch", "kind": "interface", "result": "FAIL",
                               "problems": [f"outcome={mismatch_outcome} message={mismatch_message}"]})

    staging_dir = container / "staging_case"
    staging_dir.mkdir(parents=True, exist_ok=True)
    staging_db = staging_dir / "demo.sqlite.staging"
    shutil.copyfile(valid_case.db_path, staging_db)
    staging_report = dict(valid_case.load_report_payload)
    staging_report["database_path"] = str(staging_db)
    try:
        suite.screen_raw(staging_db, staging_report)
        staging_outcome, staging_message = "computed", ""
    except screen.RuleScreenBlocked as exc:
        staging_outcome, staging_message = "blocked", str(exc)
    except screen.RuleScreenError as exc:
        staging_outcome, staging_message = "error", str(exc)
    staging_ok = staging_outcome == "blocked" and "暂存" in staging_message
    print(f"[{'PASS' if staging_ok else 'FAIL'}] I3_staging_rejected blocked={staging_outcome == 'blocked'} "
          f"reason={staging_message[:60]}")
    if not staging_ok:
        suite.failures.append({"case": "I3_staging_rejected", "kind": "interface", "result": "FAIL",
                               "problems": [f"outcome={staging_outcome} message={staging_message}"]})

    before_hash = sha256(valid_case.db_path)
    repeat_first = suite.screen_raw(valid_case.db_path, valid_case.load_report_payload)
    repeat_second = suite.screen_raw(valid_case.db_path, valid_case.load_report_payload)
    deterministic = (
        json.dumps(repeat_first["findings"], ensure_ascii=False, sort_keys=True)
        == json.dumps(repeat_second["findings"], ensure_ascii=False, sort_keys=True)
        == json.dumps(valid_case.screening["findings"], ensure_ascii=False, sort_keys=True)
    ) and sha256(valid_case.db_path) == before_hash
    print(f"[{'PASS' if deterministic else 'FAIL'}] I4_determinism 复调用结果一致且数据库字节未变")
    if not deterministic:
        suite.failures.append({"case": "I4_determinism", "kind": "interface", "result": "FAIL",
                               "problems": ["复调用 findings 或数据库字节不一致"]})

    # ---------- 6. R1 新增：必需来源集合、完整 cells 与退款来源回归 ----------
    def required_entity_set(connection: sqlite3.Connection, finding: dict) -> set[tuple[str, str]]:
        """从数据库业务表独立推导本条线索的必需来源实体（不参考现有 evidence.entities）。"""
        required: set[tuple[str, str]] = set()
        related = finding["related_ids"]
        for payment_id in related["payment_ids"]:
            required.add(("payments", payment_id))
        for allocation_id in related["allocation_ids"]:
            required.add(("allocations", allocation_id))
            refund_rows = connection.execute(
                "SELECT allocation_id, payment_id FROM allocations WHERE original_allocation_id = ?",
                (allocation_id,),
            ).fetchall()
            for refund_allocation_id, refund_payment_id in refund_rows:
                required.add(("allocations", refund_allocation_id))
                required.add(("payments", refund_payment_id))
        payable_id = related.get("payable_id")
        if payable_id:
            exists = connection.execute(
                "SELECT 1 FROM payables WHERE ap_id = ?", (payable_id,)
            ).fetchone()
            if exists:
                required.add(("payables", payable_id))
        return required

    def verify_required_sources(case: Case) -> dict:
        connection = sqlite3.connect(f"file:{case.db_path}?mode=ro", uri=True)
        problems: list[str] = []
        refund_occurrences = 0
        required_total = 0
        try:
            for finding in case.findings():
                label = f"{case.name}:{dumps_key(finding['hit_key'])}"
                entities = finding["evidence"]["entities"]
                pairs = [(item["table_name"], item["record_id"]) for item in entities]
                if len(pairs) != len(set(pairs)):
                    problems.append(f"{label} 证据实体存在重复项")
                actual = set(pairs)
                required = required_entity_set(connection, finding)
                required_total += len(required)
                missing = sorted(required - actual)
                if missing:
                    problems.append(f"{label} 缺少必需来源实体：{missing}")
                for item in finding["evidence"]["related_refunds"]:
                    refund_occurrences += 2
                for table_name, record_id in sorted(required):
                    row = connection.execute(
                        "SELECT workbook_sha256, sheet, excel_row, cells_json FROM source_records "
                        "WHERE table_name = ? AND record_id = ?",
                        (table_name, record_id),
                    ).fetchone()
                    if row is None:
                        problems.append(f"{label} 数据库缺少 {table_name}:{record_id} 的来源记录")
                        continue
                    entity = next(
                        (item for item in entities
                         if item["table_name"] == table_name and item["record_id"] == record_id),
                        None,
                    )
                    if entity is None:
                        continue
                    for field, expected in (("workbook_sha256", row[0]), ("sheet", row[1]), ("excel_row", row[2])):
                        if entity.get(field) != expected:
                            problems.append(
                                f"{label} {table_name}:{record_id}.{field} 实际 {entity.get(field)!r} != 来源 {expected!r}"
                            )
                    if json.dumps(entity.get("cells"), ensure_ascii=False, sort_keys=True) != json.dumps(
                        json.loads(row[3]), ensure_ascii=False, sort_keys=True
                    ):
                        problems.append(f"{label} {table_name}:{record_id} 的完整 cells 与来源不相等")
                    if finding["source_rows"].get(table_name, {}).get(record_id) != row[2]:
                        problems.append(f"{label} {table_name}:{record_id} 的原 Excel 行号未进入 source_rows")
                for item in finding["evidence"]["related_refunds"]:
                    for kind, record_id in (
                        ("payments", item["refund_payment_id"]),
                        ("allocations", item["refund_allocation_id"]),
                    ):
                        if (kind, record_id) not in actual:
                            problems.append(f"{label} 退款摘要 {kind}:{record_id} 没有对应来源实体")
                    expected_counted = (
                        item["settlement_status"] == "SETTLED"
                        and "2026-09-01" <= item["payment_date"] <= "2026-09-30"
                    )
                    if item["counted_in_net"] != expected_counted:
                        problems.append(f"{label} 退款 {item['refund_payment_id']} 的 counted_in_net 与状态不符")
        finally:
            connection.close()
        return {
            "findings": len(case.findings()),
            "required_entities": required_total,
            "refund_entity_occurrences": refund_occurrences,
            "problems": problems,
        }

    source_reports: dict[str, dict] = {}
    for case in suite.cases:
        if case.screen_outcome == "computed":
            source_reports[case.name] = verify_required_sources(case)
    total_missing = sum(len(report["problems"]) for report in source_reports.values())
    print(f"[{'PASS' if total_missing == 0 else 'FAIL'}] R1_REQUIRED_SOURCES 逐条必需来源集合与完整 cells 核对"
          f"（{len(source_reports)} 个成功输入、{sum(r['findings'] for r in source_reports.values())} 条线索）")
    if total_missing:
        for name, report in source_reports.items():
            for problem in report["problems"][:5]:
                print(f"        - {name}: {problem}")
        suite.failures.append({"case": "R1_REQUIRED_SOURCES", "kind": "evidence", "result": "FAIL",
                               "problems": [p for r in source_reports.values() for p in r["problems"]][:20]})

    focus_inputs = ("C06b", "C09b", "C10b", "S1_full_refund_keeps_r03", "S4_non_settled_refunds")
    focus_findings = sum(source_reports.get(name, {}).get("findings", 0) for name in focus_inputs)
    focus_refund_occurrences = sum(
        source_reports.get(name, {}).get("refund_entity_occurrences", 0) for name in focus_inputs
    )
    focus_missing = sum(len(source_reports.get(name, {}).get("problems", [])) for name in focus_inputs)
    focus_ok = focus_findings == 7 and focus_refund_occurrences == 22 and focus_missing == 0
    print(f"[{'PASS' if focus_ok else 'FAIL'}] R1_REFUND_EVIDENCE 5 个退款相关输入："
          f"线索 {focus_findings}、退款实体出现项 {focus_refund_occurrences}、缺项 {focus_missing}")
    if not focus_ok:
        suite.failures.append({"case": "R1_REFUND_EVIDENCE", "kind": "evidence", "result": "FAIL",
                               "problems": [f"findings={focus_findings} occurrences={focus_refund_occurrences} "
                                            f"missing={focus_missing}"]})

    # 来源读取故障注入：只读连接代理（不改数据库、不加生产开关）
    class _TamperedCursor:
        def __init__(self, row):
            self._row = row

        def fetchone(self):
            return self._row

        def fetchall(self):
            return [] if self._row is None else [self._row]

    class _SourceTamperingConnection:
        def __init__(self, inner, table, record_id, mode, state):
            self._inner = inner
            self._table = table
            self._record_id = record_id
            self._mode = mode
            self._state = state

        def execute(self, sql, params=()):
            if "source_records" in sql and tuple(params) == (self._table, self._record_id):
                self._state["hits"] += 1
                if self._mode == "missing":
                    return _TamperedCursor(None)
                row = self._inner.execute(sql, params).fetchone()
                return _TamperedCursor((row[0], row[1], row[2], "{not-valid-json"))
            return self._inner.execute(sql, params)

        def __getattr__(self, item):
            return getattr(self._inner, item)

    def tampering_open(table: str, record_id: str, mode: str, state: dict):
        def factory(original):
            def patched(path):
                return _SourceTamperingConnection(original(path), table, record_id, mode, state)

            return patched

        return factory

    c09b = next(item for item in suite.cases if item.name == "C09b")
    c09b_finding = c09b.findings("R04")[0]
    original_payment = c09b_finding["related_ids"]["payment_ids"][0]
    refund_payment = c09b_finding["evidence"]["related_refunds"][0]["refund_payment_id"]
    refund_allocation = c09b_finding["evidence"]["related_refunds"][0]["refund_allocation_id"]

    fault_cases = (
        ("F1_original_payment_source_missing", "payments", original_payment, "missing", "缺少来源记录"),
        ("F2_refund_payment_source_missing", "payments", refund_payment, "missing", "缺少来源记录"),
        ("F3_refund_allocation_source_missing", "allocations", refund_allocation, "missing", "缺少来源记录"),
        ("F4_refund_payment_cells_invalid", "payments", refund_payment, "bad_json", "cells JSON 非法"),
        ("F5_refund_allocation_cells_invalid", "allocations", refund_allocation, "bad_json", "cells JSON 非法"),
    )
    fault_results: list[dict] = []
    before_fault_hash = sha256(c09b.db_path)
    for name, table, record_id, mode, expect_text in fault_cases:
        state = {"hits": 0}
        original_open = screen.open_readonly
        screen.open_readonly = tampering_open(table, record_id, mode, state)(original_open)
        suite.screen_calls += 1
        try:
            screen.screen_sqlite(c09b.db_path, c09b.load_report_payload)
            outcome, message = "computed", ""
        except screen.RuleScreenError as exc:
            outcome, message = "runtime_error", str(exc)
        except screen.RuleScreenBlocked as exc:
            outcome, message = "interface_blocked", str(exc)
        finally:
            screen.open_readonly = original_open
        ok = outcome == "runtime_error" and state["hits"] > 0 and expect_text in message and record_id in message
        fault_results.append(
            {"case": name, "table": table, "record_id": record_id, "mode": mode,
             "outcome": outcome, "injection_hits": state["hits"], "message": message[:200], "ok": ok}
        )
        print(f"[{'PASS' if ok else 'FAIL'}] {name} outcome={outcome} 注入命中={state['hits']} "
              f"| {message[:70]}")
        if not ok:
            suite.failures.append({"case": name, "kind": "source_fault", "result": "FAIL",
                                   "problems": [f"outcome={outcome} hits={state['hits']} message={message}"]})
    fault_db_unchanged = sha256(c09b.db_path) == before_fault_hash
    print(f"[{'PASS' if fault_db_unchanged else 'FAIL'}] R1_FAULT_DB_UNCHANGED 故障注入未改动数据库字节")
    if not fault_db_unchanged:
        suite.failures.append({"case": "R1_FAULT_DB_UNCHANGED", "kind": "fault", "result": "FAIL",
                               "problems": ["数据库字节在故障注入后变化"]})

    # ---------- 7. 答案隔离：运行期文件读取守卫 ----------
    real_open = builtins.open
    real_io_open = io.open
    opened_paths: list[str] = []

    def guarded_open(file, *open_args, **open_kwargs):
        try:
            candidate = Path(file).resolve()
        except Exception:  # noqa: BLE001
            candidate = None
        if candidate is not None:
            opened_paths.append(str(candidate))
            try:
                candidate.relative_to(EVALUATION_DIR)
                raise RuntimeError(f"规则核心在运行期尝试读取答案册目录：{candidate}")
            except ValueError:
                pass
        return real_open(file, *open_args, **open_kwargs)

    builtins.open = guarded_open
    io.open = guarded_open
    try:
        guarded_result = suite.screen_raw(valid_case.db_path, valid_case.load_report_payload)
        guard_ok = guarded_result["batch_status"] == "DEV_RULES_COMPUTED"
        guard_message = ""
    except Exception as exc:  # noqa: BLE001
        guard_ok = False
        guard_message = f"{type(exc).__name__}: {exc}"
    finally:
        builtins.open = real_open
        io.open = real_io_open
    suite.opened_paths_under_guard = opened_paths
    generator_imported = "make_demo_data" in sys.modules
    guard_live = len(opened_paths) > 0
    print(f"[{'PASS' if guard_ok and not generator_imported and guard_live else 'FAIL'}] ANSWER_ISOLATION_runtime "
          f"守卫下成功={guard_ok} 生成器未被导入={not generator_imported} 守卫实际拦截读取={len(opened_paths)} 次 "
          f"{guard_message[:60]}")
    if not guard_ok or generator_imported or not guard_live:
        suite.failures.append({"case": "ANSWER_ISOLATION_runtime", "kind": "isolation", "result": "FAIL",
                               "problems": [guard_message, f"make_demo_data imported={generator_imported}",
                                            f"opened_paths={len(opened_paths)}（守卫必须实际生效）"]})
    print(f"        答案册首次读取发生在核心输出已保存 {suite.answer_book_first_read_at_outputs} 份之后")
    ordering_ok = suite.answer_book_first_read_at_outputs == suite.outputs_saved
    print(f"[{'PASS' if ordering_ok else 'FAIL'}] ANSWER_ISOLATION_ordering "
          f"答案册读取在全部核心输出保存之后（{suite.answer_book_first_read_at_outputs}/{suite.outputs_saved}）")
    if not ordering_ok:
        suite.failures.append({"case": "ANSWER_ISOLATION_ordering", "kind": "isolation", "result": "FAIL",
                               "problems": ["答案册读取早于核心输出保存完成"]})

    # ---------- 8. 保护核对 ----------
    hashes_after = {str(path): sha256(path) for path in PROTECTED_FILES}
    changed = [path for path, digest in hashes_before.items() if hashes_after.get(path) != digest]
    if changed:
        suite.failures.append({"case": "INTEGRITY_protected_files", "kind": "integrity", "result": "FAIL",
                               "problems": [f"受保护文件在测试前后哈希不一致：{path}" for path in changed]})
        print("[FAIL] INTEGRITY_protected_files 哈希不一致：")
        for path in changed:
            print(f"        - {path}")
    else:
        print("[PASS] INTEGRITY_protected_files 受保护文件哈希一致")

    counts = {
        "inputs": len(suite.cases),
        "screened_success": sum(1 for item in suite.cases if item.screen_outcome == "computed"),
        "upstream_blocked": sum(1 for item in suite.cases if item.normalize_outcome == "upstream_blocked"),
        "interface_blocked": sum(1 for item in suite.cases if item.screen_outcome == "interface_blocked"),
        "runtime_errors": sum(1 for item in suite.cases if item.screen_outcome == "runtime_error"),
        "test_failures": len(suite.failures),
        "checker_cli_calls": suite.cli_calls,
        "screen_calls": suite.screen_calls,
        "reinvocations": suite.screen_calls - suite.outputs_saved,
    }
    summary = {
        "task": "TASK-006",
        "out_dir": str(container),
        "python": sys.version.split()[0],
        "sqlite": sqlite3.sqlite_version,
        "rule_core": {"path": str(SCREEN), "sha256": sha256(SCREEN), "version": screen.RULE_CORE_VERSION},
        "rules_sql": {path.stem.upper(): {"path": str(path), "sha256": sha256(path)} for path in RULES_SQL},
        "counts": counts,
        "cases": [
            {
                "case": item.name,
                "checker_exit": item.checker_exit,
                "normalize_outcome": item.normalize_outcome,
                "screen_outcome": item.screen_outcome,
                "hit_counts": item.screening["rule_hit_counts"] if item.screening else None,
                "message": item.screen_message,
            }
            for item in suite.cases
        ],
        "failures": suite.failures,
        "protected_files_count": len(PROTECTED_FILES),
        "protected_files_unchanged": not changed,
        "answer_isolation": {
            "static_core_and_sql_clean": not static_hits,
            "runtime_guard_passed": guard_ok,
            "runtime_guard_intercepted_reads": len(suite.opened_paths_under_guard),
            "generator_not_imported": not generator_imported,
            "answer_book_first_read_after_outputs_saved": suite.answer_book_first_read_at_outputs,
            "outputs_saved": suite.outputs_saved,
            "opened_paths_under_guard": suite.opened_paths_under_guard,
        },
        "notes": [
            "核心输出保存完成后才读取答案册；核心不接收答案路径、不读 evaluation、不导入生成器。",
            "质量阻断与预期接口阻断不计为测试失败，也不写成“成功且零命中”。",
            "规则输出是待核实线索，不含 AI 判断、严重等级或人类判断预填值。",
        ],
    }
    (container / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (container / "command_log.json").write_text(
        json.dumps(suite.command_log, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    print()
    print(
        f"输入 {counts['inputs']} 个：规则成功 {counts['screened_success']}、上游阻断 "
        f"{counts['upstream_blocked']}、接口阻断 {counts['interface_blocked']}、运行错误 "
        f"{counts['runtime_errors']}；核心调用 {counts['screen_calls']} 次（复调用 {counts['reinvocations']}）；"
        f"checker CLI {counts['checker_cli_calls']} 次；测试失败 {counts['test_failures']}"
    )
    print(f"受保护文件前后哈希一致: {summary['protected_files_unchanged']}（{summary['protected_files_count']} 个）")
    print(f"汇总: {container / 'summary.json'}")
    if suite.failures:
        print("存在断言失败或保护失败，返回非零退出码（不输出总体成功结论）。")
        return 1
    print("全部预期成立。")
    return 0


if __name__ == "__main__":
    sys.exit(main())

