#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""TASK-005 SQLite 装载开发验证。

用法：

    & 'D:\\harness work\\audit_sa_project\\.venv\\Scripts\\python.exe' -B `
      'D:\\harness work\\audit_sa_project\\tests\\test_store_demo_sqlite.py' `
      --out-dir 'D:\\harness work\\audit_sa_project\\data\\sqlite_tests\\task005_01'

每个输入：复制到新测试容器 → 真实 checker CLI 生成质量报告 → 冻结标准化 API → 装载进新数据库 →
从数据库读取完成行数、金额、逐笔分摊、已结算总额、净额、PRAGMA 与往返核对 → 断言。
质量阻断、装载阻断、运行失败与测试失败分开统计；任何断言或保护失败使进程非零。
退出码：0 全部符合预期；1 断言或保护失败；2 参数或环境问题。
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import importlib.util
import json
import shutil
import sqlite3
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SQLITE_TESTS_ROOT = (PROJECT_ROOT / "data" / "sqlite_tests").resolve()
CHECKER = PROJECT_ROOT / "scripts" / "check_demo_data.py"
RECEIVER = PROJECT_ROOT / "scripts" / "receive_demo_batch.py"
GENERATOR = PROJECT_ROOT / "scripts" / "make_demo_data.py"
NORMALIZER = PROJECT_ROOT / "scripts" / "normalize_demo_data.py"
STORE = PROJECT_ROOT / "scripts" / "store_demo_sqlite.py"
SCHEMA = PROJECT_ROOT / "sql" / "demo_schema.sql"

DEMO_DIR = PROJECT_ROOT / "data" / "generated" / "demo_20260930_a"
MAIN_INPUT = DEMO_DIR / "input" / "erp_export.xlsx"
FIXTURE_DIR = DEMO_DIR / "fixtures"

CONVERTIBLE_FIXTURES = (
    "C01", "C02", "C03", "C04", "C05", "C06", "C07", "C08", "C09", "C10",
    "C06b", "C07b", "C09b", "C10b",
)
BLOCKED_FIXTURES = ("C11", "C12", "C13", "C14a", "C14b", "C14c", "C14d", "C14e")

SHEET_COLUMNS = {
    "suppliers": ["supplier_id", "supplier_name"],
    "payables": ["ap_id", "ap_number", "supplier_id", "ap_date", "payable_amount", "currency"],
    "payments": [
        "payment_id", "transaction_type", "original_payment_id", "payee_supplier_id",
        "payment_date", "amount", "currency", "settlement_status", "bank_reference",
    ],
    "allocations": ["allocation_id", "payment_id", "ap_id", "allocated_amount", "original_allocation_id"],
}

PROTECTED_FILES = [
    CHECKER,
    RECEIVER,
    GENERATOR,
    NORMALIZER,
    STORE,
    SCHEMA,
    PROJECT_ROOT / "docs" / "SPEC_V1.md",
    PROJECT_ROOT / "docs" / "AUDIT_CONTRACT_DRAFT.md",
    PROJECT_ROOT / "docs" / "EXEC_REPORT_004_R1.md",
    PROJECT_ROOT / "docs" / "NORMALIZATION_GUIDE.md",
    PROJECT_ROOT / "tests" / "test_normalize_demo_data.py",
    MAIN_INPUT,
    *sorted(FIXTURE_DIR.glob("*.xlsx")),
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fail(message: str) -> None:
    print(f"RUNTIME_ERROR: {message}", file=sys.stderr)
    sys.exit(2)


def load_module(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def query(db_path: Path, sql: str, params: tuple = ()):
    connection = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    try:
        return connection.execute(sql, params).fetchall()
    finally:
        connection.close()


class Case:
    def __init__(self, name: str, case_dir: Path, input_path: Path):
        self.name = name
        self.case_dir = case_dir
        self.input_path = input_path
        self.checker_exit: int | None = None
        self.report: dict | None = None
        self.normalize_outcome: str | None = None
        self.normalize_message = ""
        self.normalized: dict | None = None
        self.load_outcome: str | None = None
        self.load_message = ""
        self.reasons: list[str] = []
        self.load_report: dict | None = None
        self.db_path: Path | None = None
        self.stdout_log: Path | None = None
        self.stderr_log: Path | None = None

    def db(self) -> Path:
        assert self.db_path is not None
        return self.db_path


class Suite:
    def __init__(self, container: Path, normalizer, store):
        self.container = container
        self.normalizer = normalizer
        self.store = store
        self.cases_dir = container / "cases"
        self.cases_dir.mkdir(parents=True, exist_ok=True)
        self.cli_calls = 0
        self.cases: list[Case] = []
        self.failures: list[dict] = []
        self.command_log: list[dict] = []

    def checker_call(self, name: str, workbook: Path, case_dir: Path) -> tuple[int, dict | None]:
        self.cli_calls += 1
        report_path = case_dir / "checker_report.json"
        stdout_log = case_dir / f"{self.cli_calls:03d}_{name}.checker.stdout.txt"
        stderr_log = case_dir / f"{self.cli_calls:03d}_{name}.checker.stderr.txt"
        command = [sys.executable, "-B", str(CHECKER), "--workbook", str(workbook), "--report", str(report_path)]
        with stdout_log.open("w", encoding="utf-8") as out, stderr_log.open("w", encoding="utf-8") as err:
            exit_code = subprocess.run(command, stdout=out, stderr=err).returncode
        payload = json.loads(report_path.read_text(encoding="utf-8")) if report_path.exists() else None
        self.command_log.append(
            {
                "call_index": self.cli_calls,
                "case": name,
                "command": command,
                "exit_code": exit_code,
                "report": str(report_path),
                "stdout_log": str(stdout_log),
                "stderr_log": str(stderr_log),
            }
        )
        return exit_code, payload

    def prepare(self, name: str, source: Path | None, build=None) -> Case:
        case_dir = self.cases_dir / name
        case_dir.mkdir(parents=True, exist_ok=True)
        input_path = case_dir / "input.xlsx"
        if build is not None:
            build(input_path)
        else:
            shutil.copyfile(source, input_path)
        case = Case(name, case_dir, input_path)
        exit_code, report = self.checker_call(name, input_path, case_dir)
        case.checker_exit = exit_code
        case.report = report
        case.stdout_log = case_dir / f"{self.cli_calls:03d}_{name}.checker.stdout.txt"
        case.stderr_log = case_dir / f"{self.cli_calls:03d}_{name}.checker.stderr.txt"
        try:
            normalized = self.normalizer.normalize_workbook(input_path, report)
            case.normalize_outcome = "normalized"
            case.normalized = normalized
            (case_dir / "normalized.json").write_text(
                json.dumps(normalized, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8"
            )
        except self.normalizer.NormalizationBlocked as exc:
            case.normalize_outcome = "quality_blocked"
            case.normalize_message = str(exc)
        except self.normalizer.NormalizationError as exc:
            case.normalize_outcome = "normalization_error"
            case.normalize_message = str(exc)
        self.cases.append(case)
        return case

    def load(self, case: Case, normalized=None, patches=None, db_name: str = "demo.sqlite") -> Case:
        module = self.store
        payload = case.normalized if normalized is None else normalized
        if patches:
            originals = {}
            for attr, factory in patches.items():
                originals[attr] = getattr(module, attr)
                setattr(module, attr, factory(originals[attr]))
        else:
            originals = {}
        db_path = case.case_dir / db_name
        try:
            report = module.load_normalized_to_sqlite(payload, db_path)
            case.load_outcome = "loaded"
            case.load_report = report
            case.db_path = db_path
            (case.case_dir / "load_report.json").write_text(
                json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8"
            )
        except module.SqliteLoadBlocked as exc:
            case.load_outcome = "load_blocked"
            case.load_message = str(exc)
            case.reasons = list(getattr(exc, "reasons", []))
        except module.SqliteLoadError as exc:
            case.load_outcome = "load_error"
            case.load_message = str(exc)
            case.reasons = list(getattr(exc, "reasons", []))
        finally:
            for attr, original in originals.items():
                setattr(module, attr, original)
        return case

    def judge(self, case: Case, *, outcome: str, checks: dict, kind: str) -> Case:
        problems: list[str] = []
        if case.load_outcome != outcome:
            problems.append(f"装载结果期望 {outcome}，实际 {case.load_outcome}（{case.load_message}）")
        for label, predicate in checks.items():
            try:
                ok = bool(predicate(case))
            except Exception as exc:  # noqa: BLE001
                ok = False
                problems.append(f"{label} 断言时异常：{exc!r}")
            if not ok:
                problems.append(f"断言失败：{label}")
        record = {
            "case": case.name,
            "kind": kind,
            "checker_exit": case.checker_exit,
            "normalize_outcome": case.normalize_outcome,
            "load_outcome": case.load_outcome,
            "message": case.load_message,
            "database": str(case.db_path) if case.db_path else None,
            "result": "PASS" if not problems else "FAIL",
            "problems": problems,
        }
        if problems:
            self.failures.append(record)
        print(f"[{'PASS' if not problems else 'FAIL'}] {case.name} load={case.load_outcome}"
              + ("" if not problems else "  | " + " ; ".join(problems)))
        return case


def build_from_spec(path: Path, tables: dict) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in SHEET_COLUMNS:
        worksheet = workbook.create_sheet(title=sheet)
        worksheet.append(SHEET_COLUMNS[sheet])
        for record in tables.get(sheet, []):
            worksheet.append([record.get(column) for column in SHEET_COLUMNS[sheet]])
    workbook.save(path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="TASK-005 SQLite 装载开发验证")
    parser.add_argument("--out-dir", required=True, help="data/sqlite_tests 下尚不存在的测试容器")
    args = parser.parse_args(argv)

    container = Path(args.out_dir)
    container = container.resolve() if container.is_absolute() else (Path.cwd() / container).resolve()
    try:
        container.relative_to(SQLITE_TESTS_ROOT)
    except ValueError:
        fail(f"测试容器必须位于 {SQLITE_TESTS_ROOT} 之内：{container}")
    if container == SQLITE_TESTS_ROOT:
        fail("测试容器不能是 data/sqlite_tests 本身")
    if container.exists():
        fail(f"测试容器已存在，拒绝覆盖历史：{container}")
    for required in (CHECKER, NORMALIZER, STORE, SCHEMA, MAIN_INPUT):
        if not required.exists():
            fail(f"缺少必需输入：{required}")

    hashes_before = {str(p): sha256(p) for p in PROTECTED_FILES}
    container.mkdir(parents=True)
    normalizer = load_module(NORMALIZER, "normalizer_under_test")
    store = load_module(STORE, "store_under_test")
    suite = Suite(container, normalizer, store)
    print(f"测试容器: {container}")
    print(f"Python {sys.version.split()[0]} / SQLite {store.sqlite_version()} / 装载版本 {store.STORE_VERSION}")
    print()

    def load_checks(case: Case) -> dict:
        return {
            "状态为 DEV_SQLITE_VERIFIED": lambda c: c.load_report["batch_status"] == "DEV_SQLITE_VERIFIED",
            "开发模式且未放行": lambda c: c.load_report["mode"] == "DEVELOPMENT_TEST_ONLY"
            and c.load_report["next_step_allowed"] is False,
            "往返逐项精确相等": lambda c: c.load_report["round_trip_exact"] is True,
            "来源记录覆盖每条业务行": lambda c: c.load_report["source_record_counts"]
            == c.load_report["row_counts"],
            "PRAGMA 检查通过": lambda c: c.load_report["foreign_key_check"] == "clean"
            and c.load_report["integrity_check"] == "ok",
            "净额检查全通过": lambda c: all(c.load_report["net_checks"].values()),
            "库文件存在且哈希一致": lambda c: c.db().exists()
            and sha256(c.db()) == c.load_report["database_sha256"],
        }

    # ---------- 1. 主文件 ----------
    main_case = suite.prepare("MAIN", MAIN_INPUT)
    suite.load(main_case)
    suite.judge(
        main_case, outcome="loaded", kind="expected_success",
        checks={
            **load_checks(main_case),
            "四表行数 20/993/1000/1005": lambda c: c.load_report["row_counts"]
            == {"suppliers": 20, "payables": 993, "payments": 1000, "allocations": 1005},
            "SQL 金额与标准化一致": lambda c: all(v["equal"] for v in c.load_report["table_totals"].values()),
            "主表净额 267249329 分": lambda c: c.load_report["master_net"] == 267249329,
            "已知关联 266249329 分": lambda c: c.load_report["known_payable_link_net"] == 266249329,
            "缺失关联 1000000 分": lambda c: c.load_report["missing_payable_link_net"] == 1000000,
            "非结算组未漏": lambda c: len(c.load_report["payment_groups"]) == 5,
        },
    )
    main_normalized = main_case.normalized

    # ---------- 2. 14 份质量有效 fixture ----------
    for fixture in CONVERTIBLE_FIXTURES:
        case = suite.prepare(fixture, FIXTURE_DIR / f"{fixture}.xlsx")
        suite.load(case)
        checks = dict(load_checks(case))
        checks["checker 退出 0"] = lambda c: c.checker_exit == 0
        if fixture == "C03":
            checks["C03 主表净额 1000000 分（未因两条分摊翻倍）"] = lambda c: c.load_report["master_net"] == 1000000
            checks["C03 分摊合计仍为 1000000 分"] = (
                lambda c: c.load_report["table_totals"]["allocations"]["sql_sum"] == 1000000
            )
        suite.judge(case, outcome="loaded", kind="expected_success", checks=checks)

    def case_of(name: str) -> Case:
        return next(c for c in suite.cases if c.name == name)

    c01_case = case_of("C01")
    c01_normalized = c01_case.normalized

    # ---------- 3. 定点保留行为（直接从数据库查回） ----------
    c07 = case_of("C07")
    suite.judge(
        c07, outcome="loaded", kind="expected_success",
        checks={
            "空 ap_id 原样保留（NULL）": lambda c: query(c.db(), "SELECT ap_id FROM allocations") == [(None,)],
            "缺失净额 500000 分且来自空 ap_id": lambda c: c.load_report["missing_payable_link_net"] == 500000
            and c.load_report["missing_payable_link_detail"]["empty_ap_id"] == 500000,
        },
    )
    c07b = case_of("C07b")
    suite.judge(
        c07b, outcome="loaded", kind="expected_success",
        checks={
            "部分未匹配：一条 AP01、一条 NULL": lambda c: query(
                c.db(), "SELECT ap_id FROM allocations ORDER BY allocation_id"
            ) == [("AP01",), (None,)],
            "已知 300000 分、缺失 200000 分": lambda c: c.load_report["known_payable_link_net"] == 300000
            and c.load_report["missing_payable_link_net"] == 200000,
        },
    )
    c08 = case_of("C08")
    suite.judge(
        c08, outcome="loaded", kind="expected_success",
        checks={
            "台账外 ap_id 原样保留": lambda c: query(c.db(), "SELECT ap_id FROM allocations") == [("AP_MISSING",)],
            "缺失净额来自台账未找到": lambda c: c.load_report["missing_payable_link_detail"]["ap_id_not_in_ledger"]
            == 500000,
        },
    )
    c09 = case_of("C09")
    suite.judge(
        c09, outcome="loaded", kind="expected_success",
        checks={
            "收款对象差异入库": lambda c: query(
                c.db(), "SELECT p.payee_supplier_id, pa.supplier_id FROM payments p, payables pa"
            ) == [("S002", "S001")],
        },
    )
    for name, amount in (("C06", 500000), ("C06b", 100000)):
        case = case_of(name)
        suite.judge(
            case, outcome="loaded", kind="expected_success",
            checks={
                "退款关联与分析引用入库": lambda c: query(
                    c.db(),
                    "SELECT transaction_type, original_payment_id FROM payments WHERE payment_id = 'R01'",
                ) == [("REFUND", "P02")]
                and query(
                    c.db(),
                    "SELECT original_allocation_id FROM allocations WHERE allocation_id = 'AR01'",
                ) == [("A02",)],
                "退款金额为正数": lambda c, a=amount: query(
                    c.db(), "SELECT amount_cents FROM payments WHERE payment_id = 'R01'"
                ) == [(a,)],
            },
        )
    c10 = case_of("C10")
    suite.judge(
        c10, outcome="loaded", kind="expected_success",
        checks={
            "三种非结算状态入库": lambda c: [row[0] for row in query(
                c.db(), "SELECT settlement_status FROM payments ORDER BY payment_id"
            )] == ["PENDING", "FAILED", "VOIDED"],
            "已结算总额为 0": lambda c: c.load_report["settled_pay_master_total"] == 0
            and c.load_report["settled_refund_master_total"] == 0,
        },
    )
    c10b = case_of("C10b")
    suite.judge(
        c10b, outcome="loaded", kind="expected_success",
        checks={
            "非结算退款入库且银行标识为空": lambda c: query(
                c.db(),
                "SELECT settlement_status, bank_reference FROM payments WHERE transaction_type = 'REFUND'",
            ) == [("PENDING", None)],
            "非结算退款不计入资金净额": lambda c: c.load_report["settled_refund_master_total"] == 0
            and c.load_report["master_net"] == 1000000,
        },
    )

    # ---------- 4. 8 份质量阻断 fixture ----------
    for fixture in BLOCKED_FIXTURES:
        case = suite.prepare(fixture, FIXTURE_DIR / f"{fixture}.xlsx")
        suite.load(case)
        suite.judge(
            case, outcome="load_blocked", kind="expected_quality_block",
            checks={
                "标准化阶段已阻断": lambda c: c.normalize_outcome == "quality_blocked",
                "未建立数据库": lambda c: not (c.case_dir / "demo.sqlite").exists(),
                "保留真实 checker 问题": lambda c: c.checker_exit == 1
                and bool(c.report and c.report.get("issue_counts_by_code")),
                "无装载报告": lambda c: c.load_report is None,
            },
        )

    # ---------- 5. 小正例一：无资金事件 ----------
    def build_no_funds(path: Path) -> None:
        shutil.copyfile(FIXTURE_DIR / "C01.xlsx", path)
        workbook = load_workbook(path)
        for sheet in ("payments", "allocations"):
            del workbook[sheet]
            worksheet = workbook.create_sheet(sheet)
            worksheet.append(SHEET_COLUMNS[sheet])
        workbook.save(path)

    p1 = suite.prepare("P1_no_funds", None, build=build_no_funds)
    suite.load(p1)
    suite.judge(
        p1, outcome="loaded", kind="expected_success",
        checks={
            **load_checks(p1),
            "checker 通过": lambda c: c.checker_exit == 0,
            "资金表为空": lambda c: c.load_report["row_counts"]["payments"] == 0
            and c.load_report["row_counts"]["allocations"] == 0,
            "资金净额为 0": lambda c: c.load_report["master_net"] == 0,
            "应付保留 500000 分": lambda c: query(
                c.db(), "SELECT payable_amount_cents FROM payables"
            ) == [(500000,)],
        },
    )

    # ---------- 6. 小正例二：标识保真与来源坐标 ----------
    def build_identifiers(path: Path) -> None:
        shutil.copyfile(FIXTURE_DIR / "C01.xlsx", path)
        workbook = load_workbook(path)
        workbook["suppliers"]["A2"] = "S'01"
        workbook["payables"]["A2"] = "0009"
        workbook["payables"]["B2"] = " O'Brien "
        workbook["payables"]["C2"] = "S'01"
        workbook["payments"]["A2"] = "0001"
        workbook["payments"]["D2"] = "S'01"
        workbook["payments"]["I2"] = " B 001 "
        workbook["allocations"]["B2"] = "0001"
        workbook["allocations"]["C2"] = "0009"
        workbook.save(path)

    p2 = suite.prepare("P2_identifier_fidelity", None, build=build_identifiers)
    suite.load(p2)
    suite.judge(
        p2, outcome="loaded", kind="expected_success",
        checks={
            **load_checks(p2),
            "checker 通过": lambda c: c.checker_exit == 0,
            "前导零与单引号 ID 原样入库": lambda c: query(c.db(), "SELECT ap_id FROM payables") == [("0009",)]
            and sorted(row[0] for row in query(c.db(), "SELECT supplier_id FROM suppliers")) == ["S'01", "S002"],
            "首尾空格 ID 原样入库": lambda c: query(c.db(), "SELECT ap_number FROM payables") == [(" O'Brien ",)]
            and query(c.db(), "SELECT bank_reference FROM payments") == [(" B 001 ",)],
            "外键引用一致": lambda c: query(
                c.db(), "SELECT COUNT(*) FROM allocations a JOIN payments p ON p.payment_id = a.payment_id"
            ) == [(1,)],
            "来源坐标保留": lambda c: query(
                c.db(), "SELECT sheet, excel_row FROM source_records WHERE table_name = 'payables'"
            ) == [("payables", 2)],
        },
    )

    # ---------- 7. 小正例三：退款与原分摊引用放在原记录之前 ----------
    def build_refund_first(path: Path) -> None:
        shutil.copyfile(FIXTURE_DIR / "C06.xlsx", path)
        workbook = load_workbook(path)
        for sheet, key in (("payments", "payment_id"), ("allocations", "allocation_id")):
            worksheet = workbook[sheet]
            header = [cell.value for cell in worksheet[1]]
            rows = [[cell.value for cell in row] for row in worksheet.iter_rows(min_row=2)]
            index = header.index(key)
            rows.sort(key=lambda row: 0 if str(row[index]).startswith("R") else 1)
            del workbook[sheet]
            recreated = workbook.create_sheet(sheet)
            recreated.append(header)
            for row in rows:
                recreated.append(row)
        workbook.save(path)

    p3 = suite.prepare("P3_refund_before_original", None, build=build_refund_first)
    suite.load(p3)
    suite.judge(
        p3, outcome="loaded", kind="expected_success",
        checks={
            **load_checks(p3),
            "checker 通过": lambda c: c.checker_exit == 0,
            "Excel 中退款行在前": lambda c: query(
                c.db(),
                "SELECT payment_id FROM payments ORDER BY (SELECT excel_row FROM source_records "
                "WHERE table_name = 'payments' AND record_id = payments.payment_id) LIMIT 1",
            ) == [("R01",)],
            "自关联引用仍完整": lambda c: query(
                c.db(), "SELECT original_payment_id FROM payments WHERE payment_id = 'R01'"
            ) == [("P02",)],
        },
    )

    # ---------- 8. 路径拒绝 ----------
    same_path_case = Case("N1_same_path_again", main_case.case_dir, main_case.input_path)
    suite.cases.append(same_path_case)
    same_path_case.normalized = main_case.normalized
    same_path_case.db_path = main_case.db()
    before_hash = sha256(main_case.db())
    module = suite.store
    try:
        module.load_normalized_to_sqlite(main_case.normalized, main_case.db())
        same_path_case.load_outcome = "loaded"
    except module.SqliteLoadBlocked as exc:
        same_path_case.load_outcome = "load_blocked"
        same_path_case.load_message = str(exc)
    except module.SqliteLoadError as exc:
        same_path_case.load_outcome = "load_error"
        same_path_case.load_message = str(exc)
    suite.judge(
        same_path_case, outcome="load_blocked", kind="expected_load_block",
        checks={
            "写入前拒绝": lambda c: "已存在" in c.load_message,
            "原库字节未变": lambda c: sha256(main_case.db()) == before_hash
            and before_hash == main_case.load_report["database_sha256"],
        },
    )

    outside_case = Case("N2_out_of_range_path", main_case.case_dir, main_case.input_path)
    suite.cases.append(outside_case)
    outside_case.normalized = main_case.normalized
    outside_target = PROJECT_ROOT / "data" / "sqlite_tests_outside" / "demo.sqlite"
    try:
        module.load_normalized_to_sqlite(main_case.normalized, outside_target)
        outside_case.load_outcome = "loaded"
    except module.SqliteLoadBlocked as exc:
        outside_case.load_outcome = "load_blocked"
        outside_case.load_message = str(exc)
    except module.SqliteLoadError as exc:
        outside_case.load_outcome = "load_error"
        outside_case.load_message = str(exc)
    suite.judge(
        outside_case, outcome="load_blocked", kind="expected_load_block",
        checks={
            "越界路径在写入前拒绝": lambda c: "必须位于" in c.load_message,
            "未创建目标": lambda c: not outside_target.exists()
            and not outside_target.parent.exists(),
        },
    )

    # ---------- 9. 内存注入 ----------
    def injected(name: str, mutate, note: str, *, db_name: str = "demo.sqlite") -> Case:
        case_dir = suite.cases_dir / name
        case_dir.mkdir(parents=True, exist_ok=True)
        payload = copy.deepcopy(main_normalized)
        mutate(payload)
        (case_dir / "injected_normalized.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8"
        )
        (case_dir / "injection_note.txt").write_text(
            note + "\n（原标准化结果来自 MAIN 案例的 normalized.json，未修改上游文件）\n", encoding="utf-8"
        )
        case = Case(name, case_dir, main_case.input_path)
        case.normalized = payload
        suite.cases.append(case)
        return case

    i1 = injected(
        "I1_amount_string",
        lambda p: p["tables"]["payments"]["records"][0]["data"].__setitem__("amount_cents", "5000"),
        "内存注入：把一条付款的 amount_cents 改成字符串 \"5000\"",
    )
    suite.load(i1, normalized=i1.normalized)
    suite.judge(
        i1, outcome="load_blocked", kind="expected_load_block",
        checks={
            "金额字符串被拒绝": lambda c: "金额必须是整数分" in c.load_message,
            "未建立数据库": lambda c: not (c.case_dir / "demo.sqlite").exists(),
        },
    )

    i2 = injected(
        "I2_amount_bool",
        lambda p: p["tables"]["payables"]["records"][0]["data"].__setitem__("payable_amount_cents", True),
        "内存注入：把一条应付金额改成 bool",
    )
    suite.load(i2, normalized=i2.normalized)
    suite.judge(
        i2, outcome="load_blocked", kind="expected_load_block",
        checks={"bool 金额被拒绝": lambda c: "金额必须是整数分" in c.load_message},
    )

    i3 = injected(
        "I3_duplicate_primary_key",
        lambda p: p["tables"]["suppliers"]["records"].__setitem__(1, copy.deepcopy(
            p["tables"]["suppliers"]["records"][0]
        )),
        "内存注入：把第二条供应商记录替换为第一条的完整副本（业务主键重复）",
    )
    suite.load(i3, normalized=i3.normalized)
    suite.judge(
        i3, outcome="load_error", kind="expected_load_error",
        checks={
            "数据库约束阻止重复主键": lambda c: "insert_business_rows:suppliers" in c.load_message
            or "UNIQUE" in c.load_message,
            "未产生目标库": lambda c: not (c.case_dir / "demo.sqlite").exists(),
            "保留暂存失败现场": lambda c: (c.case_dir / "demo.sqlite.staging").exists(),
        },
    )

    i4 = injected(
        "I4_missing_funds_link",
        lambda p: (
            p["tables"]["allocations"]["records"][0]["data"].__setitem__("payment_id", "PM_NOT_EXIST"),
            p["tables"]["allocations"]["records"][0]["cells"]["payment_id"].__setitem__(
                "standard_value", "PM_NOT_EXIST"
            ),
        ),
        "内存注入：把一条分摊的 payment_id 改成不存在的付款（缺失强制资金关联）",
    )
    suite.load(i4, normalized=i4.normalized)
    suite.judge(
        i4, outcome="load_error", kind="expected_load_error",
        checks={
            "外键阻止缺失资金关联": lambda c: "insert_business_rows:allocations" in c.load_message
            or "FOREIGN KEY" in c.load_message,
            "未产生目标库": lambda c: not (c.case_dir / "demo.sqlite").exists(),
        },
    )

    def insert_fails(_original):
        def patched(connection, sheet, records):
            raise RuntimeError("模拟：写入来源记录时故障")

        return patched

    i5 = injected(
        "I5_mid_transaction_failure",
        lambda p: None,
        "内存注入：写入来源记录阶段模拟故障（业务行已插入，随后回滚）",
    )
    suite.load(i5, normalized=i5.normalized, patches={"insert_source_records": insert_fails})
    staging = i5.case_dir / "demo.sqlite.staging"
    suite.judge(
        i5, outcome="load_error", kind="expected_load_error",
        checks={
            "报运行错误且无成功报告": lambda c: c.load_report is None
            and "insert_source_records" in c.load_message,
            "未产生目标库": lambda c: not (c.case_dir / "demo.sqlite").exists(),
            "保留暂存库": lambda c: staging.exists(),
            "暂存库无半批数据": lambda c: all(
                query(staging, f"SELECT COUNT(*) FROM {table}") == [(0,)]
                for table in ("suppliers", "payables", "payments", "allocations", "source_records", "run_metadata")
            ),
        },
    )

    # ---------- 10. R1 新增：四项定点修订的回归 ----------
    def r1_injected(name: str, mutate, note: str) -> Case:
        case_dir = suite.cases_dir / name
        case_dir.mkdir(parents=True, exist_ok=True)
        payload = copy.deepcopy(c01_normalized)
        mutate(payload)
        (case_dir / "injected_normalized.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8"
        )
        (case_dir / "injection_note.txt").write_text(
            note + "\n（原标准化结果来自本容器 C01 案例的 normalized.json，未修改上游文件）\n",
            encoding="utf-8",
        )
        case = Case(name, case_dir, c01_case.input_path)
        case.normalized = payload
        suite.cases.append(case)
        return case

    def reason_of(case: Case, text: str) -> bool:
        return text in case.load_message

    def reason_any(case: Case, text: str) -> bool:
        return any(text in item for item in case.reasons)

    def staging_zero(case: Case) -> bool:
        staging = case.case_dir / "demo.sqlite.staging"
        if not staging.exists():
            return False
        return all(
            query(staging, f"SELECT COUNT(*) FROM {table}") == [(0,)]
            for table in ("suppliers", "payables", "payments", "allocations", "source_records", "run_metadata")
        )

    # Q01：核对阶段失败必须在提交前回滚
    q01 = r1_injected(
        "Q01_verify_failure_rolls_back",
        lambda p: p["amount_reconciliation"]["table_totals"]["payables_payable_amount"].__setitem__(
            "normalized", p["amount_reconciliation"]["table_totals"]["payables_payable_amount"]["normalized"] + 1
        ),
        "内存注入：把应付合计的标准化预期数加 1 分（SQL 合计 500000，预期 500001），用于在数据库核对阶段触发失败",
    )
    suite.load(q01, normalized=q01.normalized)
    suite.judge(
        q01, outcome="load_error", kind="expected_load_error",
        checks={
            "核对阶段发现 1 分差异": lambda c: reason_any(c, "500001") and reason_any(c, "500000"),
            "失败发生在数据库核对阶段": lambda c: "数据库独立核对未通过" in c.load_message,
            "事务未提交": lambda c: reason_any(c, "数据事务已提交=False"),
            "暂存库六表全为 0": staging_zero,
            "无最终目标库": lambda c: not (c.case_dir / "demo.sqlite").exists(),
            "未返回成功报告": lambda c: c.load_report is None,
        },
    )

    # Q02：缺少单元格原值／类型／坐标必须写入前阻断
    q02 = r1_injected(
        "Q02_cell_missing_provenance",
        lambda p: [
            p["tables"]["suppliers"]["records"][0]["cells"]["supplier_id"].pop(key)
            for key in list(p["tables"]["suppliers"]["records"][0]["cells"]["supplier_id"])
            if key not in ("input_field", "target_field", "standard_value")
        ],
        "内存注入：supplier_id 单元格来源只保留 input_field／target_field／standard_value",
    )
    suite.load(q02, normalized=q02.normalized)
    suite.judge(
        q02, outcome="load_blocked", kind="expected_load_block",
        checks={
            "写入前阻断并给出字段路径": lambda c: reason_of(c, "缺少必需来源字段")
            and reason_of(c, "tables.suppliers.records[0].cells.supplier_id"),
            "无目标库": lambda c: not (c.case_dir / "demo.sqlite").exists(),
            "无暂存库": lambda c: not (c.case_dir / "demo.sqlite.staging").exists(),
        },
    )

    q02b = r1_injected(
        "Q02b_cell_wrong_coordinate",
        lambda p: p["tables"]["payables"]["records"][0]["cells"]["ap_id"].__setitem__("cell", "Z9"),
        "内存注入：payables 的 ap_id 单元格坐标改成与实际列号／行号不符的 Z9",
    )
    suite.load(q02b, normalized=q02b.normalized)
    suite.judge(
        q02b, outcome="load_blocked", kind="expected_load_block",
        checks={
            "坐标关系被核验": lambda c: reason_of(c, ".cell") and reason_of(c, "坐标"),
            "无暂存库": lambda c: not (c.case_dir / "demo.sqlite.staging").exists(),
        },
    )

    # Q03：发布阶段目标晚出现必须拒绝覆盖
    def publish_hook_writes_marker(marker: bytes):
        def factory(original):
            def patched(target):
                Path(target).write_bytes(marker)
                return original(target)

            return patched

        return factory

    marker = b"MARKER-Q03-EXISTING-TARGET"
    q03 = r1_injected("Q03_target_appears_before_publish", lambda p: None, "发布前钩子注入：在核对成功、发布之前创建目标文件")
    suite.load(q03, normalized=q03.normalized, patches={"pre_publish_hook": publish_hook_writes_marker(marker)})
    suite.judge(
        q03, outcome="load_error", kind="expected_load_error",
        checks={
            "发布阶段明确拒绝覆盖": lambda c: reason_of(c, "拒绝覆盖") and reason_any(c, "阶段=publish_database"),
            "标记字节与哈希原样保留": lambda c: (c.case_dir / "demo.sqlite").read_bytes() == marker,
            "未返回成功报告": lambda c: c.load_report is None,
            "如实说明事务已提交": lambda c: reason_any(c, "数据事务已提交=True")
            and reason_any(c, "目标文件已发布=False"),
            "暂存库保留已提交数据（不假报回滚）": lambda c: (c.case_dir / "demo.sqlite.staging").exists(),
        },
    )

    # Q04：运行元数据必须实际读取核验
    def metadata_not_written(_original):
        def patched(connection, rows):
            return None

        return patched

    def metadata_fingerprint_corrupted(original):
        def patched(connection, rows):
            mutated = []
            for key, value in rows:
                if key == "normalized_payload_sha256":
                    value = value[:-1] + ("0" if value[-1] != "0" else "1")
                mutated.append((key, value))
            return original(connection, mutated)

        return patched

    q04 = r1_injected("Q04_metadata_not_written", lambda p: None, "接缝注入：insert_run_metadata 不写任何行")
    suite.load(q04, normalized=q04.normalized, patches={"insert_run_metadata": metadata_not_written})
    suite.judge(
        q04, outcome="load_error", kind="expected_load_error",
        checks={
            "发现必需元数据缺失": lambda c: reason_any(c, "run_metadata 缺少必需键"),
            "暂存库六表全为 0": staging_zero,
            "无最终目标库": lambda c: not (c.case_dir / "demo.sqlite").exists(),
            "未返回成功报告": lambda c: c.load_report is None,
        },
    )

    q04b = r1_injected("Q04b_metadata_fingerprint_wrong", lambda p: None, "接缝注入：本次写入的对象指纹末位被改错一位")
    suite.load(
        q04b, normalized=q04b.normalized,
        patches={"insert_run_metadata": metadata_fingerprint_corrupted},
    )
    suite.judge(
        q04b, outcome="load_error", kind="expected_load_error",
        checks={
            "指纹差异被实际读取发现": lambda c: reason_any(c, "normalized_payload_sha256"),
            "暂存库六表全为 0": staging_zero,
        },
    )

    # ---------- 11. R2 新增：正常关闭异常不得被吞掉 ----------
    def close_fails_after_release(original):
        """包装 connect_database：close 先释放真实句柄再抛一次 sqlite3.OperationalError。"""

        def patched(path):
            real = original(path)

            class ClosingProxy:
                def __init__(self, inner):
                    object.__setattr__(self, "_inner", inner)

                def __getattr__(self, item):
                    return getattr(object.__getattribute__(self, "_inner"), item)

                def close(self):
                    object.__getattribute__(self, "_inner").close()
                    raise sqlite3.OperationalError("注入：正常关闭时抛出 sqlite3.OperationalError")

            return ClosingProxy(real)

        return patched

    publish_calls = {"count": 0}
    original_publish = store.publish_database

    def counting_publish(_original):
        def patched(staging, target):
            publish_calls["count"] += 1
            return original_publish(staging, target)

        return patched

    r2a = r1_injected(
        "R2A_close_failure_blocks_publish", lambda p: None,
        "接缝注入：connect_database 返回的代理在 close 时先释放真实句柄再抛 sqlite3.OperationalError",
    )
    publish_calls["count"] = 0
    suite.load(
        r2a, normalized=r2a.normalized,
        patches={"connect_database": close_fails_after_release, "publish_database": counting_publish},
    )
    suite.judge(
        r2a, outcome="load_error", kind="expected_load_error",
        checks={
            "抛 SqliteLoadError 且阶段为 close_database": lambda c: reason_of(c, "close_database"),
            "保留 OperationalError 与消息": lambda c: reason_any(c, "OperationalError")
            and reason_any(c, "注入：正常关闭时抛出 sqlite3.OperationalError"),
            "保留目标与暂存路径": lambda c: reason_any(c, str(c.case_dir / "demo.sqlite")),
            "没有继续发布": lambda c: publish_calls["count"] == 0,
            "目标文件不存在": lambda c: not (c.case_dir / "demo.sqlite").exists(),
            "无成功报告": lambda c: c.load_report is None,
            "如实说明提交为真、发布为假": lambda c: reason_any(c, "数据事务已提交=True")
            and reason_any(c, "目标文件已发布=False"),
            "已提交暂存库保留完整行数": lambda c: query(
                c.case_dir / "demo.sqlite.staging",
                "SELECT (SELECT COUNT(*) FROM suppliers), (SELECT COUNT(*) FROM payables), "
                "(SELECT COUNT(*) FROM payments), (SELECT COUNT(*) FROM allocations), "
                "(SELECT COUNT(*) FROM source_records), (SELECT COUNT(*) FROM run_metadata)",
            ) == [(2, 1, 1, 1, 5, 17)],
            "不假报回滚为零行": lambda c: reason_any(c, "不声称已回滚"),
        },
    )

    r2b = r1_injected(
        "R2B_cleanup_close_failure_keeps_original_error",
        lambda p: p["amount_reconciliation"]["table_totals"]["payables_payable_amount"].__setitem__(
            "normalized", p["amount_reconciliation"]["table_totals"]["payables_payable_amount"]["normalized"] + 1
        ),
        "接缝注入：应付预期合计 +1 分触发提交前核对失败，同时 connect_database 代理在清理关闭时再抛 SQLite 异常",
    )
    suite.load(r2b, normalized=r2b.normalized, patches={"connect_database": close_fails_after_release})
    suite.judge(
        r2b, outcome="load_error", kind="expected_load_error",
        checks={
            "主错误仍是数据库核对失败": lambda c: "数据库独立核对未通过" in c.load_message,
            "保留 500000 与 500001 差异": lambda c: reason_any(c, "500000") and reason_any(c, "500001"),
            "关闭失败只作附记": lambda c: reason_any(c, "清理时关闭连接失败"),
            "附记未取代主错误": lambda c: "数据库独立核对未通过" in c.load_message
            and not any(item.startswith("阶段=close_database") for item in c.reasons),
            "事务未提交且六表全为 0": lambda c: reason_any(c, "数据事务已提交=False") and staging_zero(c),
            "无目标库、未返回成功报告": lambda c: not (c.case_dir / "demo.sqlite").exists()
            and c.load_report is None,
        },
    )

    # ---------- 12. 保护核对 ----------
    hashes_after = {str(p): sha256(p) for p in PROTECTED_FILES}
    changed = [p for p, digest in hashes_before.items() if hashes_after.get(p) != digest]
    if changed:
        suite.failures.append(
            {
                "case": "INTEGRITY_protected_files",
                "kind": "integrity",
                "result": "FAIL",
                "problems": [f"受保护文件在测试前后哈希不一致：{p}" for p in changed],
            }
        )
        print("[FAIL] INTEGRITY_protected_files 哈希不一致：")
        for path in changed:
            print(f"        - {path}")
    else:
        print("[PASS] INTEGRITY_protected_files 受保护文件哈希一致")

    counts = {
        "cases": len(suite.cases),
        "loaded": sum(1 for c in suite.cases if c.load_outcome == "loaded"),
        "quality_blocked": sum(1 for c in suite.cases if c.load_outcome == "load_blocked"
                               and c.normalize_outcome == "quality_blocked"),
        "load_blocked": sum(1 for c in suite.cases if c.load_outcome == "load_blocked"),
        "load_errors": sum(1 for c in suite.cases if c.load_outcome == "load_error"),
        "expected_negative": sum(1 for c in suite.cases if c.load_outcome in ("load_blocked", "load_error")),
        "test_failures": len(suite.failures),
        "checker_cli_calls": suite.cli_calls,
    }
    counts["expected_negative_breakdown"] = {
        "quality_blocked_fixtures": counts["quality_blocked"],
        "load_blocked_or_path_rejected": counts["load_blocked"] - counts["quality_blocked"],
        "load_errors": counts["load_errors"],
    }
    summary = {
        "task": "TASK-005",
        "out_dir": str(container),
        "python": sys.version.split()[0],
        "sqlite": sqlite3.sqlite_version,
        "store": {"path": str(STORE), "sha256": sha256(STORE), "version": store.STORE_VERSION},
        "schema": {"path": str(SCHEMA), "sha256": sha256(SCHEMA), "version": store.SCHEMA_VERSION},
        "normalizer": {"path": str(NORMALIZER), "sha256": sha256(NORMALIZER)},
        "counts": counts,
        "cases": [
            {
                "case": c.name,
                "checker_exit": c.checker_exit,
                "normalize_outcome": c.normalize_outcome,
                "load_outcome": c.load_outcome,
                "message": c.load_message,
                "database": str(c.db_path) if c.db_path else None,
            }
            for c in suite.cases
        ],
        "failures": suite.failures,
        "protected_files_count": len(PROTECTED_FILES),
        "protected_files_unchanged": not changed,
        "notes": [
            "质量报告由真实 checker CLI 对当前输入副本重新生成；未复用旧 normalized.json 或旧 checks。",
            "注入只作用于内存副本或新测试库，未修改上游文件、真实源码或既有批次。",
            "装载成功不等于没有业务线索，也不是审计结论。",
        ],
    }
    (container / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (container / "command_log.json").write_text(
        json.dumps(suite.command_log, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    print()
    print(
        f"案例 {counts['cases']} 个：入库成功 {counts['loaded']}、质量阻断 {counts['quality_blocked']}、"
        f"装载阻断 {counts['load_blocked'] - counts['quality_blocked']}、运行失败 {counts['load_errors']}；"
        f"checker CLI 调用 {counts['checker_cli_calls']} 次；测试失败 {counts['test_failures']}"
    )
    print(f"受保护文件前后哈希一致: {summary['protected_files_unchanged']}（{summary['protected_files_count']} 个）")
    print(f"汇总: {container / 'summary.json'}")
    if suite.failures:
        print("存在断言失败或保护失败，返回非零退出码（不输出总体成功结论）。")
        return 1
    print("全部预期成立。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
