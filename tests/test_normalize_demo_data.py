#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""TASK-004 标准化核心开发验证。

用法：

    & 'D:\\harness work\\audit_sa_project\\.venv\\Scripts\\python.exe' -B `
      'D:\\harness work\\audit_sa_project\\tests\\test_normalize_demo_data.py' `
      --out-dir 'D:\\harness work\\audit_sa_project\\data\\standardization_tests\\task004_01'

对每个输入：复制到新测试容器 → 用真实 checker CLI 重新生成质量报告 → 调用
scripts/normalize_demo_data.py 的 normalize_workbook() → 断言转换结果或明确的阻断／运行错误。
只转换成功的输入才写 normalized.json；预期阻断与脚本自身失败分开记录。
退出码：0 全部符合预期；1 存在断言失败或保护失败；2 参数或环境问题。
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import importlib.util
import json
import shutil
import subprocess
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
STANDARDIZATION_ROOT = (PROJECT_ROOT / "data" / "standardization_tests").resolve()
CHECKER = PROJECT_ROOT / "scripts" / "check_demo_data.py"
RECEIVER = PROJECT_ROOT / "scripts" / "receive_demo_batch.py"
GENERATOR = PROJECT_ROOT / "scripts" / "make_demo_data.py"
NORMALIZER = PROJECT_ROOT / "scripts" / "normalize_demo_data.py"

DEMO_DIR = PROJECT_ROOT / "data" / "generated" / "demo_20260930_a"
MAIN_INPUT = DEMO_DIR / "input" / "erp_export.xlsx"
FIXTURE_DIR = DEMO_DIR / "fixtures"

CONVERTIBLE_FIXTURES = (
    "C01", "C02", "C03", "C04", "C05", "C06", "C07", "C08", "C09", "C10",
    "C06b", "C07b", "C09b", "C10b",
)
BLOCKED_FIXTURES = ("C11", "C12", "C13", "C14a", "C14b", "C14c", "C14d", "C14e")

SHEET_COLUMNS = {
    "suppliers": ["supplier_id", "supplier_name"],
    "payables": ["ap_id", "ap_number", "supplier_id", "ap_date", "payable_amount", "currency"],
    "payments": [
        "payment_id", "transaction_type", "original_payment_id", "payee_supplier_id",
        "payment_date", "amount", "currency", "settlement_status", "bank_reference",
    ],
    "allocations": ["allocation_id", "payment_id", "ap_id", "allocated_amount", "original_allocation_id"],
}

PROTECTED_FILES = [
    CHECKER,
    RECEIVER,
    GENERATOR,
    PROJECT_ROOT / "docs" / "SPEC_V1.md",
    PROJECT_ROOT / "docs" / "AUDIT_CONTRACT_DRAFT.md",
    PROJECT_ROOT / "docs" / "EXEC_REPORT_002_R2.md",
    PROJECT_ROOT / "docs" / "EXEC_REPORT_003_R3.md",
    PROJECT_ROOT / "tests" / "test_check_demo_data_regression.py",
    PROJECT_ROOT / "tests" / "test_receive_demo_batch.py",
    MAIN_INPUT,
    *sorted(FIXTURE_DIR.glob("*.xlsx")),
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


_USE_REPORT = object()


def fail(message: str) -> "None":
    print(f"RUNTIME_ERROR: {message}", file=sys.stderr)
    sys.exit(2)


def load_normalizer():
    spec = importlib.util.spec_from_file_location("normalizer_under_test", NORMALIZER)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


class CaseResult:
    def __init__(self, case: str, case_dir: Path, input_path: Path):
        self.case = case
        self.case_dir = case_dir
        self.input_path = input_path
        self.checker_exit: int | None = None
        self.report: dict | None = None
        self.outcome: str | None = None  # normalized / blocked / error
        self.result: dict | None = None
        self.message: str = ""
        self.reasons: list[str] = []
        self.normalized_path: Path | None = None
        self.stdout_log: Path | None = None
        self.stderr_log: Path | None = None

    def table(self, sheet: str) -> list[dict]:
        assert self.result is not None
        return self.result["tables"][sheet]["records"]

    def data_of(self, sheet: str) -> list[dict]:
        return [record["data"] for record in self.table(sheet)]


class Suite:
    def __init__(self, container: Path, normalizer):
        self.container = container
        self.normalizer = normalizer
        self.cases_dir = container / "cases"
        self.cases_dir.mkdir(parents=True, exist_ok=True)
        self.cli_calls = 0
        self.results: list[CaseResult] = []
        self.failures: list[dict] = []
        self.command_log: list[dict] = []

    def checker_call(self, case: str, workbook: Path, case_dir: Path) -> tuple[int, dict | None]:
        self.cli_calls += 1
        report = case_dir / "checker_report.json"
        stdout_log = case_dir / f"{self.cli_calls:03d}_{case}.checker.stdout.txt"
        stderr_log = case_dir / f"{self.cli_calls:03d}_{case}.checker.stderr.txt"
        command = [
            sys.executable, "-B", str(CHECKER),
            "--workbook", str(workbook), "--report", str(report),
        ]
        with stdout_log.open("w", encoding="utf-8") as out, stderr_log.open("w", encoding="utf-8") as err:
            exit_code = subprocess.run(command, stdout=out, stderr=err).returncode
        payload = json.loads(report.read_text(encoding="utf-8")) if report.exists() else None
        self.command_log.append(
            {
                "call_index": self.cli_calls,
                "case": case,
                "command": command,
                "exit_code": exit_code,
                "report": str(report),
                "stdout_log": str(stdout_log),
                "stderr_log": str(stderr_log),
            }
        )
        return exit_code, payload

    def evaluate(
        self,
        case: str,
        source: Path | None,
        *,
        build=None,
        report_override=_USE_REPORT,
        patches=None,
        extra=None,
    ) -> CaseResult:
        case_dir = self.cases_dir / case
        case_dir.mkdir(parents=True, exist_ok=True)
        input_path = case_dir / "input.xlsx"
        if build is not None:
            build(input_path)
        else:
            shutil.copyfile(source, input_path)

        result = CaseResult(case, case_dir, input_path)
        result.extra = extra or {}
        exit_code, report = self.checker_call(case, input_path, case_dir)
        result.checker_exit = exit_code
        result.report = report

        module = self.normalizer
        if patches:
            originals = {}
            for name, factory in patches.items():
                originals[name] = getattr(module, name)
                setattr(module, name, factory(originals[name]))
        else:
            originals = {}
        try:
            payload = report if report_override is _USE_REPORT else report_override
            normalized = module.normalize_workbook(input_path, payload)
            result.outcome = "normalized"
            result.result = normalized
            result.normalized_path = case_dir / "normalized.json"
            result.normalized_path.write_text(
                json.dumps(normalized, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
                encoding="utf-8",
            )
        except module.NormalizationBlocked as exc:
            result.outcome = "blocked"
            result.message = str(exc)
            result.reasons = list(getattr(exc, "reasons", []))
        except module.NormalizationError as exc:
            result.outcome = "error"
            result.message = str(exc)
            result.reasons = list(getattr(exc, "reasons", []))
        finally:
            for name, original in originals.items():
                setattr(module, name, original)
        self.results.append(result)
        return result

    def judge(self, result: CaseResult, *, outcome: str, checks: dict, expected_block: bool) -> CaseResult:
        problems: list[str] = []
        if result.outcome != outcome:
            problems.append(f"结果期望 {outcome}，实际 {result.outcome}（{result.message}）")
        for label, predicate in checks.items():
            try:
                ok = bool(predicate(result))
            except Exception as exc:  # noqa: BLE001
                ok = False
                problems.append(f"{label} 断言时异常：{exc!r}")
            if not ok:
                problems.append(f"断言失败：{label}")
        record = {
            "case": result.case,
            "kind": "expected_block" if expected_block else "expected_success",
            "checker_exit": result.checker_exit,
            "outcome": result.outcome,
            "message": result.message,
            "normalized_json": str(result.normalized_path) if result.normalized_path else None,
            "result": "PASS" if not problems else "FAIL",
            "problems": problems,
        }
        if problems:
            self.failures.append(record)
        print(f"[{'PASS' if not problems else 'FAIL'}] {result.case} outcome={result.outcome}"
              + ("" if not problems else "  | " + " ; ".join(problems)))
        return result


# ---------- 断言小工具 ----------
def totals(result: CaseResult) -> dict:
    assert result.result is not None
    return result.result["amount_reconciliation"]


def money(records: list[dict], field: str) -> list[int]:
    return [record[field] for record in records]


def no_yuan_amount_fields(result: CaseResult) -> bool:
    """标准数据不得保留同名的“元”业务字段。"""
    for sheet, fields in (("payables", ("payable_amount",)), ("payments", ("amount",)), ("allocations", ("allocated_amount",))):
        for record in result.data_of(sheet):
            for field in fields:
                if field in record:
                    return False
    return True


def all_amounts_int(result: CaseResult) -> bool:
    for sheet, field in (("payables", "payable_amount_cents"), ("payments", "amount_cents"), ("allocations", "allocated_amount_cents")):
        for record in result.data_of(sheet):
            value = record.get(field)
            if isinstance(value, bool) or not isinstance(value, int):
                return False
    return True


def build_from_spec(path: Path, tables: dict) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in SHEET_COLUMNS:
        worksheet = workbook.create_sheet(title=sheet)
        worksheet.append(SHEET_COLUMNS[sheet])
        for record in tables.get(sheet, []):
            worksheet.append([record.get(column) for column in SHEET_COLUMNS[sheet]])
    workbook.save(path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="TASK-004 标准化核心开发验证")
    parser.add_argument("--out-dir", required=True, help="data/standardization_tests 下尚不存在的测试容器")
    args = parser.parse_args(argv)

    container = Path(args.out_dir)
    container = container.resolve() if container.is_absolute() else (Path.cwd() / container).resolve()
    try:
        container.relative_to(STANDARDIZATION_ROOT)
    except ValueError:
        fail(f"测试容器必须位于 {STANDARDIZATION_ROOT} 之内：{container}")
    if container == STANDARDIZATION_ROOT:
        fail("测试容器不能是 data/standardization_tests 本身")
    if container.exists():
        fail(f"测试容器已存在，拒绝覆盖历史：{container}")
    for required in (CHECKER, RECEIVER, NORMALIZER, MAIN_INPUT):
        if not required.exists():
            fail(f"缺少必需输入：{required}")

    hashes_before = {str(p): sha256(p) for p in PROTECTED_FILES}
    module_hashes_before = {str(NORMALIZER): sha256(NORMALIZER), str(CHECKER): sha256(CHECKER)}

    container.mkdir(parents=True)
    normalizer = load_normalizer()
    suite = Suite(container, normalizer)
    print(f"测试容器: {container}")
    print(f"被验证模块: {NORMALIZER}（版本 {normalizer.NORMALIZATION_VERSION}）")
    print()

    # ---------- 1. 主演示文件 ----------
    main_result = suite.evaluate("MAIN", MAIN_INPUT)
    suite.judge(
        main_result,
        outcome="normalized",
        expected_block=False,
        checks={
            "四表行数 20/993/1000/1005": lambda r: {
                sheet: r.result["row_accounting"][sheet]["normalized_rows"]
                for sheet in ("suppliers", "payables", "payments", "allocations")
            } == {"suppliers": 20, "payables": 993, "payments": 1000, "allocations": 1005},
            "未知关联净额 1000000 分": lambda r: totals(r)["missing_payable_link_net"] == 1000000,
            "已知关联净额 266249329 分": lambda r: totals(r)["known_payable_link_net"] == 266249329,
            "主表净额 267249329 分": lambda r: totals(r)["master_net"] == 267249329,
            "全部勾稽检查通过": lambda r: all(totals(r)["checks"].values()),
            "normalized_equals_report": lambda r: totals(r)["normalized_equals_report"] is True,
            "标准数据无“元”金额字段": no_yuan_amount_fields,
            "金额标准值均为 int": all_amounts_int,
            "mode=DEVELOPMENT_TEST_ONLY 且未允许下一步": lambda r: r.result["mode"] == "DEVELOPMENT_TEST_ONLY"
            and r.result["next_step_allowed"] is False,
            "来源哈希与输入一致": lambda r: r.result["source"]["input_sha256"] == sha256(r.input_path),
        },
    )

    # ---------- 2. 14 份质量有效 fixture ----------
    for case_id in CONVERTIBLE_FIXTURES:
        result = suite.evaluate(case_id, FIXTURE_DIR / f"{case_id}.xlsx")
        checks = {
            "全部勾稽检查通过": lambda r: all(totals(r)["checks"].values()),
            "标准数据无“元”金额字段": no_yuan_amount_fields,
            "金额标准值均为 int": all_amounts_int,
            "checker 退出 0": lambda r: r.checker_exit == 0,
        }
        if case_id == "C01":
            checks["C01 净额 500000 分"] = lambda r: totals(r)["master_net"] == 500000
        if case_id == "C03":
            checks["C03 净额 1000000 分（一笔付款分摊到两张应付）"] = lambda r: totals(r)["master_net"] == 1000000
        suite.judge(result, outcome="normalized", expected_block=False, checks=checks)

    # ---------- 3. 8 份质量阻断 fixture ----------
    for case_id in BLOCKED_FIXTURES:
        result = suite.evaluate(case_id, FIXTURE_DIR / f"{case_id}.xlsx")
        checks = {
            "明确阻断且未产出标准数据": lambda r: r.outcome == "blocked" and r.result is None,
            "没有 normalized.json": lambda r: not (r.case_dir / "normalized.json").exists(),
            "保留原质量问题": lambda r: bool(r.report)
            and (r.report.get("issue_counts_by_code") or {}) != {},
            "checker 退出 1": lambda r: r.checker_exit == 1,
        }
        if case_id == "C14c":
            checks["C14c 仍报 AMOUNT_INVALID（未补零）"] = lambda r: bool(
                r.report and "AMOUNT_INVALID" in r.report.get("issue_counts_by_code", {})
            )
        if case_id == "C14e":
            checks["C14e 仍报 AMOUNT_SCALE_INVALID（未四舍五入）"] = lambda r: bool(
                r.report and "AMOUNT_SCALE_INVALID" in r.report.get("issue_counts_by_code", {})
            )
        suite.judge(result, outcome="blocked", expected_block=True, checks=checks)

    # 保留行为的定点断言（空／未知 ap_id、对象差异、退款关联、非结算）
    def find_by(records: list[dict], field: str, value) -> dict | None:
        for record in records:
            if record.get(field) == value:
                return record
        return None

    c07 = next(r for r in suite.results if r.case == "C07")
    suite.judge(
        c07, outcome="normalized", expected_block=False,
        checks={"C07 保留空 ap_id": lambda r: r.data_of("allocations")[0]["ap_id"] is None},
    )
    c08 = next(r for r in suite.results if r.case == "C08")
    suite.judge(
        c08, outcome="normalized", expected_block=False,
        checks={"C08 保留台账外 ap_id 原值": lambda r: r.data_of("allocations")[0]["ap_id"] == "AP_MISSING"},
    )
    c09 = next(r for r in suite.results if r.case == "C09")
    suite.judge(
        c09, outcome="normalized", expected_block=False,
        checks={
            "C09 保留付款对象差异": lambda r: r.data_of("payments")[0]["payee_supplier_id"] == "S002"
            and r.data_of("payables")[0]["supplier_id"] == "S001",
            "C09 保留 R04 相关分摊金额": lambda r: r.data_of("allocations")[0]["allocated_amount_cents"] == 500000,
        },
    )
    c06 = next(r for r in suite.results if r.case == "C06")
    suite.judge(
        c06, outcome="normalized", expected_block=False,
        checks={
            "C06 保留退款关联与原分摊引用": lambda r: any(
                record["original_payment_id"] == "P02" and record["transaction_type"] == "REFUND"
                for record in r.data_of("payments")
            ) and any(
                record["original_allocation_id"] == "A02" for record in r.data_of("allocations")
            ),
            "C06 退款金额仍为正数": lambda r: all(
                record["amount_cents"] > 0 for record in r.data_of("payments")
            ),
        },
    )
    c10 = next(r for r in suite.results if r.case == "C10")
    suite.judge(
        c10, outcome="normalized", expected_block=False,
        checks={
            "C10 保留三种非结算状态": lambda r: sorted(
                record["settlement_status"] for record in r.data_of("payments")
            ) == ["FAILED", "PENDING", "VOIDED"],
            "C10 已结算资金为零": lambda r: totals(r)["settled_pay_master_total"] == 0,
        },
    )

    # ---------- 4. 开发小样 ----------
    # S1 标识保真
    def build_s1(path: Path) -> None:
        shutil.copyfile(FIXTURE_DIR / "C01.xlsx", path)
        workbook = load_workbook(path)
        payables = workbook["payables"]
        payables["A2"] = " AP01 "
        payables["B2"] = "NULL"
        payables.append(["AP0002", "#N/A", "S001", "2026-09-05", "100.00", "CNY"])
        payables["B3"].data_type = "s"  # 真正文本 #N/A，不是错误单元格
        payments = workbook["payments"]
        payments["A2"] = "0001"
        payments["I2"] = "N/A"
        workbook["allocations"]["B2"] = "0001"
        workbook["allocations"]["C2"] = " AP01 "
        workbook.save(path)

    s1 = suite.evaluate("S1_identifier_fidelity", None, build=build_s1)
    suite.judge(
        s1, outcome="normalized", expected_block=False,
        checks={
            "首尾空格 ID 原样保留": lambda r: r.data_of("payables")[0]["ap_id"] == " AP01 "
            and r.data_of("allocations")[0]["ap_id"] == " AP01 ",
            "前导零付款 ID 保留": lambda r: r.data_of("payments")[0]["payment_id"] == "0001"
            and r.data_of("allocations")[0]["payment_id"] == "0001",
            "文字 NULL 保留为文本": lambda r: r.data_of("payables")[0]["ap_number"] == "NULL",
            "真正文本 #N/A 保留为文本": lambda r: find_by(r.data_of("payables"), "ap_id", "AP0002")["ap_number"] == "#N/A",
            "文字 N/A 保留为银行流水文本": lambda r: r.data_of("payments")[0]["bank_reference"] == "N/A",
            "单元格来源记录编码为 str": lambda r: r.table("payables")[0]["cells"]["ap_id"]["encoding"] == "str"
            and r.table("payables")[0]["cells"]["ap_id"]["raw_value"] == " AP01 ",
            "真正文本 #N/A 的来源类型为 s": lambda r: [
                rec for rec in r.table("payables") if rec["data"]["ap_id"] == "AP0002"
            ][0]["cells"]["ap_number"]["excel_data_type"] == "s",
        },
    )

    # S2 精确小额（混合数值／文本）
    def build_s2(path: Path) -> None:
        build_from_spec(
            path,
            {
                "suppliers": [{"supplier_id": "S001", "supplier_name": "甲供应商"}],
                "payables": [
                    {"ap_id": "AP01", "ap_number": "AP01", "supplier_id": "S001",
                     "ap_date": "2026-09-05", "payable_amount": "0.30", "currency": "CNY"}
                ],
                "payments": [
                    {"payment_id": "PM01", "transaction_type": "PAY", "original_payment_id": None,
                     "payee_supplier_id": "S001", "payment_date": "2026-09-10", "amount": 0.1,
                     "currency": "CNY", "settlement_status": "SETTLED", "bank_reference": "B001"},
                    {"payment_id": "PM02", "transaction_type": "PAY", "original_payment_id": None,
                     "payee_supplier_id": "S001", "payment_date": "2026-09-11", "amount": "0.20",
                     "currency": "CNY", "settlement_status": "SETTLED", "bank_reference": "B002"},
                ],
                "allocations": [
                    {"allocation_id": "AL01", "payment_id": "PM01", "ap_id": "AP01",
                     "allocated_amount": 0.1, "original_allocation_id": None},
                    {"allocation_id": "AL02", "payment_id": "PM02", "ap_id": "AP01",
                     "allocated_amount": "0.20", "original_allocation_id": None},
                ],
            },
        )

    s2 = suite.evaluate("S2_small_amounts", None, build=build_s2)
    suite.judge(
        s2, outcome="normalized", expected_block=False,
        checks={
            "应付 0.30 元 = 30 分": lambda r: r.data_of("payables")[0]["payable_amount_cents"] == 30,
            "两笔付款 10 与 20 分": lambda r: money(r.data_of("payments"), "amount_cents") == [10, 20],
            "两条分摊 10 与 20 分": lambda r: money(r.data_of("allocations"), "allocated_amount_cents") == [10, 20],
            "净额精确为 30 分": lambda r: totals(r)["master_net"] == 30
            and totals(r)["known_payable_link_net"] == 30,
            "标准数据无浮点金额字段": lambda r: no_yuan_amount_fields(r) and all_amounts_int(r),
            "数值金额来源编码为 float_repr": lambda r: [
                rec for rec in r.table("payments") if rec["data"]["payment_id"] == "PM01"
            ][0]["cells"]["amount"]["encoding"] == "float_repr",
            "文本金额来源编码为 str": lambda r: [
                rec for rec in r.table("payments") if rec["data"]["payment_id"] == "PM02"
            ][0]["cells"]["amount"]["encoding"] == "str",
        },
    )

    # S3 日期、空白行与列顺序
    def build_s3(path: Path) -> None:
        shutil.copyfile(FIXTURE_DIR / "C01.xlsx", path)
        workbook = load_workbook(path)
        workbook["payables"]["D2"] = date(2026, 9, 5)
        workbook["payments"]["E2"] = datetime(2026, 9, 10, 0, 0, 0)
        # suppliers：S001 / 完全空白中间行 / S002，验证空白行与原行号跟随实际位置
        del workbook["suppliers"]
        suppliers = workbook.create_sheet("suppliers", 0)
        suppliers.append(SHEET_COLUMNS["suppliers"])
        suppliers.append(["S001", "甲供应商"])
        suppliers.append([None, None])
        suppliers.append(["S002", "乙收款对象"])
        # payments：重排列顺序（保持唯一数据行，避免主键重复）
        worksheet = workbook["payments"]
        rows = [[cell.value for cell in row] for row in worksheet.iter_rows(min_row=2)]
        reordered = ["payment_id", "transaction_type", "original_payment_id", "payee_supplier_id",
                     "payment_date", "currency", "amount", "settlement_status", "bank_reference"]
        position = workbook.sheetnames.index("payments")
        del workbook["payments"]
        worksheet = workbook.create_sheet("payments", position)
        worksheet.append(reordered)
        source_index = {name: SHEET_COLUMNS["payments"].index(name) for name in reordered}
        for row in rows:
            worksheet.append([row[source_index[name]] for name in reordered])
        workbook.save(path)

    s3 = suite.evaluate("S3_dates_blank_row_column_order", None, build=build_s3)
    suite.judge(
        s3, outcome="normalized", expected_block=False,
        checks={
            "原生日期转 ISO": lambda r: r.data_of("payables")[0]["ap_date"] == "2026-09-05",
            "午夜日期时间转 ISO": lambda r: r.data_of("payments")[0]["payment_date"] == "2026-09-10",
            "空白行被单独记录": lambda r: r.result["blank_rows"]["suppliers"] == [3],
            "空白行之后行号跟随实际位置": lambda r: [rec["source"]["excel_row"] for rec in r.table("suppliers")] == [2, 4],
            "列顺序变化后列号跟随实际位置": lambda r: r.table("payments")[0]["cells"]["amount"]["column_index"] == 7
            and r.table("payments")[0]["cells"]["currency"]["column_index"] == 6,
            "标准数据未丢行": lambda r: r.result["row_accounting"]["suppliers"]["normalized_rows"] == 2
            and r.result["row_accounting"]["payments"]["normalized_rows"] == 1,
            "勾稽仍成立": lambda r: all(totals(r)["checks"].values()),
        },
    )

    # S4 输入与报告错配
    mismatch_source = FIXTURE_DIR / "C02.xlsx"
    s4 = suite.evaluate(
        "S4_input_report_mismatch",
        mismatch_source,
        report_override=main_result.report,
        extra={"mismatch": True},
    )
    suite.judge(
        s4, outcome="blocked", expected_block=True,
        checks={
            "明确拒绝且不返回标准数据": lambda r: r.result is None,
            "没有 normalized.json": lambda r: not (r.case_dir / "normalized.json").exists(),
            "原因是报告与输入不匹配": lambda r: "不匹配" in r.message,
        },
    )

    # S5 报告不可用
    for label, override in (("null", None), ("empty_object", {})):
        s5 = suite.evaluate(
            f"S5_report_unusable_{label}",
            FIXTURE_DIR / "C01.xlsx",
            report_override=override,
        )
        suite.judge(
            s5, outcome="blocked", expected_block=True,
            checks={
                "未因默认值变成成功": lambda r: r.result is None,
                "没有 normalized.json": lambda r: not (r.case_dir / "normalized.json").exists(),
            },
        )
    c14a_result = next(r for r in suite.results if r.case == "C14a")
    s5_blocked = suite.evaluate(
        "S5_report_quality_blocked",
        FIXTURE_DIR / "C14a.xlsx",
        report_override=c14a_result.report,
    )
    suite.judge(
        s5_blocked, outcome="blocked", expected_block=True,
        checks={"质量阻断报告不被转换": lambda r: "result" in r.message or "issues" in r.message or "问题" in r.message},
    )

    # S6 转换后输入哈希读取失败 / 变化（只针对输入文件的第 2 次读取，即转换后的核对）
    def hash_read_fails_for_input(path_str: str, occurrence: int):
        def factory(_original):
            state = {"n": 0}

            def patched(path):
                if str(path) == path_str:
                    state["n"] += 1
                    if state["n"] == occurrence:
                        raise OSError("模拟：转换后读取输入哈希失败")
                return _original(path)

            return patched

        return factory

    def hash_read_changes_for_input(path_str: str, occurrence: int):
        def factory(_original):
            state = {"n": 0}

            def patched(path):
                if str(path) == path_str:
                    state["n"] += 1
                    if state["n"] == occurrence:
                        return b"tampered-bytes"
                return _original(path)

            return patched

        return factory

    for case, maker, expect_text in (
        ("S6_after_hash_read_failure", hash_read_fails_for_input, "无法读取输入哈希"),
        ("S6_after_hash_changed", hash_read_changes_for_input, "发生变化"),
    ):
        target_input = suite.cases_dir / case / "input.xlsx"
        s6 = suite.evaluate(
            case,
            FIXTURE_DIR / "C01.xlsx",
            patches={"read_file_bytes": maker(str(target_input), 2)},
        )
        suite.judge(
            s6, outcome="error", expected_block=False,
            checks={
                "报运行错误且不返回成功": lambda r, t=expect_text: r.result is None and t in r.message,
                "没有 normalized.json": lambda r: not (r.case_dir / "normalized.json").exists(),
            },
        )

    # ---------- 6. R1 新增：报告对照遗漏与类型错误 ----------
    def injected(case: str, source: Path, base_report: dict, mutate, note: str) -> CaseResult:
        payload = copy.deepcopy(base_report)
        mutate(payload)
        outcome = suite.evaluate(case, source, report_override=payload)
        (outcome.case_dir / "injected_report.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        (outcome.case_dir / "injection_note.txt").write_text(note + "\n", encoding="utf-8")
        return outcome

    c01_case = next(r for r in suite.results if r.case == "C01")
    c01_report = c01_case.report
    main_report = main_result.report

    p01 = injected(
        "R1_P01_report_groups_missing", FIXTURE_DIR / "C01.xlsx", c01_report,
        lambda p: p["amount_summary"].pop("payments_by_type_and_status", None),
        "内存注入：删除 amount_summary.payments_by_type_and_status（原始 checker 报告保留在 checker_report.json）",
    )
    suite.judge(
        p01, outcome="blocked", expected_block=True,
        checks={
            "明确阻断且不抛裸异常": lambda r: r.result is None,
            "原因为字段缺失": lambda r: "payments_by_type_and_status" in r.message and "缺失" in r.message,
            "没有 normalized.json": lambda r: not (r.case_dir / "normalized.json").exists(),
        },
    )

    p02 = injected(
        "R1_P02_report_groups_empty_but_actual", FIXTURE_DIR / "C01.xlsx", c01_report,
        lambda p: p["amount_summary"].__setitem__("payments_by_type_and_status", {}),
        "内存注入：把 payments_by_type_and_status 置为空对象，但 C01 实际有一笔已结算 PAY",
    )
    suite.judge(
        p02, outcome="error", expected_block=False,
        checks={
            "报运行错误且不返回成功": lambda r: r.result is None,
            "原因包含报告漏掉实际分组": lambda r: "漏掉实际付款分组" in r.message
            or any("漏掉实际付款分组" in item for item in r.reasons),
            "没有 normalized.json": lambda r: not (r.case_dir / "normalized.json").exists(),
        },
    )

    p03 = injected(
        "R1_P03_report_complete_master_net_plus1", FIXTURE_DIR / "C01.xlsx", c01_report,
        lambda p: p["reconciliation"]["complete_totals_cents"].__setitem__(
            "master_net", p["reconciliation"]["complete_totals_cents"]["master_net"] + 1
        ),
        "内存注入：完整净额 master_net 加 1 分",
    )
    suite.judge(
        p03, outcome="error", expected_block=False,
        checks={
            "报运行错误且不返回成功": lambda r: r.result is None,
            "原因包含完整净额差异": lambda r: "完整净额 master_net" in r.message
            or any("完整净额 master_net" in item for item in r.reasons),
        },
    )

    p04 = injected(
        "R1_P04_report_detail_not_object", FIXTURE_DIR / "C01.xlsx", c01_report,
        lambda p: p["reconciliation"]["diagnostic_known_valid_subtotals_cents"].__setitem__(
            "missing_payable_link_detail", [1, 2]
        ),
        "内存注入：missing_payable_link_detail 换成非空数组",
    )
    suite.judge(
        p04, outcome="blocked", expected_block=True,
        checks={
            "明确阻断且不抛裸异常": lambda r: r.result is None,
            "原因为类型不符": lambda r: "missing_payable_link_detail" in r.message and "必须是对象" in r.message,
        },
    )

    # P05—P08：文件操作错误必须按约定返回 NormalizationError
    def read_bytes_fails_for_path_occurrence(path_str: str, occurrence: int):
        def factory(_original):
            state = {"n": 0}

            def patched(path):
                if str(path) == path_str:
                    state["n"] += 1
                    if state["n"] == occurrence:
                        raise PermissionError(f"模拟：读取 {Path(path_str).name} 失败")
                return _original(path)

            return patched

        return factory

    def workbook_open_fails(_original):
        def patched(*args, **kwargs):
            raise PermissionError("模拟：打开工作簿失败")

        return patched

    file_error_cases = (
        ("R1_P05_initial_input_read_failure", str(suite.cases_dir / "R1_P05_initial_input_read_failure" / "input.xlsx"), 1,
         "read_input_hash", {"read_file_bytes": None}),
        ("R1_P06_open_workbook_failure", None, None, "open_workbook", {"load_workbook": workbook_open_fails}),
        ("R1_P07_spec_read_failure", str(PROJECT_ROOT / "docs" / "SPEC_V1.md"), 1, "read_spec_hash",
         {"read_file_bytes": None}),
        ("R1_P08_checker_read_failure", str(CHECKER), 1, "read_checker_hash", {"read_file_bytes": None}),
    )
    for case, path_str, occurrence, stage, patches in file_error_cases:
        if patches.get("read_file_bytes") is None and path_str is not None:
            patches = {"read_file_bytes": read_bytes_fails_for_path_occurrence(path_str, occurrence)}
        p = suite.evaluate(case, FIXTURE_DIR / "C01.xlsx", patches=patches)
        suite.judge(
            p, outcome="error", expected_block=False,
            checks={
                "为约定运行错误且无标准结果": lambda r: r.result is None,
                "保留阶段": lambda r, s=stage: s in r.message,
                "保留路径与原异常": lambda r, s=stage: any(s in item for item in r.reasons)
                and any("PermissionError" in item for item in r.reasons),
                "没有 normalized.json": lambda r: not (r.case_dir / "normalized.json").exists(),
            },
        )

    # ---------- 7. R1 新增：相邻防回归 ----------
    # 1. 只删除 PAY/PENDING 分组，其余分组保留
    p09 = injected(
        "R1_groups_partial_omission", MAIN_INPUT, main_report,
        lambda p: p["amount_summary"]["payments_by_type_and_status"].pop("PAY/PENDING", None),
        "内存注入：主文件报告只删除 PAY/PENDING 分组，其余分组保留",
    )
    suite.judge(
        p09, outcome="error", expected_block=False,
        checks={
            "发现漏项而非只看空对象": lambda r: r.result is None
            and ("漏掉实际付款分组" in r.message or any("漏掉实际付款分组" in i for i in r.reasons)),
            "原因含具体分组": lambda r: "PAY/PENDING" in r.message or any("PAY/PENDING" in i for i in r.reasons),
        },
    )

    # 2. 完整 known／missing 净额各加 1 分；字段缺失／null 也不得成功
    for field in ("known_payable_link_net", "missing_payable_link_net"):
        case = f"R1_complete_{field}_plus1"
        p = injected(
            case, FIXTURE_DIR / "C01.xlsx", c01_report,
            lambda payload, f=field: payload["reconciliation"]["complete_totals_cents"].__setitem__(
                f, payload["reconciliation"]["complete_totals_cents"][f] + 1
            ),
            f"内存注入：完整净额 {field} 加 1 分",
        )
        suite.judge(
            p, outcome="error", expected_block=False,
            checks={
                "发现 1 分差异": lambda r, f=field: r.result is None
                and (f"完整净额 {f}" in r.message or any(f"完整净额 {f}" in i for i in r.reasons)),
            },
        )
    for label, mutate in (
        ("missing", lambda p: p["reconciliation"]["complete_totals_cents"].pop("master_net", None)),
        ("null", lambda p: p["reconciliation"]["complete_totals_cents"].__setitem__("master_net", None)),
    ):
        case = f"R1_complete_master_net_{label}"
        p = injected(
            case, FIXTURE_DIR / "C01.xlsx", c01_report, mutate,
            f"内存注入：完整净额 master_net 字段{label}",
        )
        suite.judge(
            p, outcome="blocked", expected_block=True,
            checks={
                "不得成功": lambda r: r.result is None,
                "原因为字段不可用": lambda r: "master_net" in r.message,
            },
        )

    # 3. 某付款分组只改 complete_amount_cents 差 1 分
    p10 = injected(
        "R1_group_complete_amount_plus1", FIXTURE_DIR / "C01.xlsx", c01_report,
        lambda p: p["amount_summary"]["payments_by_type_and_status"]["PAY/SETTLED"].__setitem__(
            "complete_amount_cents",
            p["amount_summary"]["payments_by_type_and_status"]["PAY/SETTLED"]["complete_amount_cents"] + 1,
        ),
        "内存注入：PAY/SETTLED 分组只把 complete_amount_cents 加 1 分",
    )
    suite.judge(
        p10, outcome="error", expected_block=False,
        checks={
            "发现分组完整金额差异": lambda r: r.result is None
            and ("完整金额" in r.message or any("完整金额" in i for i in r.reasons)),
        },
    )

    # 4. 正例：真实无资金事件时分组为空仍可转换
    def build_empty_payments(path: Path) -> None:
        shutil.copyfile(FIXTURE_DIR / "C01.xlsx", path)
        workbook = load_workbook(path)
        for sheet in ("payments", "allocations"):
            del workbook[sheet]
            worksheet = workbook.create_sheet(sheet)
            worksheet.append(SHEET_COLUMNS[sheet])
        workbook.save(path)

    positive = suite.evaluate("R1_empty_payments_positive", None, build=build_empty_payments)
    suite.judge(
        positive, outcome="normalized", expected_block=False,
        checks={
            "checker 通过": lambda r: r.checker_exit == 0,
            "报告分组为空": lambda r: r.report["amount_summary"]["payments_by_type_and_status"] == {},
            "仍然转换成功": lambda r: r.result is not None,
            "资金净额为 0": lambda r: totals(r)["master_net"] == 0
            and totals(r)["known_payable_link_net"] == 0
            and totals(r)["missing_payable_link_net"] == 0,
            "应付金额保留且不强求相等": lambda r: totals(r)["table_totals"]["payables_payable_amount"]["normalized"] == 500000,
            "勾稽检查全通过": lambda r: all(totals(r)["checks"].values()),
        },
    )

    # ---------- 8. 保护核对 ----------
    hashes_after = {str(p): sha256(p) for p in PROTECTED_FILES}
    module_hashes_after = {str(NORMALIZER): sha256(NORMALIZER), str(CHECKER): sha256(CHECKER)}
    changed = [p for p, digest in hashes_before.items() if hashes_after.get(p) != digest]
    module_changed = [p for p, digest in module_hashes_before.items() if module_hashes_after.get(p) != digest]
    integrity_problems = changed + module_changed
    if integrity_problems:
        suite.failures.append(
            {
                "case": "INTEGRITY_protected_files",
                "kind": "integrity",
                "result": "FAIL",
                "problems": [f"受保护文件在测试前后哈希不一致：{p}" for p in integrity_problems],
            }
        )
        print("[FAIL] INTEGRITY_protected_files 哈希不一致：")
        for path in integrity_problems:
            print(f"        - {path}")
    else:
        print("[PASS] INTEGRITY_protected_files 受保护文件与被验证模块哈希一致")

    counts = {
        "inputs": len(suite.results),
        "normalized": sum(1 for r in suite.results if r.outcome == "normalized"),
        "blocked": sum(1 for r in suite.results if r.outcome == "blocked"),
        "errors": sum(1 for r in suite.results if r.outcome == "error"),
        "expected_block_or_error": sum(
            1 for r in suite.results if r.outcome in ("blocked", "error")
        ),
        "failures": len(suite.failures),
        "checker_cli_calls": suite.cli_calls,
    }
    summary = {
        "task": "TASK-004",
        "out_dir": str(container),
        "normalizer": {"path": str(NORMALIZER), "sha256": sha256(NORMALIZER)},
        "checker": {"path": str(CHECKER), "sha256": sha256(CHECKER)},
        "counts": counts,
        "cases": [
            {
                "case": r.case,
                "outcome": r.outcome,
                "message": r.message,
                "checker_exit": r.checker_exit,
                "normalized_json": str(r.normalized_path) if r.normalized_path else None,
            }
            for r in suite.results
        ],
        "failures": suite.failures,
        "protected_files_count": len(PROTECTED_FILES),
        "protected_files_unchanged": not integrity_problems,
        "protected_files_changed": integrity_problems,
        "notes": [
            "只转换成功的输入写 normalized.json；阻断与运行错误不产出标准数据。",
            "质量报告由真实 checker CLI 对当前输入副本重新生成，未使用旧 checks／答案册／生成器内存数据。",
            "转换后哈希失败与变化只在内存中模拟，未修改真实源码或原文件。",
        ],
    }
    (container / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (container / "command_log.json").write_text(
        json.dumps(suite.command_log, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    print()
    print(
        f"输入 {counts['inputs']} 个：转换 {counts['normalized']}、预期阻断／错误 "
        f"{counts['expected_block_or_error']}；checker CLI 调用 {counts['checker_cli_calls']} 次；"
        f"测试失败 {counts['failures']}"
    )
    print(f"受保护文件前后哈希一致: {summary['protected_files_unchanged']}（{summary['protected_files_count']} 个）")
    print(f"汇总: {container / 'summary.json'}")
    if suite.failures:
        print("存在断言失败或保护失败，返回非零退出码（不输出总体成功结论）。")
        return 1
    print("全部预期成立。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
